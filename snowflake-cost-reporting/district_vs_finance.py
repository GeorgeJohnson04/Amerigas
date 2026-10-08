"""
One district: Snowflake data side-by-side with Finance landed cost (LIC). No outlier marking.

Uses every load for the --match district from the archived snapshots (snapshots\\*.parquet) so all
months we have are covered, de-duplicated by OrderID+Type+YearMonth (newest snapshot wins). For each
load it puts the Snowflake numbers next to the Finance "Total Cost Above Basis" (LIC) for that load's
month, and shows the pre-basis comparison (differential + freight + fees vs LIC).

Finance workbook / price index / fiscal year: same options as district_cost_check.py.

Usage (from this folder):
    py district_vs_finance.py --match "SPRINGFIELD"
    py district_vs_finance.py --match "SPRINGFIELD" --pull               # refresh snapshots first
    py district_vs_finance.py --match "SPRINGFIELD" --include-forecast   # also Type=Forcast loads
"""

import argparse
import datetime as dt
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font

import district_cost_check as lic
import zone_detail_report as base

HERE = Path(__file__).resolve().parent
P5, CENTS, USD = "0.00000", '0.0000', "#,##0.00"
BLUE, GREY = base.BLUE, "A5A5A5"
HI = "FFF2CC"


def load_all_snapshots(match):
    files = sorted(base.SNAP_DIR.glob(f"{base.SNAP_PREFIX}*.parquet"))
    if not files:
        raise SystemExit("No archived snapshots. Run with --pull first.")
    frames = []
    for f in files:
        d = pd.read_parquet(f)
        d = d[d["District"].astype(str).str.contains(match, case=False, na=False)]
        if len(d):
            d["SourceSnapshotID"] = int(pd.to_numeric(d["SnapshotID"], errors="coerce").max())
            frames.append(d)
    allrows = pd.concat(frames, ignore_index=True)
    allrows["YearMonth"] = pd.to_datetime(allrows["YearMonth"])
    allrows = (allrows.sort_values("SourceSnapshotID")
               .drop_duplicates(subset=["OrderID", "Type", "YearMonth"], keep="last"))
    print(f"{len(allrows)} {match} rows across {len(files)} snapshot(s); "
          f"months {sorted(allrows['YearMonth'].dt.strftime('%Y-%m').unique())}")
    return allrows.reset_index(drop=True)


def write(sub, fin_row, source, finance_path, match, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Side by side"
    dist = sub["District"].mode()[0]

    ws["A1"] = f"{dist}: Snowflake data vs Finance LIC"
    ws["A1"].font = Font(bold=True, size=14)
    for i, (k, v) in enumerate([
        ("Snowflake", source),
        ("Finance file", f"{Path(finance_path).name} ({fin_row['FinDistrict']}, basis {fin_row['FinBasis']})"),
        ("Built", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Loads", f"{len(sub)} ({sub['OrderDate'].min():%Y-%m-%d} to {sub['OrderDate'].max():%Y-%m-%d})"),
    ], start=3):
        base.put(ws, i, 1, k, bold=True)
        base.put(ws, i, 2, v)

    # column groups: identity | Snowflake raw | pre-basis build-up | Finance LIC | difference
    groups = [("", ["Month", "OrderDate", "Type", "OrderID", "MoveType", "Lane", "Qty", "Basis"]),
              ("Snowflake raw ($/gal)", ["Product", "Overhead", "Storage", "Transportation", "GrandTotal",
                                         "Index (Price)", "Snowflake LIC col"]),
              ("Pre-basis actual ($/gal)", ["Differential", "Freight", "Fees", "Actual above basis"]),
              ("Finance", ["Finance LIC (month)"]),
              ("Difference", ["Actual - LIC", "Variance $"])]
    flat = [c for _, cols in groups for c in cols]

    top = 8
    hdr = top + 1
    c = 1
    for title, cols in groups:
        if title:
            cell = base.put(ws, top, c, title, bold=True)
            ws.merge_cells(start_row=top, start_column=c, end_row=top, end_column=c + len(cols) - 1)
            cell.alignment = Alignment(horizontal="center")
            cell.fill = base.HDR_FILL
            cell.font = Font(bold=True, color="FFFFFF")
        c += len(cols)
    base.header(ws, hdr, 1, flat)

    src = {"Product": "Product", "Overhead": "Overhead", "Storage": "Storage",
           "Transportation": "Transportation", "GrandTotal": "GrandTotal", "Index (Price)": "Index",
           "Snowflake LIC col": "Snowflake LIC", "Differential": "Differential", "Freight": "Freight",
           "Fees": "Fees", "Actual above basis": "Actual above basis", "Finance LIC (month)": "Finance LIC",
           "Actual - LIC": "Variance vs LIC", "Variance $": "Variance $"}
    r = hdr
    for _, rec in sub.iterrows():
        r += 1
        ym = pd.to_datetime(rec["YearMonth"])
        vals = {"Month": f"{ym:%b %Y}", "OrderDate": rec["OrderDate"], "Type": rec["Type"],
                "OrderID": rec["OrderID"], "MoveType": rec["MoveType"], "Lane": rec["Lane"],
                "Qty": rec["Qty"], "Basis": rec["Basis"]}
        vals.update({k: rec[v] for k, v in src.items()})
        for j, col in enumerate(flat, start=1):
            v = vals[col]
            fmt = None
            if col in ("Product", "Overhead", "Storage", "Transportation", "GrandTotal", "Index (Price)",
                       "Snowflake LIC col", "Differential", "Freight", "Fees", "Actual above basis",
                       "Finance LIC (month)", "Actual - LIC"):
                fmt = P5
            elif col == "Qty":
                fmt = "#,##0"
            elif col == "Variance $":
                fmt = "$#,##0.00;[Red]-$#,##0.00"
            elif col == "OrderDate":
                fmt = "yyyy-mm-dd"
            color = HI if col in ("Actual above basis", "Finance LIC (month)", "Actual - LIC") else None
            base.put(ws, r, j, v, fmt=fmt, color=color)
    last = r

    # weighted-average total row
    r += 1
    base.put(ws, r, 1, "Wtd avg / total", bold=True)
    def wa(col):
        return lic.wavg(sub.rename(columns={}), col) if False else \
            float((sub[col] * sub["Qty"]).where(sub[col].notna()).sum() /
                  sub["Qty"].where(sub[col].notna()).sum()) if sub["Qty"].where(sub[col].notna()).sum() else np.nan
    for col, key in [("Qty", None), ("Product", "Product"), ("Overhead", "Overhead"), ("Storage", "Storage"),
                     ("Transportation", "Transportation"), ("GrandTotal", "GrandTotal"),
                     ("Index (Price)", "Index"), ("Differential", "Differential"), ("Freight", "Freight"),
                     ("Fees", "Fees"), ("Actual above basis", "Actual above basis"),
                     ("Finance LIC (month)", "Finance LIC"), ("Actual - LIC", "Variance vs LIC")]:
        j = flat.index(col) + 1
        if col == "Qty":
            base.put(ws, r, j, sub["Qty"].sum(), fmt="#,##0", bold=True)
        else:
            base.put(ws, r, j, wa(key), fmt=P5, bold=True, color=HI if "LIC" in col or col == "Actual above basis"
                     or col == "Actual - LIC" else None)
    base.put(ws, r, flat.index("Variance $") + 1, sub["Variance $"].sum(),
             fmt="$#,##0.00;[Red]-$#,##0.00", bold=True)

    ws.freeze_panes = ws.cell(row=hdr + 1, column=2)
    ws.auto_filter.ref = f"A{hdr}:{base.get_column_letter(len(flat))}{last}"
    widths = {"Lane": 52, "OrderDate": 12, "Type": 9, "OrderID": 10, "MoveType": 9, "Month": 10}
    for j, col in enumerate(flat, start=1):
        ws.column_dimensions[base.get_column_letter(j)].width = widths.get(col, 13)
    ws.row_dimensions[top].height = 15

    # by-month summary
    wm = wb.create_sheet("By month")
    base.header(wm, 1, 1, ["Month", "Loads", "Gallons", "Actual above basis", "Finance LIC",
                           "Actual - LIC", "Variance $", "GrandTotal (all-in)", "Snowflake LIC col"])
    order = {m: i for i, m in enumerate(lic.MONTHS)}
    months = sorted(sub["YearMonth"].dropna().unique())
    mr = 1
    for ym in months:
        g = sub[sub["YearMonth"] == ym]
        mr += 1
        def w(col):
            d = g[g[col].notna()]
            return float((d[col] * d["Qty"]).sum() / d["Qty"].sum()) if len(d) and d["Qty"].sum() else np.nan
        vals = [f"{pd.to_datetime(ym):%b %Y}", len(g), g["Qty"].sum(), w("Actual above basis"),
                w("Finance LIC"), w("Variance vs LIC"), g["Variance $"].sum(), w("GrandTotal"),
                g["Snowflake LIC"].dropna().iloc[0] if g["Snowflake LIC"].notna().any() else None]
        fmts = [None, "0", "#,##0", P5, P5, P5, "$#,##0.00;[Red]-$#,##0.00", P5, P5]
        for j, (v, f) in enumerate(zip(vals, fmts), start=1):
            base.put(wm, mr, j, v, fmt=f)
    mlast = mr
    ch = BarChart()
    ch.type, ch.title = "col", "Actual above basis vs Finance LIC by month ($/gal)"
    ch.height, ch.width = 8, 16
    ch.add_data(Reference(wm, min_col=4, max_col=5, min_row=1, max_row=mlast), titles_from_data=True)
    ch.set_categories(Reference(wm, min_col=1, min_row=2, max_row=mlast))
    ch.series[0].graphicalProperties.solidFill = BLUE
    ch.series[1].graphicalProperties.solidFill = GREY
    base._axes_on(ch)
    ch.legend.position = "b"
    wm.add_chart(ch, "K1")
    for col, wdt in {"A": 11, "B": 8, "C": 12, "D": 16, "E": 12, "F": 12, "G": 14, "H": 16, "I": 15}.items():
        wm.column_dimensions[col].width = wdt

    # finance reference
    wf = wb.create_sheet("Finance LIC")
    base.put(wf, 1, 1, f"Finance Total Cost Above Basis ($/gal) - {fin_row['FinDistrict']} "
                       f"(basis {fin_row['FinBasis']})", bold=True)
    base.header(wf, 3, 1, ["Month", "Finance LIC", "Loads"])
    used = set(sub["YearMonth"].dt.strftime("%b"))
    for k, m in enumerate(lic.MONTHS, start=4):
        on = m in used
        base.put(wf, k, 1, m, color=HI if on else None)
        base.put(wf, k, 2, fin_row.get(m), fmt=P5, color=HI if on else None)
        base.put(wf, k, 3, int((sub["YearMonth"].dt.strftime("%b") == m).sum()) or None, fmt="0")
    base.widths(wf, {"A": 10, "B": 13, "C": 8})

    try:
        wb.save(out_path)
    except PermissionError:
        out_path = out_path.with_name(f"{out_path.stem} {dt.datetime.now():%Y%m%d_%H%M%S}{out_path.suffix}")
        wb.save(out_path)
        print("Report was open in Excel - saved a copy instead.")
    return out_path


def main():
    ap = argparse.ArgumentParser(description="One district: Snowflake vs Finance landed cost (LIC), "
                                             "side by side.")
    ap.add_argument("--match", required=True,
                    help='text to find in District, e.g. "SPRINGFIELD"')
    ap.add_argument("--finance", default=str(lic.DEFAULT_FINANCE),
                    help="Finance landed-cost workbook (default: FINANCE_WORKBOOK or %(default)s)")
    ap.add_argument("--fiscal-year", type=int,
                    help="first calendar year of the Apr-Mar fiscal year (default: read from file name)")
    ap.add_argument("--price-index", default=str(lic.ARGUS),
                    help="index price parquet used to fill blank prices (default: PRICE_INDEX_FILE or "
                         "%(default)s; skipped if missing)")
    ap.add_argument("--pull", action="store_true", help="refresh snapshots from Snowflake first")
    ap.add_argument("--include-forecast", action="store_true")
    ap.add_argument("--out")
    a = ap.parse_args()
    lic.ARGUS = Path(a.price_index)

    if a.pull:
        import pull_history
        pull_history.pull_snowflake()

    raw = load_all_snapshots(a.match)
    if "Type" in raw.columns and not a.include_forecast:
        n_fc = int((raw["Type"] != "Actual").sum())
        raw = raw[raw["Type"] == "Actual"]
        if n_fc:
            print(f"Left out {n_fc} forecast rows (use --include-forecast to keep them)")
    source = (f"archived snapshots {sorted(raw['SourceSnapshotID'].unique())}; "
              f"{len(raw)} actual load(s)")

    fin, fin_name = lic.load_finance(a.finance)
    sub, _, offsets = lic.compute(raw, a.match, fin, lic.fiscal_start(a.finance, a.fiscal_year))
    fin_row = sub.iloc[0]
    name = a.match.title()
    out = Path(a.out) if a.out else HERE / f"{name} Snowflake vs Finance LIC.xlsx"
    out = write(sub, fin_row, source, a.finance, a.match, out)

    print(f"Index offsets: {offsets}; index sources: {sub['Index source'].value_counts().to_dict()}")
    for ym, g in sub.groupby(sub["YearMonth"].dt.strftime("%Y-%m")):
        d = g[g["Actual above basis"].notna()]
        aab = (d["Actual above basis"] * d["Qty"]).sum() / d["Qty"].sum() if len(d) else float("nan")
        flic = g["Finance LIC"].dropna().iloc[0] if g["Finance LIC"].notna().any() else float("nan")
        print(f"  {ym}: {len(g)} loads, {g['Qty'].sum():,.0f} gal | actual above basis {aab:.5f} | "
              f"Finance LIC {flic:.5f} | diff {aab - flic:+.5f} | Snowflake LIC col "
              f"{g['Snowflake LIC'].dropna().iloc[0] if g['Snowflake LIC'].notna().any() else float('nan'):.5f}")
    print(f"Workbook: {out}")


if __name__ == "__main__":
    main()
