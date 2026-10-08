"""
Zone Detail - irregular data report
===================================
Pulls a zone-detail landed-cost table from Snowflake (connection comes from
snowflake_connect.py in this folder) and writes "Zone Detail Irregular Data.xlsx"
with charts and highlighted rows for anything that looks irregular:

  * landed-cost rules: TotalZoneValue != Qty x GrandTotal, GrandTotal != sum of cost parts,
    zero Product/Overhead/freight, rows outside the report month, PO / TMW numbers
    shared by several orders, missing LIC (landed cost), LIC that differs within a district-month
  * numeric outliers (beyond 3 x IQR of rows with the same MoveType/State),
    stray negatives, text in number columns
  * blanks in columns that are otherwise filled
  * placeholder values (N/A, NULL, -, ?...), leading/trailing spaces,
    inconsistent spelling/case ("Zone A" vs "ZONE-A"), very rare category values
  * future / very old / invalid dates, gaps and spikes in rows per day/week/month
  * fully duplicated rows

Table / warehouse come from environment variables (placeholders used if unset):
    SNOWFLAKE_TABLE      DB.SCHEMA.TABLE to pull   (default ANALYTICS_DB.REPORTING.ZONE_DETAIL)
                         a bare name (no dots) is searched for across the account
    SNOWFLAKE_WAREHOUSE  warehouse for queries     (default ANALYTICS_WH)

Usage (from this folder):
    py zone_detail_report.py                       # pull from Snowflake + build report
    py zone_detail_report.py --limit 200000        # cap rows pulled
    py zone_detail_report.py --table MYDB.MYSCHEMA.ZONE_DETAIL
    py zone_detail_report.py --from-file zone_detail_raw.parquet   # rebuild, no pull
"""

import argparse
import datetime as dt
import os
import re
import sys
import warnings
from decimal import Decimal
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.marker import DataPoint
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# pandas 2.3 + NumPy 2.5 emit this on datetime comparisons; harmless
warnings.filterwarnings("ignore", category=DeprecationWarning, message=".*'generic' unit.*")

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))  # snowflake_connect.py lives in this folder

DEFAULT_TABLE = os.environ.get("SNOWFLAKE_TABLE", "ANALYTICS_DB.REPORTING.ZONE_DETAIL")
DEFAULT_WAREHOUSE = os.environ.get("SNOWFLAKE_WAREHOUSE", "ANALYTICS_WH")
RAW_SNAPSHOT = HERE / "zone_detail_raw.parquet"      # always the latest pull
SNAP_DIR = HERE / "snapshots"                        # every pull, kept by SnapshotID
SNAP_PREFIX = "zone_detail_snapshot_"                # snapshots\<prefix><id>_<date>.parquet
REPORT_PATH = HERE / "Zone Detail Irregular Data.xlsx"

IQR_K = 3.0                 # outlier fence = Q1 - k*IQR .. Q3 + k*IQR
RARE_SHARE = 0.001          # category values below 0.1% of rows are "rare"
BLANK_MAX_SHARE = 0.05      # blanks are irregular only if the column is >= 95% filled
DATE_FLAG_MAX_SHARE = 0.05  # future/old dates are irregular only if they are uncommon
MIN_DATE = pd.Timestamp("1990-01-01")
MAX_FLAGGED_ROWS = 20_000   # rows copied to the "Flagged Rows" sheet
MAX_LOG_ROWS = 50_000       # rows in the "Issue Log" sheet
MAX_CHARTS = 40             # per sheet
PEER_COLS = ["MoveType", "State"]  # outliers are judged against rows sharing these (if present)
PEER_MIN = 20               # smallest peer group used; smaller groups fall back to a wider one
RARE_MAX_DISTINCT = 30      # rare-value check only for columns with few distinct values
PLACEHOLDERS = {"", "N/A", "NA", "#N/A", "NULL", "NONE", "NAN", "-", "--", "?",
                "UNKNOWN", "TBD", "XXX", "0000-00-00"}
ID_NAME = re.compile(r"(ID|KEY|NUMBER|NBR|NUM|ZIP|ZIPCODE|PHONE)$", re.I)
DATE_LIKE = re.compile(r"^\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}")

# issue -> (fill colour, meaning). Order = colour priority when a cell has several issues.
ISSUES = {
    # landed-cost checks (only run when the columns exist)
    "Value != Qty x Total": ("FF0000", "TotalZoneValue differs from Qty x GrandTotal by > $1 and > 0.5%"),
    "Total != sum of parts": ("FF0000", "GrandTotal differs from Product + Overhead + Storage + Transportation"),
    "Zero cost":            ("FF7C80", "Product or Overhead is 0, or Transportation is 0 on a non-delivered move "
                                       "between different locations"),
    "Outside report month": ("9BC2E6", "YearMonth is not the month most rows belong to"),
    "Duplicate PO/order #": ("CDB4DB", "PONumber or TMWOrderNumber is used by more than one OrderID"),
    "Missing LIC":          ("FFD966", "LIC / Basis is blank"),
    "LIC varies in district": ("FFD966", "LIC differs from the rest of the District for the same YearMonth"),
    # generic checks
    "Outlier (high)":       ("FFC7CE", f"Above Q3 + {IQR_K:g} x IQR for both its State+MoveType and its MoveType"),
    "Outlier (low)":        ("FFC7CE", f"Below Q1 - {IQR_K:g} x IQR for both its State+MoveType and its MoveType"),
    "Negative value":       ("FFC7CE", "Negative in a column that is otherwise >= 0"),
    "Non-numeric text":     ("F4B084", "Text in a column that is otherwise numbers"),
    "Invalid date":         ("BDD7EE", "Could not be read as a date, or out of range (e.g. 9999-12-31)"),
    "Future date":          ("BDD7EE", "Date after today (when such dates are uncommon)"),
    "Very old date":        ("BDD7EE", f"Before {MIN_DATE:%Y-%m-%d}, often a placeholder like 1900-01-01"),
    "Placeholder value":    ("FCE4D6", "Stand-in for missing data: N/A, NULL, -, ?, UNKNOWN..."),
    "Spelling/case variant": ("FCE4D6", "Same value written differently (case, spaces, punctuation)"),
    "Leading/trailing spaces": ("FCE4D6", "Text has extra spaces at the start or end"),
    "Rare value":           ("FFF2CC", f"Category value used once or by < {RARE_SHARE:.1%} of rows"),
    "Unexpected blank":     ("FFEB9C", f"Blank in a column that is >= {1 - BLANK_MAX_SHARE:.0%} filled"),
    "Duplicate row":        ("E4DFEC", "Every column matches another row"),
}
ISSUE_ORDER = {k: i for i, k in enumerate(ISSUES)}
TEXT_ISSUES = {"Placeholder value", "Spelling/case variant", "Leading/trailing spaces", "Rare value"}
KIND_LABEL = {"numeric": "Number", "id": "ID / key", "category": "Category", "text": "Text",
              "date": "Date", "empty": "Empty"}

BLUE, RED, ORANGE = "4472C4", "C00000", "ED7D31"
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF")
_FILLS = {}


# --------------------------------------------------------------------------- loading
def _q(ident):
    return '"' + str(ident).replace('"', '""') + '"'


def _find_table(cur, name):
    """Search every database this role can see for a table/view called `name`."""
    key = re.sub(r"[^A-Za-z0-9]", "", name).upper()
    # CostZoneDetail -> COST, ZONE, DETAIL, so COST_ZONE_DETAIL / COSTZONEDETAIL both match
    parts = [p.upper() for p in re.findall(r"[A-Z][a-z]+|[A-Z]+(?![a-z])|[a-z]+|\d+", name)] or [key]
    patterns = ["%" + "%".join(parts) + "%"]
    if len(parts) > 2:
        patterns.append("%" + "%".join(parts[-2:]) + "%")  # e.g. ZONE%DETAIL without the prefix
    hits = []
    for pattern in patterns:
        for kind in ("TABLES", "VIEWS"):
            try:
                cur.execute(f"SHOW TERSE {kind} LIKE '{pattern}' IN ACCOUNT")
            except Exception as e:  # role may not allow account-wide SHOW
                print(f"  SHOW {kind} failed: {e}")
                continue
            cols = [d[0].lower() for d in cur.description]
            for row in cur.fetchall():
                r = dict(zip(cols, row))
                hits.append((r["database_name"], r["schema_name"], r["name"], kind[:-1].lower()))
        if hits:
            break
    if not hits:
        sys.exit(f"No table or view like '{name}' is visible to this role. "
                 "Pass --table DB.SCHEMA.TABLE (and set SNOWFLAKE_ROLE if needed).")
    print("  Candidates:")
    for h in hits:
        print(f"    {h[0]}.{h[1]}.{h[2]}  ({h[3]})")
    exact = [h for h in hits if re.sub(r"[^A-Za-z0-9]", "", h[2]).upper() == key]
    pick = (exact or hits)[0]
    if not exact:
        print(f"  No exact match for {name}; using the closest name.")
    elif len(exact) > 1:
        print("  Several exact matches; using the first. Override with --table DB.SCHEMA.TABLE")
    return f"{_q(pick[0])}.{_q(pick[1])}.{_q(pick[2])}"


def load_from_snowflake(table, limit, max_rows):
    if not re.fullmatch(r'[A-Za-z0-9_$."]+', table):
        sys.exit(f"Table name looks invalid: {table!r}")
    from snowflake_connect import get_connection

    print("Connecting to Snowflake (a browser window may open for SSO sign-in)...")
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("SELECT CURRENT_USER(), CURRENT_ROLE(), CURRENT_WAREHOUSE(), "
                    "CURRENT_DATABASE(), CURRENT_SCHEMA()")
        user, role, wh, db, sch = cur.fetchone()
        print(f"Connected as {user}  role={role}  warehouse={wh}  database={db}  schema={sch}")

        if not wh:  # queries need a running warehouse; use the first one this role can use
            cur.execute("SHOW WAREHOUSES")
            names = [r[0] for r in cur.fetchall()]
            if not names:
                sys.exit("No warehouse available to this role - set SNOWFLAKE_WAREHOUSE.")
            cur.execute(f"USE WAREHOUSE {_q(names[0])}")
            print(f"No default warehouse; using {names[0]}")

        if "." not in table:
            print(f"Looking for {table}...")
            table = _find_table(cur, table)
        print(f"Using {table}")

        cur.execute(f"SELECT COUNT(*) FROM {table}")
        total = cur.fetchone()[0]
        print(f"{total:,} rows in table")

        sql = f"SELECT * FROM {table}"
        if limit:
            sql += f" LIMIT {int(limit)}"
        elif total > max_rows:
            sql += f" SAMPLE ({int(max_rows)} ROWS)"
            print(f"Large table - pulling a random sample of {max_rows:,} rows "
                  "(use --max-rows to change)")
        print(f"Querying: {sql}")
        cur.execute(sql)
        df = cur.fetch_pandas_all()
    finally:
        cur.close()
        conn.close()
    return df, table, total


def archive_pull(df):
    """Save a pull as the latest copy AND under snapshots/ by SnapshotID.

    The zone-detail table only holds the current snapshot (it is replaced on each refresh and Time
    Travel keeps 1 day), so prior months exist only in these saved files.
    """
    df.to_parquet(RAW_SNAPSHOT, index=False)
    print(f"Saved latest pull to {RAW_SNAPSHOT.name}")
    if "SnapshotID" not in df.columns or df.empty:
        return None
    SNAP_DIR.mkdir(exist_ok=True)
    sid = int(pd.to_numeric(df["SnapshotID"], errors="coerce").max())
    when = pd.to_datetime(df["SnapshotInstanceDateTime"]).max() if "SnapshotInstanceDateTime" in df.columns \
        else pd.Timestamp.now()
    path = SNAP_DIR / f"{SNAP_PREFIX}{sid}_{when:%Y-%m-%d}.parquet"
    if not path.exists():
        df.to_parquet(path, index=False)
        print(f"Archived snapshot {sid} -> snapshots\\{path.name}")
    return path


def load_from_file(path):
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".parquet":
        return pd.read_parquet(p)
    # keep "N/A", "NULL" etc. as text so they can be flagged as placeholders
    if suffix in (".xlsx", ".xlsm", ".xls"):
        return pd.read_excel(p, keep_default_na=False, na_values=[""])
    return pd.read_csv(p, keep_default_na=False, na_values=[""], low_memory=False)


# --------------------------------------------------------------------------- typing
def _is_missing(v):
    try:
        return v is None or bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def _to_float(v):
    if _is_missing(v):
        return np.nan
    if isinstance(v, (int, float, np.integer, np.floating, Decimal)) and not isinstance(v, bool):
        return float(v)
    try:
        return float(str(v).strip().replace(",", "").replace("$", ""))
    except ValueError:
        return np.nan


def _naive(s):
    if isinstance(s.dtype, pd.DatetimeTZDtype):
        return s.dt.tz_convert(None)
    return s


def _to_datetime(s):
    try:
        out = pd.to_datetime(s, errors="coerce", format="mixed")
    except (TypeError, ValueError):
        out = pd.to_datetime(s, errors="coerce", utc=True)
    return _naive(out)


def _numeric_kind(name, x):
    v = x.dropna()
    if v.empty:
        return "empty"
    is_int = bool(np.all(np.mod(v, 1) == 0))
    if is_int and (ID_NAME.search(str(name)) or (len(v) > 50 and v.nunique() >= 0.95 * len(v))):
        return "id"
    if is_int and v.nunique() <= 10:
        return "category"
    return "numeric"


def classify(df):
    """Return (analysis frame, kind per column, note per column)."""
    work, kinds, notes = {}, {}, {}
    for c in df.columns:
        s = df[c]
        nn = s.dropna()
        if nn.empty:
            kinds[c], work[c] = "empty", s
        elif pd.api.types.is_bool_dtype(s):
            kinds[c], work[c] = "category", s
        elif pd.api.types.is_datetime64_any_dtype(s):
            kinds[c], work[c] = "date", _naive(s)
        elif pd.api.types.is_numeric_dtype(s):
            x = pd.Series(s.astype("Float64").to_numpy(dtype="float64", na_value=np.nan), index=s.index)
            kinds[c], work[c] = _numeric_kind(c, x), x
        elif nn.map(lambda v: isinstance(v, (dt.date, dt.datetime))).all():
            kinds[c], work[c] = "date", _to_datetime(s)
        elif nn.map(lambda v: isinstance(v, (int, float, Decimal, np.number))
                    and not isinstance(v, bool)).all():
            x = s.map(_to_float).astype("float64")
            kinds[c], work[c] = _numeric_kind(c, x), x
        else:
            txt = nn.astype(str).str.strip()
            as_num = pd.to_numeric(txt.str.replace(",", "", regex=False), errors="coerce")
            if as_num.notna().mean() >= 0.95:
                x = s.map(_to_float).astype("float64")
                kinds[c], work[c] = _numeric_kind(c, x), x
                notes[c] = "numbers stored as text"
            elif txt.head(1000).str.match(DATE_LIKE).mean() >= 0.95:
                kinds[c], work[c] = "date", _to_datetime(s)
                notes[c] = "dates stored as text"
            else:
                nunq = txt.nunique()
                kinds[c] = "category" if nunq <= 200 and nunq <= 0.5 * len(txt) else "text"
                work[c] = s
    return pd.DataFrame(work, index=df.index), kinds, notes


# --------------------------------------------------------------------------- checks
class Flags:
    def __init__(self):
        self.rows, self.cols, self.issues, self.details = [], [], [], []

    def add(self, idx, col, issue, detail=""):
        idx = list(idx)
        if not idx:
            return
        self.rows += idx
        self.cols += [col] * len(idx)
        self.issues += [issue] * len(idx)
        self.details += list(detail) if isinstance(detail, (list, pd.Series)) else [detail] * len(idx)

    def frame(self):
        return pd.DataFrame({"row": self.rows, "column": self.cols,
                             "issue": self.issues, "detail": self.details})


def fnum(v):
    if v is None or pd.isna(v):
        return ""
    a = abs(v)
    if a >= 10_000:
        return f"{v:,.0f}"
    if a >= 100:
        return f"{v:,.1f}"
    return f"{v:,.3g}"


def _fence(v, min_n=PEER_MIN):
    """(low, high) outlier fence for a set of values, or None if it can't be judged."""
    v = v.dropna()
    if len(v) < min_n:
        return None
    q1, q3 = v.quantile(0.25), v.quantile(0.75)
    iqr = q3 - q1
    if iqr > 0:
        return q1 - IQR_K * iqr, q3 + IQR_K * iqr
    med = v.median()  # mostly one value: anything else is odd only if it is rare
    return (med, med) if (v != med).mean() <= 0.01 else None


def peer_fences(x, df, peers):
    """Per-row fences from the narrowest peer group (e.g. State+MoveType) big enough to judge."""
    lo = pd.Series(np.nan, index=x.index)
    hi = pd.Series(np.nan, index=x.index)
    label = pd.Series("", index=x.index, dtype=object)
    done = pd.Series(False, index=x.index)
    for k in range(len(peers), -1, -1):
        lvl = peers[:k]
        if lvl:
            keys = [df[c].astype(str) for c in lvl]
            groups = x.groupby(keys, sort=False).groups.items()
        else:
            groups = [((), x.index)]
        for key, idx in groups:
            todo = idx[~done.loc[idx].to_numpy()]
            if not len(todo):
                continue
            f = _fence(x.loc[idx])
            if f is None:
                continue
            key = key if isinstance(key, tuple) else (key,)
            lo.loc[todo], hi.loc[todo] = f
            label.loc[todo] = ", ".join(f"{c}={v}" for c, v in zip(lvl, key)) or "all rows"
            done.loc[todo] = True
        if done.all():
            break
    return lo, hi, label


def lic_checks(df, fl):
    """Business-rule checks for the zone-detail table; each runs only if its columns exist."""
    has = lambda *cols: all(c in df.columns for c in cols)
    num = lambda c: pd.to_numeric(df[c], errors="coerce")

    if has("Qty", "GrandTotal", "TotalZoneValue"):
        exp = num("Qty") * num("GrandTotal")
        diff = num("TotalZoneValue") - exp
        bad = (diff.abs() > 1) & (diff.abs() > 0.005 * exp.abs())
        idx = df.index[bad]
        fl.add(idx, "TotalZoneValue", "Value != Qty x Total",
               [f"Qty x GrandTotal = {exp[i]:,.2f}; off by {diff[i]:+,.2f} ({diff[i] / exp[i]:+.1%})"
                for i in idx])

    parts = [c for c in ("Product", "Overhead", "Storage", "Transportation") if c in df.columns]
    if has("GrandTotal") and len(parts) == 4:
        tot = sum(num(c).fillna(0) for c in parts)
        diff = num("GrandTotal") - tot
        idx = df.index[diff.abs() > 0.0005]
        fl.add(idx, "GrandTotal", "Total != sum of parts",
               [f"parts add to {tot[i]:.5f}; off by {diff[i]:+.5f}" for i in idx])

    for c in ("Product", "Overhead"):
        if has(c):
            fl.add(df.index[num(c) == 0], c, "Zero cost", f"{c} is 0")
    if has("Transportation", "OriginLocation", "DestinationLocation"):
        site = lambda s: s.astype(str).str.split(" - ").str[0].str.replace(r"\s+RT$", "", regex=True).str.strip()
        delivered = df["OriginLocation"].astype(str).str.contains("DELIVERED", case=False)
        same = site(df["OriginLocation"]) == site(df["DestinationLocation"])
        bad = (num("Transportation") == 0) & ~delivered & ~same
        fl.add(df.index[bad], "Transportation", "Zero cost",
               "no freight on a move between different locations")

    if has("YearMonth"):
        ym = pd.to_datetime(df["YearMonth"], errors="coerce")
        main = ym.mode()
        if len(main) and (ym == main[0]).mean() >= 0.8:
            idx = df.index[ym.notna() & (ym != main[0])]
            fl.add(idx, "YearMonth", "Outside report month",
                   f"{(ym == main[0]).mean():.1%} of rows are {main[0]:%Y-%m}")

    if has("OrderID"):
        for c in ("PONumber", "TMWOrderNumber"):
            if has(c):
                s = df[c].astype("string").str.strip()
                n_orders = df.groupby(s)["OrderID"].transform("nunique")
                idx = df.index[s.notna() & (n_orders > 1)]
                fl.add(idx, c, "Duplicate PO/order #",
                       [f"{c} {s[i]} is on {n_orders[i]} orders" for i in idx])

    if has("LIC"):
        miss = df["LIC"].isna()
        if 0 < miss.mean() < 0.5:
            fl.add(df.index[miss], "LIC", "Missing LIC", "LIC is blank")
        if has("District", "YearMonth"):
            lic = num("LIC")
            mode = lic.groupby([df["District"], df["YearMonth"]]).transform(
                lambda g: g.mode().iloc[0] if g.notna().any() else np.nan)
            idx = df.index[lic.notna() & mode.notna() & ((lic - mode).abs() > 1e-6)]
            fl.add(idx, "LIC", "LIC varies in district",
                   [f"district usually {mode[i]:.4f}" for i in idx])


def find_issues(df, work, kinds, notes):
    fl, fences = Flags(), {}
    today = pd.Timestamp.today().normalize()
    peers = [c for c in PEER_COLS if c in df.columns and kinds.get(c) == "category"]
    lic_checks(df, fl)

    try:
        dup = df.duplicated(keep=False)
    except TypeError:
        dup = df.astype(str).duplicated(keep=False)
    fl.add(df.index[dup], "(whole row)", "Duplicate row", "identical to another row")

    for c in df.columns:
        kind, s, x = kinds[c], df[c], work[c]
        if kind == "empty":
            continue

        null = s.isna()
        share = null.mean()
        if 0 < share <= BLANK_MAX_SHARE:
            fl.add(s.index[null], c, "Unexpected blank", f"column is {1 - share:.1%} filled")

        if notes.get(c) == "numbers stored as text":
            bad = s.notna() & x.isna()
            fl.add(s.index[bad], c, "Non-numeric text", "expected a number")

        if kind == "numeric":
            v = x.dropna()
            f = _fence(v, min_n=1)
            if f:
                fences[c] = f  # table-wide range, shown on Summary / Distributions
            # unusual locally (e.g. same MoveType+State) AND broadly (same MoveType), so normal
            # regional differences (CA overhead, MI truck sizes) are not flagged
            lo, hi, label = peer_fences(x, df, peers)
            wlo, whi, wlabel = peer_fences(x, df, peers[:1]) if len(peers) > 1 else (lo, hi, label)

            def why(i):
                txt = f"normal for {label[i]}: {fnum(lo[i])} to {fnum(hi[i])}"
                if wlabel[i] != label[i]:
                    txt += f"; for {wlabel[i]}: {fnum(wlo[i])} to {fnum(whi[i])}"
                return txt

            for name, mask in (("Outlier (high)", (x > hi) & (x > whi)),
                               ("Outlier (low)", (x < lo) & (x < wlo))):
                idx = x.index[mask]
                fl.add(idx, c, name, [why(i) for i in idx])
            if (v < 0).any() and (v >= 0).mean() >= 0.95:
                fl.add(x.index[x < 0], c, "Negative value", f"{(v >= 0).mean():.1%} of values are >= 0")

        elif kind == "date":
            fl.add(s.index[s.notna() & x.isna()], c, "Invalid date", "not a readable date")
            fut = x > today + pd.Timedelta(days=1)
            if 0 < fut.mean() <= DATE_FLAG_MAX_SHARE:
                fl.add(x.index[fut], c, "Future date", f"after {today:%Y-%m-%d}")
            old = x < MIN_DATE
            if 0 < old.mean() <= DATE_FLAG_MAX_SHARE:
                fl.add(x.index[old], c, "Very old date", f"before {MIN_DATE:%Y-%m-%d}")

        if kind in ("category", "text"):
            st = s[s.map(lambda v: isinstance(v, str))]
            if not st.empty:
                stripped = st.str.strip()
                fl.add(st.index[st != stripped], c, "Leading/trailing spaces", "extra spaces")
                ph = stripped.str.upper().isin(PLACEHOLDERS)
                fl.add(st.index[ph], c, "Placeholder value", "stands in for missing data")
                if kind == "category":
                    real = stripped[~ph]
                    key = real.str.lower().str.replace(r"[^a-z0-9]", "", regex=True)
                    canon = real.groupby(key).agg(lambda g: g.value_counts().idxmax())
                    nvar = real.groupby(key).nunique()
                    multi = nvar.index[nvar > 1]
                    odd = key.isin(multi) & (real != key.map(canon))
                    fl.add(real.index[odd], c, "Spelling/case variant",
                           [f"usually written '{canon[k]}'" for k in key[odd]])

        if kind == "category":
            nonnull = s.dropna()
            if len(nonnull) >= 200:
                vals = nonnull.map(str)
                vc = vals.value_counts()
                if 3 <= len(vc) <= RARE_MAX_DISTINCT:
                    rare = vc[(vc / len(nonnull) < RARE_SHARE) | (vc == 1)]
                    hit = vals.isin(rare.index)
                    fl.add(vals.index[hit], c, "Rare value",
                           [f"appears {vc[v]:,}x in {len(nonnull):,} rows" for v in vals[hit]])

    return fl.frame(), fences


def build_timeline(df, work, kinds):
    dates = [c for c in df.columns if kinds[c] == "date" and work[c].nunique() >= 3]
    if not dates:
        return None
    dc = max(dates, key=lambda c: (work[c].notna().sum(), work[c].nunique()))
    x = work[dc].dropna()
    today = pd.Timestamp.today().normalize()
    sane = x[(x >= MIN_DATE) & (x <= today + pd.Timedelta(days=366))]
    if sane.nunique() < 3:
        sane = x
    span = (sane.max() - sane.min()).days
    freq, label = ("D", "day") if span <= 400 else ("W", "week") if span <= 2500 else ("M", "month")
    period = sane.dt.to_period(freq).dt.start_time
    rng = pd.period_range(sane.min(), sane.max(), freq=freq).start_time
    if len(rng) > 3000:
        return None
    counts = period.value_counts().reindex(rng, fill_value=0)

    c = counts.astype(float)
    med = c.median()
    mad = (c - med).abs().median()
    flag = pd.Series("", index=counts.index, dtype=object)
    if mad > 0:
        rz = 0.6745 * (c - med) / mad
        # must be statistically unusual AND a big move (>= 2x or <= 0.5x typical), not just noise
        flag[(rz > 3.5) & (c >= 2 * med)] = "Spike"
        flag[(rz < -3.5) & (c <= 0.5 * med) & (counts > 0)] = "Drop"
    flag[counts == 0] = "Gap (no rows)"

    numcols = [k for k in df.columns if kinds[k] == "numeric"][:12]
    means = (work.loc[sane.index, numcols].groupby(period).mean().reindex(rng)
             if numcols else pd.DataFrame(index=rng))
    return {"date_col": dc, "label": label, "counts": counts, "flag": flag, "means": means}


# --------------------------------------------------------------------------- excel helpers
def fill(color):
    if color not in _FILLS:
        _FILLS[color] = PatternFill("solid", fgColor=color)
    return _FILLS[color]


def xl(v):
    if _is_missing(v):
        return None
    if isinstance(v, np.datetime64):
        v = pd.Timestamp(v)
    if isinstance(v, pd.Timestamp):
        if v.tzinfo is not None:
            v = v.tz_convert(None)
        return v.to_pydatetime()
    if isinstance(v, dt.datetime) and v.tzinfo is not None:
        return v.replace(tzinfo=None)
    if isinstance(v, np.generic):
        return v.item()
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (int, float, bool, dt.date, dt.time)):
        return v
    return ILLEGAL_CHARACTERS_RE.sub("", str(v))[:32000]


def put(ws, r, c, v, color=None, bold=False, fmt=None):
    val = xl(v)
    cell = ws.cell(row=r, column=c, value=val)
    if isinstance(val, str) and val.startswith("="):
        cell.data_type = "s"  # never treat data as a formula
    if isinstance(val, dt.datetime):
        cell.number_format = "yyyy-mm-dd hh:mm" if (val.hour or val.minute or val.second) else "yyyy-mm-dd"
    elif isinstance(val, dt.date):
        cell.number_format = "yyyy-mm-dd"
    if fmt:
        cell.number_format = fmt
    if color:
        cell.fill = fill(color)
    if bold:
        cell.font = Font(bold=True)
    return cell


def header(ws, r, c0, names):
    for j, name in enumerate(names):
        cell = ws.cell(row=r, column=c0 + j, value=name)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def _axes_on(ch):
    ch.x_axis.delete = False
    ch.y_axis.delete = False


def bar_chart(ws, title, cat_col, val_col, r_hdr, r_last, horizontal=False, colors=None,
              x_title=None, y_title=None):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title, ch.legend, ch.gapWidth = title, None, 25
    ch.height, ch.width = 7.5, 17
    ch.add_data(Reference(ws, min_col=val_col, min_row=r_hdr, max_row=r_last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cat_col, min_row=r_hdr + 1, max_row=r_last))
    ser = ch.series[0]
    ser.graphicalProperties.solidFill = BLUE
    ser.graphicalProperties.line.solidFill = BLUE
    for i, color in (colors or {}).items():
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = color
        pt.graphicalProperties.line.solidFill = color
        ser.dPt.append(pt)
    _axes_on(ch)
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"  # first row at the top
    if x_title:
        ch.x_axis.title = x_title
    if y_title:
        ch.y_axis.title = y_title
    return ch


# --------------------------------------------------------------------------- report
def write_report(df, work, kinds, notes, flags, fences, timeline, source, out_path):
    wb = Workbook()
    n = len(df)
    flagged_rows = flags["row"].nunique() if not flags.empty else 0
    by_col = flags.groupby("column").size() if not flags.empty else pd.Series(dtype=int)
    by_issue = flags.groupby("issue").size() if not flags.empty else pd.Series(dtype=int)

    # ---------------- Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Zone Detail - irregular data report"
    ws["A1"].font = Font(bold=True, size=14)
    tl_text = "no usable date column"
    if timeline:
        f = timeline["flag"]
        tl_text = (f"{(f == 'Gap (no rows)').sum()} gaps, {(f == 'Spike').sum()} spikes, "
                   f"{(f == 'Drop').sum()} drops in rows per {timeline['label']} "
                   f"({timeline['date_col']}) - see Timeline")
    facts = [
        ("Source", source),
        ("Built", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Rows", n),
        ("Columns", df.shape[1]),
        ("Duplicate rows", int(by_issue.get("Duplicate row", 0))),
        ("Rows with at least one flag", f"{flagged_rows:,} ({flagged_rows / n:.1%})"),
        ("Flagged cells", int(len(flags))),
        ("Timeline", tl_text),
    ]
    for i, (k, v) in enumerate(facts, start=3):
        put(ws, i, 1, k, bold=True)
        put(ws, i, 2, v, fmt="#,##0" if isinstance(v, int) else None)

    r = 3 + len(facts) + 1
    put(ws, r, 1, "Colour key (used on Flagged Rows / Issue Log)", bold=True)
    r += 1
    header(ws, r, 1, ["Issue", "What it means", "Cells flagged"])
    legend_hdr = r
    for issue, (color, meaning) in ISSUES.items():
        r += 1
        put(ws, r, 1, issue, color=color)
        put(ws, r, 2, meaning)
        put(ws, r, 3, int(by_issue.get(issue, 0)), fmt="#,##0")
    legend_last = r
    if len(by_issue):
        ch = bar_chart(ws, "Flagged cells by issue type", 1, 3, legend_hdr, legend_last,
                       horizontal=True,
                       colors={i: RED for i, k in enumerate(ISSUES) if by_issue.get(k, 0) > 0})
        ch.width = 15
        ws.add_chart(ch, "E3")

    r += 3
    put(ws, r, 1, "Column profile", bold=True)
    r += 1
    prof_hdr = r
    header(ws, r, 1, ["Column", "Type", "Filled", "Blank %", "Distinct", "Min", "Max", "Mean",
                      "Median", "Normal low", "Normal high", "Flagged cells", "Flagged %", "Notes"])
    for c in df.columns:
        r += 1
        kind, s, x = kinds[c], df[c], work[c]
        nnull = int(s.isna().sum())
        try:
            distinct = int(s.nunique(dropna=True))
        except TypeError:
            distinct = int(s.astype(str).nunique())
        note = [notes[c]] if c in notes else []
        if kind == "empty":
            note.append("all blank")
        elif distinct == 1:
            note.append("constant - only one value")
        if kind == "id":
            note.append("ID-like: outlier checks skipped")
        if kind == "numeric" and c not in fences:
            note.append("too spread / mostly one value: outlier check skipped")
        stats = [None] * 4
        if kind in ("numeric", "id"):
            v = x.dropna()
            if len(v):
                stats = [v.min(), v.max(), v.mean(), v.median()]
        elif kind == "date":
            v = x.dropna()
            if len(v):
                stats = [v.min(), v.max(), None, None]
        lo, hi = fences.get(c, (None, None))
        nflag = int(by_col.get(c, 0))
        row = [c, KIND_LABEL[kind], n - nnull, nnull / n, distinct, *stats, lo, hi, nflag, nflag / n,
               "; ".join(note)]
        fmts = [None, None, "#,##0", "0.0%", "#,##0", None, None, None, None, None, None, "#,##0",
                "0.0%", None]
        for j, (v, f) in enumerate(zip(row, fmts), start=1):
            if f is None and isinstance(v, (float, np.floating)) and not isinstance(v, bool):
                f = "#,##0.00"
            put(ws, r, j, v, fmt=f)
    if r > prof_hdr:
        ws.conditional_formatting.add(f"D{prof_hdr + 1}:D{r}",
                                      ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                     end_type="num", end_value=1, end_color="F8696B"))
        ws.conditional_formatting.add(f"L{prof_hdr + 1}:L{r}", DataBarRule(start_type="min", end_type="max",
                                                                         color="F8696B"))
    ws.freeze_panes = "A2"
    widths(ws, {"A": 30, "B": 48, "C": 13, "D": 9, "E": 10, "F": 13, "G": 13, "H": 12, "I": 12,
                "J": 12, "K": 12, "L": 13, "M": 10, "N": 50})

    # ---------------- Issues by column (pivot + stacked chart)
    wi = wb.create_sheet("Issues by Column")
    if flags.empty:
        put(wi, 1, 1, "No irregular data found.", bold=True)
    else:
        issue_cols = [k for k in ISSUES if k in by_issue.index]
        pv = (flags.pivot_table(index="column", columns="issue", values="row", aggfunc="count", fill_value=0)
              .reindex(columns=issue_cols, fill_value=0))
        pv["Total"] = pv.sum(axis=1)
        pv = pv.sort_values("Total", ascending=False)
        header(wi, 1, 1, ["Column", *issue_cols, "Total"])
        for i, (col, vals) in enumerate(pv.iterrows(), start=2):
            put(wi, i, 1, col)
            for j, v in enumerate(vals, start=2):
                put(wi, i, j, int(v) or None, fmt="#,##0")
        last = len(pv) + 1
        ch = BarChart()
        ch.type, ch.grouping, ch.overlap = "bar", "stacked", 100
        ch.title = "Flagged cells by column and issue"
        ch.height = max(8, min(30, 0.6 * len(pv) + 3))
        ch.width = 22
        ch.add_data(Reference(wi, min_col=2, max_col=1 + len(issue_cols), min_row=1, max_row=last),
                    titles_from_data=True)
        ch.set_categories(Reference(wi, min_col=1, min_row=2, max_row=last))
        palette = ["C00000", "FF6F61", "F4B084", "ED7D31", "5B9BD5", "2E75B6", "9DC3E6", "FFC000",
                   "A9D18E", "70AD47", "BF9000", "FFD966", "8064A2"]
        for i, ser in enumerate(ch.series):
            ser.graphicalProperties.solidFill = palette[i % len(palette)]
        _axes_on(ch)
        ch.x_axis.scaling.orientation = "maxMin"
        wi.add_chart(ch, f"{get_column_letter(len(issue_cols) + 4)}2")
        wi.freeze_panes = "B2"
        widths(wi, {"A": 30})

    # ---------------- Flagged rows (original data, irregular cells coloured)
    wf = wb.create_sheet("Flagged Rows")
    if flags.empty:
        put(wf, 1, 1, "No irregular rows found.", bold=True)
    else:
        cell_issues = flags.groupby(["row", "column"])["issue"].agg(list).to_dict()
        row_issues = flags.groupby("row")["issue"].agg(
            lambda s: ", ".join(sorted(set(s), key=ISSUE_ORDER.get)))
        rows = sorted(row_issues.index)
        truncated = len(rows) > MAX_FLAGGED_ROWS
        rows = rows[:MAX_FLAGGED_ROWS]
        cols = list(df.columns)
        header(wf, 1, 1, ["Row #", "Issues", *cols])
        for i, ridx in enumerate(rows, start=2):
            dup = (ridx, "(whole row)") in cell_issues
            put(wf, i, 1, int(ridx) + 1, color=ISSUES["Duplicate row"][0] if dup else None)
            put(wf, i, 2, row_issues[ridx])
            rec = df.iloc[ridx]
            for j, c in enumerate(cols, start=3):
                iss = cell_issues.get((ridx, c))
                color = ISSUES[min(iss, key=ISSUE_ORDER.get)][0] if iss else None
                put(wf, i, j, rec[c], color=color)
        if truncated:
            put(wf, len(rows) + 3, 1, f"Showing the first {MAX_FLAGGED_ROWS:,} of "
                                      f"{len(row_issues):,} flagged rows.", bold=True)
        wf.freeze_panes = "C2"
        wf.auto_filter.ref = f"A1:{get_column_letter(len(cols) + 2)}{len(rows) + 1}"
        widths(wf, {"A": 8, "B": 34})
        for j in range(3, len(cols) + 3):
            wf.column_dimensions[get_column_letter(j)].width = 16

    # ---------------- Issue log
    wl = wb.create_sheet("Issue Log")
    header(wl, 1, 1, ["Row #", "Column", "Value", "Issue", "Detail"])
    log = flags.sort_values(["row", "column"]).head(MAX_LOG_ROWS) if not flags.empty else flags
    for i, rec in enumerate(log.itertuples(index=False), start=2):
        put(wl, i, 1, int(rec.row) + 1)
        put(wl, i, 2, rec.column)
        if rec.column in df.columns:
            raw = df.at[rec.row, rec.column]
            put(wl, i, 3, f"'{raw}'" if isinstance(raw, str) and raw != raw.strip() else raw)
        put(wl, i, 4, rec.issue, color=ISSUES[rec.issue][0])
        put(wl, i, 5, rec.detail)
    if len(flags) > MAX_LOG_ROWS:
        put(wl, len(log) + 3, 1, f"Showing the first {MAX_LOG_ROWS:,} of {len(flags):,} issues.", bold=True)
    wl.freeze_panes = "A2"
    wl.auto_filter.ref = f"A1:E{max(len(log), 1) + 1}"
    widths(wl, {"A": 8, "B": 28, "C": 26, "D": 24, "E": 44})

    # ---------------- Distributions (numeric columns)
    numcols = [c for c in df.columns if kinds[c] == "numeric"][:MAX_CHARTS]
    if numcols:
        wd = wb.create_sheet("Distributions")
        put(wd, 1, 1, "Histogram of each number column. Red bars = outside the table-wide range "
                      f"(Q1 - {IQR_K:g} x IQR to Q3 + {IQR_K:g} x IQR). Rows are only flagged as outliers "
                      "against similar rows (same MoveType/State), so e.g. railcar volumes are not "
                      "flagged just for being bigger than trucks - see Issue Log.", bold=True)
        r = 3
        for c in numcols:
            v = work[c].dropna()
            lo, hi = fences.get(c, (v.min(), v.max()))
            a, b = max(v.min(), lo), min(v.max(), hi)
            inner = v[(v >= lo) & (v <= hi)]
            bins = []
            below, above = int((v < lo).sum()), int((v > hi).sum())
            if below:
                bins.append((f"< {fnum(lo)} (low outliers)", below, True))
            if b > a:
                cnt, edges = np.histogram(inner, bins=20, range=(a, b))
                bins += [(f"{fnum(edges[i])} to {fnum(edges[i + 1])}", int(cnt[i]), False)
                         for i in range(len(cnt))]
            else:
                bins.append((f"= {fnum(a)}", len(inner), False))
            if above:
                bins.append((f"> {fnum(hi)} (high outliers)", above, True))
            header(wd, r, 1, [f"{c} range", "Rows"])
            for i, (lab, cnt, out) in enumerate(bins, start=1):
                put(wd, r + i, 1, lab, color="FFC7CE" if out else None)
                put(wd, r + i, 2, cnt, fmt="#,##0")
            ch = bar_chart(wd, f"{c}  (normal range {fnum(lo)} to {fnum(hi)})", 1, 2, r, r + len(bins),
                           colors={i: RED for i, (_, _, out) in enumerate(bins) if out}, y_title="Rows")
            wd.add_chart(ch, f"D{r}")
            r += max(len(bins) + 2, 17) + 1
        widths(wd, {"A": 30, "B": 10})

    # ---------------- Categories
    catcols = [c for c in df.columns if kinds[c] == "category"][:MAX_CHARTS]
    if catcols:
        wc = wb.create_sheet("Categories")
        put(wc, 1, 1, "Most common values per category column, plus every flagged value. "
                      "Orange bars = flagged (rare, misspelled/case variant, placeholder, extra spaces).",
            bold=True)
        r = 3
        for c in catcols:
            s = df[c]
            vals = s.dropna().map(str)
            vc = vals.value_counts()
            flagged_vals = set()
            if not flags.empty:
                fr = flags[(flags["column"] == c) & flags["issue"].isin(TEXT_ISSUES)]["row"]
                flagged_vals = set(df.loc[fr.unique(), c].dropna().map(str))
            keep = list(vc.index[:10]) + [k for k in vc.index if k in flagged_vals and k not in vc.index[:10]][:10]
            keep = sorted(keep, key=lambda k: -vc[k])
            bins = [(f"'{k}'" if k != k.strip() or k == "" else k, int(vc[k]), k in flagged_vals) for k in keep]
            other = int(vc.drop(keep).sum())
            if other:
                bins.append((f"(all {len(vc) - len(keep):,} other values)", other, False))
            if s.isna().any():
                bins.append(("(blank)", int(s.isna().sum()), s.isna().mean() <= BLANK_MAX_SHARE))
            header(wc, r, 1, [f"{c} value", "Rows"])
            for i, (lab, cnt, bad) in enumerate(bins, start=1):
                put(wc, r + i, 1, lab, color="FCE4D6" if bad else None)
                put(wc, r + i, 2, cnt, fmt="#,##0")
            ch = bar_chart(wc, f"{c}  ({len(vc):,} distinct values)", 1, 2, r, r + len(bins), horizontal=True,
                           colors={i: ORANGE for i, (_, _, bad) in enumerate(bins) if bad})
            ch.height = max(7.5, 0.55 * len(bins) + 2)
            wc.add_chart(ch, f"D{r}")
            r += max(len(bins) + 2, int(ch.height / 0.5) + 1) + 1
        widths(wc, {"A": 34, "B": 10})

    # ---------------- Timeline
    if timeline:
        wt = wb.create_sheet("Timeline")
        counts, flag, means = timeline["counts"], timeline["flag"], timeline["means"]
        unit = timeline["label"]
        mcols = list(means.columns)
        header(wt, 1, 1, [f"{unit.title()} starting", "Rows", "Flag", "Flagged rows",
                          *[f"Avg {m}" for m in mcols]])
        for i, p in enumerate(counts.index, start=2):
            bad = bool(flag[p])
            put(wt, i, 1, p.date(), color="FFC7CE" if bad else None)
            put(wt, i, 2, int(counts[p]), fmt="#,##0")
            put(wt, i, 3, flag[p] or None, color="FFC7CE" if bad else None)
            put(wt, i, 4, int(counts[p]) if bad else None, fmt="#,##0")
            for j, m in enumerate(mcols, start=5):
                put(wt, i, j, means.at[p, m], fmt="#,##0.00")
        last = len(counts) + 1
        anchor_col = get_column_letter(len(mcols) + 6)

        ch = LineChart()
        ch.title = f"Rows per {unit} by {timeline['date_col']} (red = gap / spike / drop)"
        ch.height, ch.width = 8, 26
        ch.add_data(Reference(wt, min_col=2, min_row=1, max_row=last), titles_from_data=True)
        ch.add_data(Reference(wt, min_col=4, min_row=1, max_row=last), titles_from_data=True)
        ch.set_categories(Reference(wt, min_col=1, min_row=2, max_row=last))
        s0, s1 = ch.series
        s0.smooth = False
        s0.graphicalProperties.line.solidFill = BLUE
        s1.smooth = False
        s1.marker.symbol, s1.marker.size = "circle", 7
        s1.marker.graphicalProperties.solidFill = RED
        s1.marker.graphicalProperties.line.solidFill = RED
        s1.graphicalProperties.line.noFill = True
        ch.x_axis.number_format = "yyyy-mm-dd"
        ch.legend.position = "b"
        _axes_on(ch)
        wt.add_chart(ch, f"{anchor_col}2")

        for k, m in enumerate(mcols):
            lc = LineChart()
            lc.title = f"Average {m} per {unit}"
            lc.height, lc.width, lc.legend = 7, 26, None
            lc.add_data(Reference(wt, min_col=5 + k, min_row=1, max_row=last), titles_from_data=True)
            lc.set_categories(Reference(wt, min_col=1, min_row=2, max_row=last))
            lc.series[0].smooth = False
            lc.series[0].graphicalProperties.line.solidFill = BLUE
            lc.x_axis.number_format = "yyyy-mm-dd"
            _axes_on(lc)
            wt.add_chart(lc, f"{anchor_col}{19 + 15 * k}")
        wt.freeze_panes = "B2"
        widths(wt, {"A": 14, "B": 9, "C": 14, "D": 12})

    try:
        wb.save(out_path)
    except PermissionError:
        out_path = out_path.with_name(f"{out_path.stem} {dt.datetime.now():%Y%m%d_%H%M%S}{out_path.suffix}")
        wb.save(out_path)
        print(f"Report was open in Excel - saved a copy instead.")
    return out_path


# --------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Find and chart irregular data in the zone-detail table.")
    ap.add_argument("--table", default=DEFAULT_TABLE,
                    help="table to pull, optionally DB.SCHEMA.TABLE (default: %(default)s)")
    ap.add_argument("--limit", type=int, help="pull only the first N rows")
    ap.add_argument("--max-rows", type=int, default=1_000_000,
                    help="tables bigger than this are randomly sampled down to it (default: %(default)s)")
    ap.add_argument("--from-file", help="analyse a saved .parquet/.csv/.xlsx instead of querying Snowflake")
    ap.add_argument("--out", default=str(REPORT_PATH), help="output workbook")
    a = ap.parse_args()

    if a.from_file:
        df = load_from_file(a.from_file)
        source = f"file: {Path(a.from_file).name}"
    else:
        df, table, total = load_from_snowflake(a.table, a.limit, a.max_rows)
        source = f"Snowflake: {table.replace(chr(34), '')} - {len(df):,} of {total:,} rows"
        if a.limit:
            source += f" (first {a.limit:,})"
        elif len(df) < total:
            source += " (random sample)"
        try:
            archive_pull(df)
        except Exception as e:  # snapshot is a convenience only
            print(f"Could not save raw snapshot: {e}")

    if df.empty:
        sys.exit("No rows returned - nothing to analyse.")
    df = df.reset_index(drop=True)
    print(f"Loaded {len(df):,} rows x {df.shape[1]} columns")

    work, kinds, notes = classify(df)
    flags, fences = find_issues(df, work, kinds, notes)
    timeline = build_timeline(df, work, kinds)
    out = write_report(df, work, kinds, notes, flags, fences, timeline, source, Path(a.out))

    print(f"\nFlagged {len(flags):,} cells across {flags['row'].nunique() if len(flags) else 0:,} rows")
    if len(flags):
        print(flags.groupby("issue").size().sort_values(ascending=False).to_string())
    print(f"\nReport: {out}")


if __name__ == "__main__":
    main()
