"""
District - actual cost above basis vs Finance landed cost (LIC)
===============================================================
LIC = planned landed cost (Finance "Total Cost Above Basis") = differential above index + freight
+ PERC + supply fee.
So each load is compared on a pre-basis (differential) basis:

    Index               MTB/CON index for the order date (Snowflake "Price"; filled when blank, see below)
    Differential        Product - Index
    Freight             Transportation
    Fees (PERC+supply)  Overhead
    Actual above basis  Differential + Freight + Fees            <- compared with Finance LIC
    Storage             shown separately (not in Finance's definition); "incl. storage" column too

Index fill when Snowflake Price is blank for a date: Argus Mont Belvieu LST prompt-month price
(parquet given by --price-index / PRICE_INDEX_FILE, default price_index.parquet in this folder;
columns priceindex, pricedate, delivdate, price) + its average offset to Snowflake Price on days
where both exist; otherwise straight-line between the nearest Snowflake prices.

Finance workbook: --finance / FINANCE_WORKBOOK (default finance_costs.xlsx in this folder). First
sheet: a header row starting "Dist..." (DeptID, district name, basis, then month columns Apr..Mar).
The fiscal year (Apr-Mar) is read from a "YYYY-YYYY" in the file name, or from --fiscal-year.

Usage (from this folder):
    py district_cost_check.py --match "SPRINGFIELD"
    py district_cost_check.py --match "SPRINGFIELD" --finance finance_costs.xlsx --fiscal-year 2026
    py district_cost_check.py --match "SPRINGFIELD" --from-file zone_detail_raw.parquet
"""

import argparse
import datetime as dt
import os
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font

import district_report as dr
import zone_detail_report as base

HERE = Path(__file__).resolve().parent
DEFAULT_FINANCE = Path(os.environ.get("FINANCE_WORKBOOK", HERE / "finance_costs.xlsx"))
ARGUS = Path(os.environ.get("PRICE_INDEX_FILE", HERE / "price_index.parquet"))
ARGUS_INDEX = {"MTB": "NGL fixed price - Financial Mont Belvieu Propane LST",
               "CON": "NGL fixed price - Financial Conway Propane"}
MONTHS = ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]
RED, GREEN, LIGHT_RED, LIGHT_GREEN, GREY = "C00000", "548235", "FFC7CE", "C6EFCE", "D9D9D9"
P5 = "0.00000"
CENTS = '+0.0" ¢";-0.0" ¢";0.0" ¢"'


# --------------------------------------------------------------------------- inputs
def load_finance(path):
    """Finance LIC sheet -> long table: DeptID, District, Basis, Month (Apr..Mar), LIC."""
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and r[0] and str(r[0]).strip().lower().startswith("dist"))
    hdr = [str(h).strip() if h is not None else "" for h in rows[hdr_i]]
    fin = pd.DataFrame(rows[hdr_i + 1:], columns=hdr).dropna(subset=[hdr[0]])
    fin = fin.rename(columns={hdr[0]: "DeptID", hdr[1]: "FinDistrict", hdr[2]: "FinBasis"})
    fin["DeptID"] = fin["DeptID"].astype(str).str.strip().str.zfill(4)
    keep = ["DeptID", "FinDistrict", "FinBasis", *[m for m in MONTHS if m in fin.columns]]
    return fin[keep], ws.title


def index_series(df, basis, dates):
    """Index price per date for one basis, with its source label."""
    sf = (df[df["Basis"] == basis].dropna(subset=["Price"]).groupby("OrderDate")["Price"].first())
    lo, hi = min(dates.min(), sf.index.min()), max(dates.max(), sf.index.max())
    days = pd.date_range(lo, hi, freq="D")
    price = sf.reindex(days)
    source = pd.Series(np.where(price.notna(), "Snowflake Price", ""), index=days, dtype=object)

    offset = None
    if ARGUS.exists() and basis in ARGUS_INDEX:
        a = pd.read_parquet(ARGUS, columns=["priceindex", "pricedate", "delivdate", "price"])
        a = a[(a["priceindex"] == ARGUS_INDEX[basis]) &
              (a["delivdate"] == a["pricedate"].dt.to_period("M").dt.start_time)]
        arg = a.groupby("pricedate")["price"].mean().reindex(days)
        both = price.notna() & arg.notna()
        if both.sum() >= 2:
            offset = float((price[both] - arg[both]).mean())
            fill = price.isna() & arg.notna()
            price[fill] = arg[fill] + offset
            source[fill] = f"Argus LST {offset:+.4f} (est.)" if basis == "MTB" else f"Argus Conway {offset:+.4f} (est.)"
    gap = price.isna()
    price = price.interpolate(method="time", limit_area="inside")
    source[gap & price.notna()] = "Interpolated (est.)"
    return price, source, offset


# --------------------------------------------------------------------------- calc
def fiscal_start(finance_path, year=None):
    """--fiscal-year 2026, or a file named '2026-2027 ...xlsx' -> 2026-04-01 (months run Apr..Mar).

    None (no year given or in the name) means every load is compared."""
    if year:
        return pd.Timestamp(int(year), 4, 1)
    m = re.search(r"(20\d\d)\s*-\s*(20\d\d)", Path(finance_path).name)
    return pd.Timestamp(int(m.group(1)), 4, 1) if m else None


def compute(df, match, fin, fy_start=None):
    sub, cell_out, _, _, _ = dr.find_outliers(df, match)
    sub = sub.copy()
    sub["DeptID"] = sub["DeptID"].astype(str).str.strip().str.zfill(4)
    sub = sub.merge(fin, on="DeptID", how="left", validate="many_to_one").set_index(sub.index)
    if sub["FinDistrict"].isna().any():
        missing = sub.loc[sub["FinDistrict"].isna(), "District"].unique()
        print(f"WARNING: no Finance LIC row for {', '.join(missing)}")

    sub["Index"], sub["Index source"] = np.nan, ""
    offsets = {}
    for b in sub["Basis"].dropna().unique():
        m = sub["Basis"] == b
        px, src, off = index_series(df, b, sub.loc[m, "OrderDate"])
        sub.loc[m, "Index"] = sub.loc[m, "OrderDate"].map(px).to_numpy()
        sub.loc[m, "Index source"] = sub.loc[m, "OrderDate"].map(src).to_numpy()
        offsets[b] = off

    # Each load gets the Finance LIC for its own month (YearMonth), e.g. a Sep load -> Sep column,
    # an Oct load -> Oct column. Loads outside the Finance file's fiscal year (Apr-Mar) get no LIC.
    ym = pd.to_datetime(sub["YearMonth"])
    month = ym.dt.strftime("%b")
    sub["LIC month"] = month
    in_fy = ym.between(fy_start, fy_start + pd.DateOffset(months=11)) if fy_start is not None \
        else pd.Series(True, index=sub.index)
    sub["Finance LIC"] = [sub.at[i, m] if ok and m in sub.columns else np.nan
                          for i, m, ok in zip(sub.index, month, in_fy)]
    if (~in_fy).any():
        print(f"WARNING: {int((~in_fy).sum())} load(s) fall outside the Finance file's fiscal year "
              f"({fy_start:%b %Y} - {fy_start + pd.DateOffset(months=11):%b %Y}); no LIC compared for them")
    sub["Differential"] = sub["Product"] - sub["Index"]
    sub["Freight"] = sub["Transportation"]
    sub["Fees"] = sub["Overhead"]
    sub["Actual above basis"] = sub["Differential"] + sub["Freight"] + sub["Fees"]
    sub["Actual incl. storage"] = sub["Actual above basis"] + sub["Storage"].fillna(0)
    sub["Variance vs LIC"] = sub["Actual above basis"] - sub["Finance LIC"]
    sub["Variance $"] = sub["Variance vs LIC"] * sub["Qty"]
    sub["Variance incl. storage"] = sub["Actual incl. storage"] - sub["Finance LIC"]
    sub["Snowflake LIC"] = sub["LIC"]
    return sub, cell_out, offsets


def wavg(g, col):
    w = g["Qty"].where(g[col].notna())
    return float((g[col] * g["Qty"]).sum() / w.sum()) if w.sum() else np.nan


def rollup(g):
    return {
        "Loads": len(g), "Gallons": g["Qty"].sum(),
        "Index": wavg(g, "Index"), "Product": wavg(g, "Product"),
        "Differential": wavg(g, "Differential"), "Freight": wavg(g, "Freight"), "Fees": wavg(g, "Fees"),
        "Actual above basis": wavg(g, "Actual above basis"),
        "Finance LIC": wavg(g, "Finance LIC"),
        "Variance vs LIC": wavg(g, "Variance vs LIC"),
        "Variance $": g["Variance $"].sum(),
        "Storage": wavg(g, "Storage"), "Actual incl. storage": wavg(g, "Actual incl. storage"),
        "Variance incl. storage": wavg(g, "Variance incl. storage"),
    }


def snowflake_vs_finance_months(df, fin):
    """How the Snowflake LIC column lines up with Finance months, across all districts."""
    d = df.dropna(subset=["LIC"]).copy()
    d["DeptID"] = d["DeptID"].astype(str).str.strip().str.zfill(4)
    d["mon"] = pd.to_datetime(d["YearMonth"]).dt.strftime("%b")
    d = d.drop_duplicates(["DeptID", "mon"]).merge(fin, on="DeptID", how="inner")
    out = {}
    for m in ("same month", *MONTHS):
        if m == "same month":
            ref = pd.Series([r[r["mon"]] if r["mon"] in r else np.nan for _, r in d.iterrows()], index=d.index)
        else:
            ref = d[m]
        out[m] = int(((d["LIC"] - pd.to_numeric(ref, errors="coerce")).abs() < 0.0006).sum())
    return len(d), out


# --------------------------------------------------------------------------- excel
def _bold(ws, cell, size=None):
    ws[cell].font = Font(bold=True, size=size) if size else Font(bold=True)


def write(sub, cell_out, offsets, fin_name, sf_check, source, out_path, finance_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    dist = sub["District"].mode()[0]
    fin_row = sub.iloc[0]
    order = {m: i for i, m in enumerate(MONTHS)}
    months = sorted(sub["LIC month"].dropna().unique(), key=lambda m: order.get(m, 99))
    month = months[0] if len(months) == 1 else ", ".join(months)
    lic_label = f"Finance LIC ({month})" if len(months) == 1 else "Finance LIC (each load's month)"
    total = rollup(sub)
    clean = rollup(sub[sub["Status"] == "Non-Outlier"])

    ws["A1"] = f"{dist} - actual cost above basis vs Finance LIC ({month})"
    _bold(ws, "A1", 14)
    facts = [
        ("Snowflake data", source),
        ("Finance file", f"{Path(finance_path).name}  ({fin_name}: {fin_row['FinDistrict']}, basis {fin_row['FinBasis']})"),
        ("Built", dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Loads / gallons", f"{total['Loads']} loads, {total['Gallons']:,.0f} gal  "
                            f"(order dates {sub['OrderDate'].min():%m/%d} to {sub['OrderDate'].max():%m/%d})"),
    ]
    for i, (k, v) in enumerate(facts, start=3):
        base.put(ws, i, 1, k, bold=True)
        base.put(ws, i, 2, v)

    # headline
    r = 8
    base.put(ws, r, 1, "Headline ($/gal, volume-weighted)", bold=True)
    r += 1
    base.header(ws, r, 1, ["", "All loads", "Excl. outliers"])
    lines = [(lic_label, "Finance LIC"), ("Actual above basis", "Actual above basis"),
             ("Variance (actual - LIC)", "Variance vs LIC"), ("Variance $ (x gallons)", "Variance $"),
             ("   Differential (Product - Index)", "Differential"), ("   Freight", "Freight"),
             ("   Fees (PERC + supply)", "Fees"),
             ("Storage (not in LIC definition)", "Storage"), ("Variance incl. storage", "Variance incl. storage")]
    for label, k in lines:
        r += 1
        base.put(ws, r, 1, label, bold=not label.startswith("   "))
        for j, src in ((2, total), (3, clean)):
            v = src[k]
            fmt = "$#,##0;[Red]-$#,##0" if k == "Variance $" else P5
            color = None
            if k.startswith("Variance") and pd.notna(v):
                color = LIGHT_RED if v > 0 else LIGHT_GREEN
            base.put(ws, r, j, v, fmt=fmt, color=color)
    r += 1
    base.put(ws, r, 1, "Positive variance = actual cost above basis was higher than LIC (we paid more than planned).")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)

    # by month - each month is compared with that month's Finance LIC
    r += 2
    base.put(ws, r, 1, "By month (each load vs the Finance LIC for its own month)", bold=True)
    r += 1
    base.header(ws, r, 1, ["Month", "Loads", "Gallons", "Finance LIC", "Actual above basis",
                           "Variance vs LIC", "Variance $"])
    for m in months:
        g = sub[sub["LIC month"] == m]
        ru = rollup(g)
        ym = pd.to_datetime(g["YearMonth"]).iloc[0]
        r += 1
        vals = [f"{ym:%b %Y}", ru["Loads"], ru["Gallons"], ru["Finance LIC"], ru["Actual above basis"],
                ru["Variance vs LIC"], ru["Variance $"]]
        for j, (v, f) in enumerate(zip(vals, [None, "0", "#,##0", P5, P5, P5, "$#,##0;[Red]-$#,##0"]), start=1):
            color = (LIGHT_RED if v > 0 else LIGHT_GREEN) if j in (6, 7) and pd.notna(v) else None
            base.put(ws, r, j, v, fmt=f, color=color)

    # by lane
    r += 2
    base.put(ws, r, 1, "By lane ($/gal, volume-weighted)", bold=True)
    r += 1
    lane_hdr = r
    cols = ["Lane", "Loads", "Gallons", "Index", "Product", "Differential", "Freight", "Fees",
            "Actual above basis", "Finance LIC", "Variance vs LIC", "Variance $", "Storage", "Outliers"]
    base.header(ws, r, 1, cols)
    for lane, g in sub.groupby("Lane", sort=False):
        r += 1
        ru = rollup(g)
        vals = [lane, ru["Loads"], ru["Gallons"], ru["Index"], ru["Product"], ru["Differential"], ru["Freight"],
                ru["Fees"], ru["Actual above basis"], ru["Finance LIC"], ru["Variance vs LIC"], ru["Variance $"],
                ru["Storage"], int((g["Status"] == "Outlier").sum())]
        fmts = [None, "0", "#,##0", P5, P5, P5, P5, P5, P5, P5, P5, "$#,##0;[Red]-$#,##0", P5, "0"]
        for j, (v, f) in enumerate(zip(vals, fmts), start=1):
            color = (LIGHT_RED if v > 0 else LIGHT_GREEN) if j in (11, 12) and pd.notna(v) else None
            base.put(ws, r, j, v, fmt=f, color=color)
    lane_last = r
    for mt, g in sub.groupby("MoveType"):
        r += 1
        ru = rollup(g)
        vals = [f"All {mt}", ru["Loads"], ru["Gallons"], ru["Index"], ru["Product"], ru["Differential"],
                ru["Freight"], ru["Fees"], ru["Actual above basis"], ru["Finance LIC"], ru["Variance vs LIC"],
                ru["Variance $"], ru["Storage"], int((g["Status"] == "Outlier").sum())]
        fmts = [None, "0", "#,##0", P5, P5, P5, P5, P5, P5, P5, P5, "$#,##0;[Red]-$#,##0", P5, "0"]
        for j, (v, f) in enumerate(zip(vals, fmts), start=1):
            base.put(ws, r, j, v, fmt=f, bold=True)

    ch = BarChart()
    ch.type, ch.title = "bar", "Actual above basis vs Finance LIC by lane ($/gal)"
    ch.height, ch.width = 8, 22
    ch.add_data(Reference(ws, min_col=9, max_col=10, min_row=lane_hdr, max_row=lane_last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=lane_hdr + 1, max_row=lane_last))
    ch.series[0].graphicalProperties.solidFill = base.BLUE
    ch.series[1].graphicalProperties.solidFill = "A5A5A5"
    base._axes_on(ch)
    ch.x_axis.scaling.orientation = "maxMin"
    ch.legend.position = "b"
    ws.add_chart(ch, "P3")

    # notes
    r += 2
    base.put(ws, r, 1, "Method and notes", bold=True)
    n_rows, hits = sf_check
    same, best = hits["same month"], max(MONTHS, key=lambda m: hits[m])
    notes = [
        "LIC = differential above index + freight + PERC + supply fee (Finance definition). Actual above basis = "
        "(Product - Index) + Transportation + Overhead. Storage is shown but not counted.",
        "Index = Snowflake 'Price' for the load's basis and order date. Blank dates were estimated: "
        + "; ".join(f"{b}: Argus prompt-month + {o:+.4f} offset" if o is not None else f"{b}: straight-line"
                    for b, o in offsets.items())
        + ", then straight-line where Argus has no price. See 'Index source' on the Loads sheet.",
        (f"Snowflake's own LIC column ({fin_row['LIC']:.4f} here) matches Finance {month} "
         f"({fin_row['Finance LIC']:.4f}); across all districts it matches the load's month on {same} of "
         f"{n_rows} district-months."
         if pd.notna(fin_row["Finance LIC"]) and abs(fin_row["LIC"] - fin_row["Finance LIC"]) < 0.0006 else
         f"Snowflake's own LIC column ({fin_row['LIC']:.4f} here) is NOT the Finance {month} value "
         f"({fin_row['Finance LIC']:.4f}). Across all districts it matches Finance {best} on {hits[best]} of "
         f"{n_rows} district-months and the load's month on only {same}: Snowflake stamps every row with "
         f"the LIC for the snapshot's current month ({best}), so leftover {month} loads carry {best} LIC. "
         f"This report compares against Finance {month}."),
        "Loads with a flat (fixed) product cost have a differential that moves opposite the index: "
        "high when the index is low and low when it is high.",
        "Each load is priced against the index for its own OrderDate.",
        "Outlier status comes from the same tests as the Outliers vs Non-Outliers report (district_report.py).",
    ]
    for t in notes:
        r += 1
        c = base.put(ws, r, 1, t)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=13)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 32
    base.widths(ws, {"A": 64, "B": 15, "C": 15, "D": 10, "E": 10, "F": 12, "G": 10, "H": 10, "I": 12,
                     "J": 11, "K": 12, "L": 12, "M": 10, "N": 9})

    # loads
    wl = wb.create_sheet("Loads")
    cols = ["Status", "OrderDate", "OrderID", "MoveType", "Lane", "Qty", "Basis", "Index", "Index source",
            "Product", "Differential", "Freight", "Fees", "Actual above basis", "LIC month", "Finance LIC",
            "Variance vs LIC", "Variance $", "Storage", "Actual incl. storage", "Variance incl. storage",
            "Snowflake LIC", "Outlier reasons"]
    base.header(wl, 1, 1, cols)
    fm = {"Qty": "#,##0", "Variance $": "$#,##0.00;[Red]-$#,##0.00"}
    for r, (i, rec) in enumerate(sub.iterrows(), start=2):
        for j, c in enumerate(cols, start=1):
            v = rec[c]
            color = None
            if c == "Status":
                color = LIGHT_RED if v == "Outlier" else LIGHT_GREEN
            elif c in ("Variance vs LIC", "Variance $", "Variance incl. storage") and pd.notna(v):
                color = LIGHT_RED if v > 0 else LIGHT_GREEN
            elif c == "Index source" and "est." in str(v):
                color = "FFEB9C"
            elif c == "Product" and (i, "Product") in cell_out:
                color = LIGHT_RED
            f = fm.get(c, P5 if isinstance(v, (float, np.floating)) else None)
            cell = base.put(wl, r, j, v if not (isinstance(v, str) and v == "") else None, color=color, fmt=f)
            if c == "Outlier reasons":
                cell.alignment = Alignment(wrap_text=False)
    last = len(sub) + 1
    tr = last + 1
    base.put(wl, tr, 1, "Total / wtd avg", bold=True)
    base.put(wl, tr, 6, total["Gallons"], fmt="#,##0", bold=True)
    for c in ("Index", "Product", "Differential", "Freight", "Fees", "Actual above basis", "Finance LIC",
              "Variance vs LIC", "Storage", "Actual incl. storage", "Variance incl. storage"):
        base.put(wl, tr, cols.index(c) + 1, total[c], fmt=P5, bold=True)
    base.put(wl, tr, cols.index("Variance $") + 1, total["Variance $"], fmt=fm["Variance $"], bold=True)
    wl.freeze_panes = "F2"
    wl.auto_filter.ref = f"A1:{base.get_column_letter(len(cols))}{last}"
    widths = {"Status": 12, "OrderDate": 11, "OrderID": 10, "MoveType": 9, "Lane": 58, "Index source": 22,
              "Outlier reasons": 80}
    for j, c in enumerate(cols, start=1):
        wl.column_dimensions[base.get_column_letter(j)].width = widths.get(c, 12)

    # chart: per-load actual vs LIC
    wc = wb.create_sheet("Chart")
    base.put(wc, 1, 1, "Each load's actual cost above basis vs Finance LIC ($/gal). Bars above the grey line = "
                       "paid more than LIC.", bold=True)
    base.header(wc, 3, 1, ["Load", "Actual above basis", "Finance LIC", "Index"])
    for r, (i, rec) in enumerate(sub.iterrows(), start=4):
        base.put(wc, r, 1, f"{str(rec['MoveType'])[:4]} {rec['OrderDate']:%m/%d} #{rec['OrderID']}")
        base.put(wc, r, 2, rec["Actual above basis"], fmt=P5)
        base.put(wc, r, 3, rec["Finance LIC"], fmt=P5)
        base.put(wc, r, 4, rec["Index"], fmt=P5)
    lastc = 3 + len(sub)
    bar = BarChart()
    bar.title, bar.height, bar.width = "Actual cost above basis by load vs Finance LIC", 10, 32
    bar.add_data(Reference(wc, min_col=2, min_row=3, max_row=lastc), titles_from_data=True)
    bar.set_categories(Reference(wc, min_col=1, min_row=4, max_row=lastc))
    bar.series[0].graphicalProperties.solidFill = base.BLUE
    from openpyxl.chart.marker import DataPoint
    for k, (_, rec) in enumerate(sub.iterrows()):
        if rec["Actual above basis"] > rec["Finance LIC"]:
            pt = DataPoint(idx=k)
            pt.graphicalProperties.solidFill = RED
            bar.series[0].dPt.append(pt)
    line = LineChart()
    line.add_data(Reference(wc, min_col=3, min_row=3, max_row=lastc), titles_from_data=True)
    line.series[0].graphicalProperties.line.solidFill = "7F7F7F"
    line.series[0].graphicalProperties.line.width = 28000
    line.series[0].smooth = False
    bar += line
    base._axes_on(bar)
    bar.x_axis.tickLblSkip = 1
    bar.legend.position = "b"
    wc.add_chart(bar, "F3")
    base.widths(wc, {"A": 24, "B": 16, "C": 12, "D": 10})

    # finance reference
    wf = wb.create_sheet("Finance LIC")
    base.put(wf, 1, 1, f"Finance Total Cost Above Basis ($/gal) for {fin_row['FinDistrict']} "
                       f"(from {Path(finance_path).name})", bold=True)
    base.header(wf, 3, 1, ["Month", "Finance LIC", "Loads compared"])
    for k, m in enumerate(MONTHS, start=4):
        used = m in months
        base.put(wf, k, 1, m, color=LIGHT_RED if used else None)
        base.put(wf, k, 2, fin_row.get(m), fmt=P5, color=LIGHT_RED if used else None)
        base.put(wf, k, 3, int((sub["LIC month"] == m).sum()) or None, fmt="0")
    base.widths(wf, {"A": 12, "B": 14, "C": 16})

    try:
        wb.save(out_path)
    except PermissionError:
        out_path = out_path.with_name(f"{out_path.stem} {dt.datetime.now():%Y%m%d_%H%M%S}{out_path.suffix}")
        wb.save(out_path)
        print("Report was open in Excel - saved a copy instead.")
    return out_path, total, clean


def main():
    global ARGUS
    ap = argparse.ArgumentParser(description="Actual cost above basis vs Finance landed cost (LIC) "
                                             "for one district.")
    ap.add_argument("--match", required=True,
                    help='text to find in District/Origin/Destination, e.g. "SPRINGFIELD"')
    ap.add_argument("--finance", default=str(DEFAULT_FINANCE),
                    help="Finance landed-cost workbook (default: FINANCE_WORKBOOK or %(default)s)")
    ap.add_argument("--fiscal-year", type=int,
                    help="first calendar year of the Apr-Mar fiscal year (default: read from file name)")
    ap.add_argument("--price-index", default=str(ARGUS),
                    help="index price parquet used to fill blank prices (default: PRICE_INDEX_FILE or "
                         "%(default)s; skipped if missing)")
    ap.add_argument("--from-file", default=str(base.RAW_SNAPSHOT),
                    help="saved Snowflake pull (default: zone_detail_raw.parquet); use --pull for a fresh one")
    ap.add_argument("--pull", action="store_true", help="pull the zone-detail table fresh from Snowflake first")
    ap.add_argument("--include-forecast", action="store_true",
                    help="also compare Type='Forcast' (planned) loads, not just actuals")
    ap.add_argument("--out")
    a = ap.parse_args()
    ARGUS = Path(a.price_index)

    if a.pull:
        df, table, total = base.load_from_snowflake(base.DEFAULT_TABLE, None, 10_000_000)
        base.archive_pull(df)
        source = f"Snowflake {table.replace(chr(34), '')}, pulled {dt.datetime.now():%Y-%m-%d %H:%M}"
    else:
        df = base.load_from_file(a.from_file)
        pulled = dt.datetime.fromtimestamp(Path(a.from_file).stat().st_mtime)
        source = f"{Path(a.from_file).name} (Snowflake pull {pulled:%Y-%m-%d %H:%M})"
    if "SnapshotInstanceDateTime" in df.columns:
        source += f"; snapshot {pd.to_datetime(df['SnapshotInstanceDateTime']).max():%Y-%m-%d %H:%M}"
    # Snowflake mixes Actual loads with Forcast (planned) loads for the rest of the month.
    # LIC variance is about what we actually paid, so forecasts are left out unless asked for.
    if "Type" in df.columns and not a.include_forecast:
        n_fc = int((df["Type"] != "Actual").sum())
        df = df[df["Type"] == "Actual"]
        if n_fc:
            print(f"Left out {n_fc:,} forecast rows (use --include-forecast to keep them)")
            source += f"; actuals only ({n_fc:,} forecast rows excluded)"
    df = df.reset_index(drop=True)

    fin, fin_name = load_finance(a.finance)
    sub, cell_out, offsets = compute(df, a.match, fin, fiscal_start(a.finance, a.fiscal_year))
    sf_check = snowflake_vs_finance_months(df, fin)
    name = a.match.title()
    out = Path(a.out) if a.out else HERE / f"{name} LIC Check vs Finance.xlsx"
    out, total, clean = write(sub, cell_out, offsets, fin_name, sf_check, source, out, a.finance)

    print(f"Index offsets: {offsets}")
    print(f"Snowflake LIC vs Finance months: {sf_check}")
    for label, t in (("ALL", total), ("EXCL OUTLIERS", clean)):
        print(f"{label}: loads={t['Loads']} gal={t['Gallons']:,.0f} actual={t['Actual above basis']:.5f} "
              f"LIC={t['Finance LIC']:.5f} var={t['Variance vs LIC']:+.5f} (${t['Variance $']:,.0f}) "
              f"diff={t['Differential']:.5f} frt={t['Freight']:.5f} fees={t['Fees']:.5f} "
              f"stor={t['Storage']:.5f}")
    for lane, g in sub.groupby("Lane", sort=False):
        ru = rollup(g)
        print(f"  {lane}: n={ru['Loads']} actual={ru['Actual above basis']:.5f} var={ru['Variance vs LIC']:+.5f} "
              f"(${ru['Variance $']:,.0f}) diff={ru['Differential']:.5f} idx={ru['Index']:.5f}")
    for (m, lic), g in sub.groupby(["LIC month", "Finance LIC"], dropna=False, sort=False):
        ru = rollup(g)
        print(f"  month {m}: {len(g)} loads use Finance LIC {lic:.5f}; actual={ru['Actual above basis']:.5f} "
              f"var={ru['Variance vs LIC']:+.5f}")
    print(sub["Index source"].value_counts().to_dict())
    print(f"Report: {out}")


if __name__ == "__main__":
    main()
