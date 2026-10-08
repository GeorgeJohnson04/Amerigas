"""
Hedge product tracker: daily blended forward-strip prices for each hedge
program configured in hedge_products.json, with a multi-year look-back that
grows as new price days arrive.

For every price date the tracker:
  1. Builds the origin blend for each delivery month, e.g.
         35% Mt B TET + 35% Mt B Non-TET + 30% Conway
     (origins and weights come from "origin_blend" in the config)
  2. Weights those monthly blends by the product's demand shape
     (the "shape_workbook" in the config) to get one volume-weighted
     strip price per product year.

Outputs (next to this script):
    Hedge Product Tracker.xlsx   Summary, charts, daily history, latest curve,
                                 methodology
    hedge_tracker_history.csv    long-format daily history for BI / Power Query

Daily update:
    1. Drop the new ArgusFC_YYYYMMDD.csv file(s) into the LPG folder
       (configured in hedge_products.json), or refresh the price parquet
       with build_price_index.py.
    2. Run update_tracker.bat (or: py hedge_tracker.py).

Try it without real data:
    py make_sample_data.py      # writes a synthetic price parquet + shape workbook
    py hedge_tracker.py
    Any Argus file newer than the parquet is parsed and appended to the
    parquet (a .bak copy is kept), so the price store stays current.

Usage:
    py hedge_tracker.py
    py hedge_tracker.py --config other.json
    py hedge_tracker.py --no-save-prices     # read new Argus files, don't touch the parquet
"""
from __future__ import annotations

import argparse
import json
import os
import shutil
import sys
from dataclasses import dataclass, field
from datetime import datetime

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_CONFIG = os.path.join(HERE, "hedge_products.json")

MONTH_NAMES = {m.lower(): i for i, m in enumerate(
    ["January", "February", "March", "April", "May", "June", "July",
     "August", "September", "October", "November", "December"], start=1)}
MON3 = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
        "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------
def load_config(path: str) -> dict:
    with open(path, "r", encoding="utf-8") as fh:
        cfg = json.load(fh)
    base = os.path.dirname(os.path.abspath(path))
    for key in ("price_parquet", "lpg_dir", "shape_workbook",
                "output_workbook", "history_csv"):
        if cfg.get(key):
            cfg[key] = os.path.normpath(os.path.join(base, cfg[key]))

    total = sum(o["weight"] for o in cfg["origin_blend"])
    if abs(total - 1.0) > 1e-9:
        raise ValueError(f"origin_blend weights sum to {total}, expected 1.0")
    return cfg


# ---------------------------------------------------------------------------
# Demand shapes
# ---------------------------------------------------------------------------
def load_shapes(xlsx_path: str) -> dict[str, dict[int, float]]:
    """Read every shape block from the workbook.

    Layout per block (any column/row offset):
        <label>
        <month name> <month name> ...
        <weight>     <weight>     ...
    Returns {label: {month_number: weight}}; labels are stripped.
    """
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    def month_of(v):
        return MONTH_NAMES.get(str(v).strip().lower()) if isinstance(v, str) else None

    shapes: dict[str, dict[int, float]] = {}
    for r, row in enumerate(grid):
        for c, v in enumerate(row):
            if not isinstance(v, str) or month_of(v) or not v.strip():
                continue
            if r + 2 >= len(grid):
                continue
            months_row, weights_row = grid[r + 1], grid[r + 2]
            shape = {}
            cc = c
            while cc < len(months_row) and month_of(months_row[cc]):
                w = weights_row[cc] if cc < len(weights_row) else None
                if isinstance(w, (int, float)):
                    shape[month_of(months_row[cc])] = float(w)
                cc += 1
            if shape:
                shapes[v.strip()] = shape
    return shapes


# ---------------------------------------------------------------------------
# Prices
# ---------------------------------------------------------------------------
def _import_argus_parser(cfg: dict):
    """parse_propane()/build_blend() from build_price_index.py.

    Looked for next to the price parquet first, then this script's folder and
    its parent, so the tracker keeps working if the folders are rearranged.
    """
    candidates = [os.path.dirname(cfg["price_parquet"]), HERE, os.path.dirname(HERE)]
    for folder in candidates:
        if os.path.exists(os.path.join(folder, "build_price_index.py")):
            if folder not in sys.path:
                sys.path.insert(0, folder)
            import build_price_index as bup  # noqa: E402
            return bup
    raise FileNotFoundError(
        "build_price_index.py not found in: " + "; ".join(candidates))


def load_prices(cfg: dict, save: bool) -> tuple[pd.DataFrame, list[str]]:
    """Component curves (long format) from the parquet, plus any newer Argus files.

    Returns (df, log_lines). df columns: priceindex, pricedate, delivdate, price.
    """
    log: list[str] = []
    component_names = [o["index"] for o in cfg["origin_blend"]]

    path = cfg["price_parquet"]
    if not os.path.exists(path):
        raise FileNotFoundError(f"price parquet not found: {path}")
    store = pd.read_parquet(path)
    log.append(f"price store: {path}")
    log.append(f"  {len(store):,} rows, price dates "
               f"{store['pricedate'].min().date()} .. {store['pricedate'].max().date()}")

    last_date = store.loc[store["priceindex"].isin(component_names), "pricedate"].max()

    lpg_dir = cfg.get("lpg_dir")
    new_frames = []
    if lpg_dir and os.path.isdir(lpg_dir):
        bup = _import_argus_parser(cfg)
        pending = []
        for name in os.listdir(lpg_dir):
            d = bup.file_date(name)
            if d and pd.Timestamp(d) > last_date:
                pending.append((d, os.path.join(lpg_dir, name)))
        pending.sort()
        for d, fp in pending:
            try:
                part = bup.parse_propane(fp)
            except Exception as exc:  # keep going; report the bad file
                log.append(f"  WARNING: could not parse {os.path.basename(fp)}: {exc}")
                continue
            if part.empty:
                log.append(f"  WARNING: no US propane rows in {os.path.basename(fp)}")
                continue
            new_frames.append(part)
        log.append(f"new Argus files after {last_date.date()}: "
                   f"{len(pending)} found, {len(new_frames)} parsed")
    else:
        log.append(f"LPG folder not found ({lpg_dir}); using the parquet as-is")

    if new_frames:
        new = pd.concat(new_frames, ignore_index=True)
        new = new.drop_duplicates(
            subset=["commodity", "priceindex", "pricedate", "delivdate"], keep="last")
        blend = bup.build_blend(new, verbose=False)
        new_all = pd.concat([new, blend], ignore_index=True)
        store = pd.concat([store, new_all], ignore_index=True)
        store = store.drop_duplicates(
            subset=["commodity", "priceindex", "pricedate", "delivdate"], keep="last")
        store = store.sort_values(["priceindex", "pricedate", "delivdate"]).reset_index(drop=True)
        log.append(f"  added {new['pricedate'].nunique()} price date(s), "
                   f"now through {store['pricedate'].max().date()}")
        if save:
            shutil.copy2(path, path + ".bak")
            store.to_parquet(path, index=False)
            log.append(f"  parquet updated (previous copy: {os.path.basename(path)}.bak)")

    df = store.loc[store["priceindex"].isin(component_names),
                   ["priceindex", "pricedate", "delivdate", "price"]].copy()
    missing = set(component_names) - set(df["priceindex"].unique())
    if missing:
        raise ValueError(f"blend component(s) missing from price data: {sorted(missing)}")
    df["pricedate"] = pd.to_datetime(df["pricedate"]).dt.normalize()
    df["delivdate"] = pd.to_datetime(df["delivdate"]).dt.to_period("M").dt.to_timestamp()
    return df, log


# ---------------------------------------------------------------------------
# Strips
# ---------------------------------------------------------------------------
@dataclass
class Strip:
    code: str
    product: str
    label: str
    months: list[pd.Timestamp]
    weights: np.ndarray            # normalised shape weights, one per month
    deadline: pd.Timestamp
    annual_volume: float | None
    coverage: float
    notes: str = ""
    hedge_gal: np.ndarray | None = field(default=None)

    @property
    def start(self):
        return self.months[0]

    @property
    def end(self):
        return self.months[-1]


def build_strips(cfg: dict, shapes: dict[str, dict[int, float]]) -> list[Strip]:
    strips = []
    for p in cfg["products"]:
        label = p["shape_label"].strip()
        if label not in shapes:
            raise ValueError(f"shape '{label}' not in workbook; found {sorted(shapes)}")
        shape = shapes[label]
        for y in p["strip_start_years"]:
            start = pd.Timestamp(year=y, month=p["start_month"], day=1)
            months = [start + pd.DateOffset(months=i) for i in range(p["months"])]
            raw = []
            for m in months:
                if m.month not in shape:
                    raise ValueError(f"shape '{label}' has no weight for {MON3[m.month - 1]}")
                raw.append(shape[m.month])
            w = np.array(raw, dtype=float)
            w = w / w.sum()

            dm = p["hedge_deadline_month"]
            dy = start.year if dm < start.month else start.year - 1
            deadline = pd.Timestamp(year=dy, month=dm, day=1) + pd.offsets.MonthEnd(0)

            end = months[-1]
            name = (f"{p['code']} {MON3[start.month - 1]}{start.year % 100:02d}-"
                    f"{MON3[end.month - 1]}{end.year % 100:02d}")
            vol = p.get("annual_volume_gal")
            cov = float(p.get("hedge_coverage", 1.0))
            # annual_volume_gal is the total for the strip's delivery period,
            # spread by the normalised shape and scaled by hedge coverage
            hedge = w * vol * cov if vol else None
            strips.append(Strip(p["code"], p["name"], name, months, w, deadline,
                                vol, cov, p.get("notes", ""), hedge))
    return strips


def origin_curves(prices: pd.DataFrame, cfg: dict) -> dict[str, pd.DataFrame]:
    """{origin: pricedate x delivdate matrix} plus 'Blend'.

    A delivery month that has rolled off the curve (delivdate before the price
    date's month) keeps its last quoted forward price, so a strip stays priced
    once it enters delivery. A month not yet listed stays blank and the strip is
    reported as incomplete for that day.
    """
    out = {}
    for o in cfg["origin_blend"]:
        sub = prices[prices["priceindex"] == o["index"]]
        mat = sub.pivot_table(index="pricedate", columns="delivdate",
                              values="price", aggfunc="last").sort_index()
        mat = mat.reindex(columns=sorted(mat.columns))
        pd_month = mat.index.to_period("M").to_timestamp()
        expired = np.array([[dcol < pm for dcol in mat.columns] for pm in pd_month])
        ff = mat.ffill()
        fill = expired & mat.isna().to_numpy() & ff.notna().to_numpy()
        vals = mat.to_numpy(copy=True)
        vals[fill] = ff.to_numpy()[fill]
        out[o["origin"]] = pd.DataFrame(vals, index=mat.index, columns=mat.columns)

    idx = sorted(set().union(*[m.index for m in out.values()]))
    cols = sorted(set().union(*[m.columns for m in out.values()]))
    for k in out:
        out[k] = out[k].reindex(index=idx, columns=cols)
    blend = sum(out[o["origin"]] * o["weight"] for o in cfg["origin_blend"])  # NaN if any leg missing
    out["Blend"] = blend
    return out


def strip_history(strip: Strip, curves: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Daily shaped strip price for each origin and the blend."""
    rows = {}
    months = pd.DatetimeIndex(strip.months)
    for name, mat in curves.items():
        sub = mat.reindex(columns=months)
        vals = sub.to_numpy()
        complete = ~np.isnan(vals).any(axis=1)
        price = np.where(complete, np.nan_to_num(vals) @ strip.weights, np.nan)
        rows[name] = price
        if name == "Blend":
            rows["months_quoted"] = (~np.isnan(vals)).sum(axis=1)
            pd_month = sub.index.to_period("M").to_timestamp()
            rows["months_expired"] = np.array(
                [(months < pm).sum() for pm in pd_month])
    df = pd.DataFrame(rows, index=mat.index)
    df.index.name = "pricedate"
    return df


# ---------------------------------------------------------------------------
# Summary statistics
# ---------------------------------------------------------------------------
def _asof(series: pd.Series, when: pd.Timestamp):
    s = series.dropna()
    s = s[s.index <= when]
    return (s.iloc[-1], s.index[-1]) if len(s) else (np.nan, pd.NaT)


def summarise(strip: Strip, hist: pd.DataFrame, asof: pd.Timestamp,
              latest_blend: np.ndarray) -> dict:
    s = hist["Blend"].dropna()
    row = {
        "Product": strip.code,
        "Strip": strip.label,
        "Delivery": f"{strip.start:%b %Y} - {strip.end:%b %Y}",
        "Hedge deadline": strip.deadline.date(),
    }
    if s.empty:
        row.update({"Price date": None, "Blended $/gal": None,
                    "Status": "No complete strip quote yet"})
        return row

    last_px, last_dt = s.iloc[-1], s.index[-1]
    prev_px, _ = _asof(s, s.index[-2]) if len(s) > 1 else (np.nan, None)
    wk_px, _ = _asof(s, last_dt - pd.Timedelta(days=7))
    mo_px, _ = _asof(s, last_dt - pd.DateOffset(months=1))
    yr_px, _ = _asof(s, last_dt - pd.DateOffset(years=1))

    def chg(x):
        return None if pd.isna(x) else last_px - x

    days_left = (strip.deadline - asof).days
    if asof >= strip.start:
        status = "In delivery"
    elif days_left < 0:
        status = "Past hedge deadline"
    else:
        status = "Hedge window open"

    row.update({
        "Price date": last_dt.date(),
        "Blended $/gal": last_px,
        "Chg 1D": chg(prev_px),
        "Chg 1W": chg(wk_px),
        "Chg 1M": chg(mo_px),
        "Chg 1Y": chg(yr_px),
        "Look-back avg": s.mean(),
        "Look-back high": s.max(),
        "High date": s.idxmax().date(),
        "Look-back low": s.min(),
        "Low date": s.idxmin().date(),
        "Percentile": (s < last_px).mean() + 0.5 * (s == last_px).mean(),
        "First full quote": s.index[0].date(),
        "Days priced": int(s.size),
        "Days to deadline": max(days_left, 0),
        "Status": status,
    })
    for col in hist.columns:
        if col in ("Blend", "months_quoted", "months_expired"):
            continue
        # per-origin strip price on the same day as the blend
        row[f"{col} $/gal"] = hist.at[last_dt, col]
    if strip.hedge_gal is not None:
        row["Hedge gal"] = float(strip.hedge_gal.sum())
        row["Notional $"] = float(np.nansum(strip.hedge_gal * latest_blend))
    else:
        row["Hedge gal"] = None
        row["Notional $"] = None
    return row


def latest_curve(strip: Strip, curves: dict[str, pd.DataFrame], asof: pd.Timestamp) -> pd.DataFrame:
    months = pd.DatetimeIndex(strip.months)
    rows = {"Delivery month": [m.strftime("%b %Y") for m in months],
            "Shape weight": strip.weights}
    for name, mat in curves.items():
        rows[f"{name} $/gal"] = mat.reindex(columns=months).loc[asof].to_numpy()
    df = pd.DataFrame(rows)
    df["Weighted $/gal"] = df["Shape weight"] * df["Blend $/gal"]
    if strip.hedge_gal is not None:
        df["Hedge gal"] = strip.hedge_gal
        df["Notional $"] = df["Hedge gal"] * df["Blend $/gal"]
    return df


# ---------------------------------------------------------------------------
# Excel output
# ---------------------------------------------------------------------------
def write_workbook(path, summary, daily, origin_daily, curves_by_strip, strips,
                   cfg, shapes, asof, log_lines, blend_check):
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.axis import DateAxis
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14)
    thin = Side(style="thin", color="BFBFBF")

    def style_header(ws, row, ncols, col0=1):
        for c in range(col0, col0 + ncols):
            cell = ws.cell(row=row, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)

    def autosize(ws, min_w=9, max_w=38):
        widths = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                n = len(str(cell.value)) if not isinstance(cell.value, float) else 9
                widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), n)
        for col, w in widths.items():
            ws.column_dimensions[col].width = max(min_w, min(max_w, w + 2))

    fmt = {"$/gal": "$0.0000", "Chg": "+0.0000;-0.0000;0.0000", "gal": "#,##0",
           "Notional": "$#,##0", "Percentile": "0%", "weight": "0.00%"}

    def number_format_for(header: str):
        h = str(header)
        if h.startswith("Chg"):
            return fmt["Chg"]
        if "Notional" in h:
            return fmt["Notional"]
        if "gal" in h and "$" not in h:
            return fmt["gal"]
        if "$/gal" in h or h.startswith("Look-back") or h in [s.label for s in strips]:
            return fmt["$/gal"]
        if h == "Percentile":
            return fmt["Percentile"]
        if "weight" in h.lower():
            return fmt["weight"]
        return None

    def apply_formats(ws, header_row, first_data_row, last_row):
        for c in range(1, ws.max_column + 1):
            nf = number_format_for(ws.cell(row=header_row, column=c).value)
            if not nf:
                continue
            for r in range(first_data_row, last_row + 1):
                ws.cell(row=r, column=c).number_format = nf

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        # ---- Summary -----------------------------------------------------
        summary.to_excel(xw, sheet_name="Summary", index=False, startrow=3)
        ws = xw.sheets["Summary"]
        ws["A1"] = "Hedge Product Tracker - blended forward strip prices ($/gal)"
        ws["A1"].font = title_font
        blend_txt = " + ".join(f"{o['weight']:.1%} {o['origin']}" for o in cfg["origin_blend"])
        ws["A2"] = (f"As of {asof:%a %d %b %Y}.  Origin blend: {blend_txt}.  "
                    f"Look-back from {daily.index.min():%d %b %Y}.  "
                    f"Strip price = demand-shape-weighted average of monthly blend prices.")
        ws["A2"].font = Font(italic=True, color="595959")
        style_header(ws, 4, summary.shape[1])
        apply_formats(ws, 4, 5, 4 + len(summary))
        ws.freeze_panes = "C5"
        autosize(ws)
        ws.row_dimensions[4].height = 32

        # ---- Charts ------------------------------------------------------
        daily_out = daily.reset_index()
        daily_out.to_excel(xw, sheet_name="Daily Blend", index=False)
        wsd = xw.sheets["Daily Blend"]
        style_header(wsd, 1, daily_out.shape[1])
        for r in range(2, len(daily_out) + 2):
            wsd.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        apply_formats(wsd, 1, 2, len(daily_out) + 1)
        wsd.freeze_panes = "B2"
        autosize(wsd)
        wsd.column_dimensions["A"].width = 12

        wsc = xw.book.create_sheet("Charts", 1)
        wsc["A1"] = "Blended strip price history ($/gal)"
        wsc["A1"].font = title_font
        n = len(daily_out)
        dates = Reference(wsd, min_col=1, min_row=2, max_row=n + 1)
        codes = list(dict.fromkeys(s.code for s in strips))
        anchor_row = 3
        for code in codes + ["ALL"]:
            chart = LineChart()
            chart.title = ("All strips - blended $/gal" if code == "ALL"
                           else f"{code} - {next(s.product for s in strips if s.code == code)} - blended $/gal")
            chart.y_axis.number_format = "$0.00"
            # DateAxis has axId 500; the value axis must cross it, or Excel
            # rejects the file as corrupt
            chart.y_axis.crossAx = 500
            chart.x_axis = DateAxis(crossAx=100)
            chart.x_axis.number_format = "mmm-yy"
            chart.x_axis.majorTimeUnit = "months"
            chart.x_axis.majorUnit = 3
            # Scale the value axis to the data (Excel defaults to $0, which
            # flattens a $0.50-$0.80 range into near-straight lines)
            cols = [s.label for s in strips if code == "ALL" or s.code == code]
            lo, hi = daily[cols].min().min(), daily[cols].max().max()
            chart.y_axis.scaling.min = float(np.floor(lo * 20) / 20)
            chart.y_axis.scaling.max = float(np.ceil(hi * 20) / 20)
            chart.y_axis.majorUnit = 0.05
            # right-hand side: a bottom legend collides with the date labels
            # and a top legend collides with the chart title
            chart.legend.position = "r"
            chart.legend.overlay = False
            # axis titles collide with tick labels in Excel's auto layout;
            # the chart title and $ number format already say what's shown
            chart.x_axis.title = None
            chart.y_axis.title = None
            chart.height, chart.width = 9, 26
            chart.display_blanks = "gap"
            for ci, s in enumerate(strips, start=2):
                if code != "ALL" and s.code != code:
                    continue
                ref = Reference(wsd, min_col=ci, min_row=1, max_row=n + 1)
                chart.add_data(ref, titles_from_data=True)
            for ser in chart.series:
                ser.smooth = False
                ser.graphicalProperties.line.width = 15000
            chart.set_categories(dates)
            # openpyxl hides axes by default in newer Excel builds
            chart.x_axis.delete = False
            chart.y_axis.delete = False
            wsc.add_chart(chart, f"A{anchor_row}")
            anchor_row += 19

        # ---- Origin detail ----------------------------------------------
        od = origin_daily.reset_index()
        od.to_excel(xw, sheet_name="Origin Detail", index=False)
        wso = xw.sheets["Origin Detail"]
        style_header(wso, 1, od.shape[1])
        for r in range(2, len(od) + 2):
            wso.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        for c in range(2, od.shape[1] + 1):
            h = str(wso.cell(row=1, column=c).value)
            if "months" in h:
                continue
            for r in range(2, len(od) + 2):
                wso.cell(row=r, column=c).number_format = fmt["$/gal"]
        wso.freeze_panes = "B2"
        autosize(wso, max_w=24)
        wso.row_dimensions[1].height = 45

        # ---- Latest curve per strip -------------------------------------
        wsl = xw.book.create_sheet("Latest Curve")
        wsl["A1"] = f"Month-by-month prices on {asof:%d %b %Y} ($/gal)"
        wsl["A1"].font = title_font
        r = 3
        for s in strips:
            df = curves_by_strip[s.label]
            wsl.cell(row=r, column=1, value=f"{s.label}   (hedge by {s.deadline:%d %b %Y}"
                     + (f", {s.coverage:.0%} coverage)" if s.coverage else ")")).font = Font(bold=True, size=12)
            r += 1
            for ci, col in enumerate(df.columns, start=1):
                wsl.cell(row=r, column=ci, value=col)
            style_header(wsl, r, df.shape[1])
            hdr = r
            for rec in df.itertuples(index=False):
                r += 1
                for ci, v in enumerate(rec, start=1):
                    wsl.cell(row=r, column=ci, value=None if (isinstance(v, float) and np.isnan(v)) else v)
            # totals row
            r += 1
            wsl.cell(row=r, column=1, value="Strip").font = Font(bold=True)
            for ci, col in enumerate(df.columns, start=1):
                letter = get_column_letter(ci)
                rng = f"{letter}{hdr + 1}:{letter}{r - 1}"
                if col in ("Shape weight", "Weighted $/gal", "Hedge gal", "Notional $"):
                    wsl.cell(row=r, column=ci, value=f"=SUM({rng})")
                elif col.endswith("$/gal"):
                    wcol = get_column_letter(2)
                    wsl.cell(row=r, column=ci,
                             value=f"=SUMPRODUCT({wcol}{hdr + 1}:{wcol}{r - 1},{rng})")
                wsl.cell(row=r, column=ci).font = Font(bold=True)
                wsl.cell(row=r, column=ci).border = Border(top=thin)
            apply_formats(wsl, hdr, hdr + 1, r)
            r += 3
        autosize(wsl, max_w=20)
        wsl.column_dimensions["A"].width = 16

        # ---- Methodology -------------------------------------------------
        wsm = xw.book.create_sheet("Methodology")
        lines = [
            ("Hedge Product Tracker - methodology", title_font),
            ("", None),
            ("Origin blend (per delivery month)", Font(bold=True)),
        ]
        for o in cfg["origin_blend"]:
            lines.append((f"   {o['weight']:.1%}  {o['origin']}  <-  Argus {o['index']}", None))
        lines += [
            ("", None),
            ("Strip price", Font(bold=True)),
            ("   Blended strip $/gal = SUM(shape weight(month) x origin blend $/gal(month)) over the "
             "strip's delivery months. Shape weights are normalised to sum to 100% within each strip.", None),
            ("   A day is only priced when every month of the strip is quoted on the forward curve. "
             "Far strips therefore start later in the look-back (see 'First full quote' on Summary).", None),
            ("   Once a delivery month rolls off the curve it holds its last quoted forward price, so a "
             "strip stays priced through delivery.", None),
            ("   Hedge coverage scales gallons only; it does not change the $/gal strip price.", None),
            ("", None),
            (f"Demand shapes (from {os.path.basename(cfg['shape_workbook'])})", Font(bold=True)),
        ]
        for p in cfg["products"]:
            sh = shapes[p["shape_label"].strip()]
            order = [((p["start_month"] - 1 + i) % 12) + 1 for i in range(p["months"])]
            lines.append((f"   {p['code']}: " + ", ".join(
                f"{MON3[m - 1]} {sh[m]:.2%}" for m in order), None))
        lines += [("", None), ("Products and assumptions", Font(bold=True))]
        for p in cfg["products"]:
            vol = f"{p['annual_volume_gal']:,.0f} gal" if p.get("annual_volume_gal") else "volume not set"
            lines.append((f"   {p['code']} ({p['name']}): {vol}, coverage "
                          f"{p.get('hedge_coverage', 1):.0%}, hedge deadline month "
                          f"{MON3[p['hedge_deadline_month'] - 1]}. {p.get('notes', '')}", None))
        lines += [("", None), ("Data check", Font(bold=True))]
        lines.append((f"   {blend_check}", None))
        lines += [("", None), ("Last run", Font(bold=True))]
        lines.append((f"   {datetime.now():%Y-%m-%d %H:%M}", None))
        for ln in log_lines:
            lines.append((f"   {ln}", None))
        for i, (text, font) in enumerate(lines, start=1):
            cell = wsm.cell(row=i, column=1, value=text)
            if font:
                cell.font = font
        wsm.column_dimensions["A"].width = 160

        order = ["Summary", "Charts", "Latest Curve", "Daily Blend",
                 "Origin Detail", "Methodology"]
        for target, name in enumerate(order):
            current = xw.book.sheetnames.index(name)
            if current != target:
                xw.book.move_sheet(name, offset=target - current)


# ---------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser(description="Hedge product blended-strip tracker")
    ap.add_argument("--config", default=DEFAULT_CONFIG)
    ap.add_argument("--no-save-prices", action="store_true",
                    help="Parse new Argus files but do not write them back to the parquet")
    args = ap.parse_args()

    cfg = load_config(args.config)
    print("=" * 72)
    print("Hedge Product Tracker")
    print("=" * 72)

    shapes = load_shapes(cfg["shape_workbook"])
    print(f"shapes: {', '.join(f'{k} ({len(v)} months)' for k, v in shapes.items())}")
    strips = build_strips(cfg, shapes)

    prices, log = load_prices(cfg, save=not args.no_save_prices)
    for ln in log:
        print(ln)

    curves = origin_curves(prices, cfg)
    asof = curves["Blend"].index.max()

    if cfg.get("history_start"):
        start = pd.Timestamp(cfg["history_start"])
    else:
        start = asof - pd.DateOffset(years=int(cfg.get("lookback_years", 3)))

    # Cross-check against the pre-built blend index in the parquet (if present)
    blend_check = "blend check skipped"
    ref_name = cfg.get("blend_check_index", "US Propane Blend")
    try:
        ref = pd.read_parquet(cfg["price_parquet"], filters=[("priceindex", "==", ref_name)])
        if ref.empty:
            raise ValueError(f"no '{ref_name}' rows in the parquet")
        ref = ref.pivot_table(index="pricedate", columns="delivdate", values="price", aggfunc="last")
        mine = curves["Blend"].reindex(index=ref.index, columns=ref.columns)
        diff = (mine - ref).abs().stack().dropna()
        blend_check = (f"Recomputed origin blend vs '{ref_name}' in the parquet: "
                       f"{len(diff):,} points, max abs diff {diff.max():.2e} $/gal")
    except Exception as exc:
        blend_check = f"blend check not run: {exc}"
    print(blend_check)

    daily_cols, origin_cols, summary_rows, curves_by_strip = {}, {}, [], {}
    history_long = []
    for s in strips:
        h = strip_history(s, curves)
        h = h[h.index >= start]
        daily_cols[s.label] = h["Blend"]
        for o in [o["origin"] for o in cfg["origin_blend"]] + ["Blend"]:
            origin_cols[f"{s.label} | {o}"] = h[o]
        origin_cols[f"{s.label} | months quoted"] = h["months_quoted"]

        lc = latest_curve(s, curves, asof)
        curves_by_strip[s.label] = lc
        summary_rows.append(summarise(s, h, asof, lc["Blend $/gal"].to_numpy()))

        hl = h.reset_index()
        hl.insert(1, "product", s.code)
        hl.insert(2, "strip", s.label)
        history_long.append(hl)

    daily = pd.DataFrame(daily_cols)
    daily = daily.dropna(how="all")
    origin_daily = pd.DataFrame(origin_cols).loc[daily.index]
    summary = pd.DataFrame(summary_rows)

    # ---- console summary -------------------------------------------------
    print()
    print(f"as of {asof.date()}   look-back from {daily.index.min().date()} "
          f"({len(daily):,} price dates)")
    print(f"{'Strip':<18}{'$/gal':>9}{'1D':>9}{'1W':>9}{'1M':>9}"
          f"{'Avg':>9}{'Low':>9}{'High':>9}  First quote   Status")
    for r in summary_rows:
        if r.get("Blended $/gal") is None:
            print(f"{r['Strip']:<18}{'n/a':>9}  {r['Status']}")
            continue
        f = lambda v: f"{v:>+9.4f}" if v is not None and not pd.isna(v) else f"{'':>9}"
        print(f"{r['Strip']:<18}{r['Blended $/gal']:>9.4f}{f(r['Chg 1D'])}{f(r['Chg 1W'])}"
              f"{f(r['Chg 1M'])}{r['Look-back avg']:>9.4f}{r['Look-back low']:>9.4f}"
              f"{r['Look-back high']:>9.4f}  {r['First full quote']}    {r['Status']}")

    # ---- write outputs ---------------------------------------------------
    hist_df = pd.concat(history_long, ignore_index=True)
    hist_df = hist_df.rename(columns={"Blend": "blend", "Mt B TET": "mtb_tet",
                                      "Mt B Non-TET": "mtb_non_tet", "Conway": "conway"})
    hist_df = hist_df.dropna(subset=["blend"])
    csv_path = cfg["history_csv"]
    out_path = cfg["output_workbook"]
    try:
        hist_df.to_csv(csv_path, index=False, float_format="%.6f")
        write_workbook(out_path, summary, daily, origin_daily, curves_by_strip, strips,
                       cfg, shapes, asof, log, blend_check)
    except PermissionError:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        alt = os.path.splitext(out_path)[0] + f" {stamp}.xlsx"
        print(f"\nWARNING: {os.path.basename(out_path)} is open in Excel; writing {os.path.basename(alt)} instead")
        write_workbook(alt, summary, daily, origin_daily, curves_by_strip, strips,
                       cfg, shapes, asof, log, blend_check)
        out_path = alt

    print()
    print(f"wrote {out_path}")
    print(f"wrote {csv_path} ({len(hist_df):,} rows)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
