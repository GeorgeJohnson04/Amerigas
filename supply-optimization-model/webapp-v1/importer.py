"""
importer.py - bringing data in from Excel (paste or upload) and back out.

The data arrives by hand, in whatever layout the sender used, so the importer
works the layout out instead of asking anyone to reformat it:

  * A header row is found by its words (Plant, Location, From, To, Hub, State,
    Area, Rate, Min, Max, ...) and by month headings (Apr, Apr-26, 4/1/2026,
    April 2026, P1 ...). Extra columns (Total, Summer, Ratio ...) are ignored,
    and where a sheet has two sets of month columns you can pick which.
  * Values can be one per month (months across, or a Month column) or a single
    value used for the months you choose.
  * Freight can also come as a grid: supply points down, plants across.
  * Names are matched loosely - exact name, district number ("1001"), the name
    without its code, supplier + location, a hub's code or name - and every
    match shows how sure it is. Anything can be re-pointed, skipped or added
    as new before saving.
  * Units are worked out (gallons or thousand gallons, $/gal or $/thousand
    gal, 90% or 0.9) by comparing with what's already there, and can be changed.

parse() never changes anything: it returns a preview plus the list of
operations; apply_ops() carries those out on a copy the store then saves.

Every kind of data has a sheet in the Excel workbook that export_workbook()
writes, so data can go out to Excel, be edited offline, and come back in.
"""

from __future__ import annotations

import csv
import datetime as _dt
import difflib
import io
import math
import re
import uuid
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
_FULL = ["january", "february", "march", "april", "may", "june", "july", "august", "september",
         "october", "november", "december"]
SEASON_ORDER = ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]
MAX_ROWS, MAX_CHARS = 30000, 12_000_000
AUTO_OK, AUTO_CHECK = 0.86, 0.70

# kind -> what it is. entity: plant | supply | lane | hub | state | weather
# field: where monthly values go; unit: volume | price | pct | raw | None
KINDS: Dict[str, dict] = {
    "plants":      {"label": "Plants", "group": "demand", "entity": "plant", "field": None, "unit": None,
                    "help": "One row per plant: Area, Plant, State (and opening inventory if any). Adds new plants."},
    "budget":      {"label": "Budget demand", "group": "demand", "entity": "plant", "field": "budget", "unit": "volume",
                    "help": "Demand per plant, months across (gallons). The contract % by state and weather are applied on top."},
    "plant_cap":   {"label": "Plant storage capacity", "group": "demand", "entity": "plant", "field": "cap", "unit": "volume",
                    "help": "Most a plant can hold at month end (gallons)."},
    "plant_min":   {"label": "Plant minimum inventory", "group": "demand", "entity": "plant", "field": "inv_min", "unit": "volume",
                    "help": "Least that must stay in a plant's storage at month end (gallons)."},
    "state_pct":   {"label": "Contract % by state", "group": "demand", "entity": "state", "field": None, "unit": "pct",
                    "help": "% of budget demand to contract for, by state and month (the State Adjustments sheet)."},
    "supply":      {"label": "Supply points", "group": "supply", "entity": "supply", "field": None, "unit": None,
                    "help": "One row per supply point: Area, Location, Supplier, Hub (and opening inventory). Adds new ones."},
    "contract":    {"label": "Contract volumes", "group": "supply", "entity": "supply", "field": "contract", "unit": "volume",
                    "help": "Contracted gallons per supply point, months across."},
    "diff":        {"label": "Differentials", "group": "supply", "entity": "supply", "field": "diff", "unit": "price",
                    "help": "Differential to the hub, $/gal (price = hub price + differential)."},
    "min_pct":     {"label": "Minimum lift %", "group": "supply", "entity": "supply", "field": "min_pct", "unit": "pct",
                    "help": "Least share of the contract that must be taken (90 = 90%)."},
    "max_pct":     {"label": "Maximum lift %", "group": "supply", "entity": "supply", "field": "max_pct", "unit": "pct",
                    "help": "Most that may be taken (150 = half as much again as the contract)."},
    "supply_cap":  {"label": "Supply storage capacity", "group": "supply", "entity": "supply", "field": "cap", "unit": "volume",
                    "help": "Most a supply point can hold at month end (gallons)."},
    "supply_min":  {"label": "Supply minimum inventory", "group": "supply", "entity": "supply", "field": "inv_min", "unit": "volume",
                    "help": "Least that must stay in a supply point's storage (gallons)."},
    "hubs":        {"label": "Hub prices", "group": "supply", "entity": "hub", "field": "price", "unit": "price",
                    "help": "Monthly price per hub (CON, MTB ...), $/gal. Applies to every area with that hub unless an Area is given."},
    "lanes":       {"label": "Lanes & freight", "group": "lanes", "entity": "lane", "field": "rate", "unit": "price",
                    "help": "Routes (OD pairs): From, Supplier, To, freight $/gal (one rate or months across), optional Min / Max gallons. "
                            "Adds new lanes. A blank rate means the lane isn't used that month."},
    "weather":     {"label": "Weather history", "group": "weather", "entity": "weather", "field": None, "unit": "raw",
                    "help": "Place (plant, state, area or All), Season (2006-07), then Apr..Mar: gallons or degree days."},
}
ORDER = ["plants", "budget", "state_pct", "plant_cap", "plant_min", "supply", "contract", "diff", "min_pct", "max_pct",
         "supply_cap", "supply_min", "hubs", "lanes", "weather"]

ROLES = {
    "area": ["area", "model", "areaname", "modelarea", "optimizationarea"],
    "region": ["region", "supplyregion"],
    "plant": ["plant", "plantname", "district", "districtname", "geography", "demandpoint", "lookup", "plantdistrict",
              "location" ],
    "location": ["location", "supplypoint", "terminal", "supply", "supplylocation", "origin", "source", "facility"],
    "supplier": ["supplier", "vendor", "counterparty", "company", "seller"],
    "hub": ["hub", "pricehub", "index", "hubcode", "pricingpoint"],
    "hubname": ["hubname", "name", "description"],
    "state": ["state", "st", "statecode"],
    "from": ["from", "origin", "supplypoint", "terminal", "source", "shipfrom", "fromlocation", "loadpoint", "location"],
    "to": ["to", "destination", "dest", "shipto", "plant", "district", "deliverto", "toplant"],
    "value": ["value", "rate", "freight", "freightrate", "price", "volume", "gallons", "gal", "demand", "contract",
              "diff", "differential", "pct", "percent", "amount", "cost", "ratepergal"],
    "min": ["min", "minimum", "mingal", "mingallons", "minshipment", "minvolume"],
    "max": ["max", "maximum", "maxgal", "maxgallons", "maxshipment", "maxvolume", "capacity"],
    "open": ["open", "opening", "openinginventory", "startinventory", "initialinventory", "beginninginventory", "openinv"],
    "season": ["season", "year", "winter", "contractyear", "heatingseason"],
    "month": ["month", "period", "date", "monthyear"],
    "place": ["place", "key", "location", "area", "state", "plant", "district", "region", "geography", "name"],
}
# which roles each kind looks for, in priority order for the "name" of a row
KIND_ROLES = {
    "plant": ["area", "region", "plant", "state", "open", "value", "month"],
    "supply": ["area", "region", "location", "supplier", "hub", "open", "value", "month"],
    "lane": ["area", "from", "supplier", "to", "value", "min", "max", "month"],
    "hub": ["area", "hub", "hubname", "value", "month"],
    "state": ["state", "value", "month"],
    "weather": ["place", "season", "value", "month"],
}
_BLANK = {"", "-", "--", "\u2014", "n/a", "na", "#n/a", "x", "none", "null", "#value!", "#div/0!", "#ref!"}
_SKIP = re.compile(r"^(grand\s*)?(sub\s*)?totals?\b|^summer$|^winter$|^ratio$|^average$|^avg$|^peak", re.I)
_PLACEHOLDER = re.compile(r"^(supplier|demand|supply|plant)\s*\d+$", re.I)


def fold(s) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(s or "").lower())


def col_name(c: int) -> str:
    s, c = "", c + 1
    while c:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


def new_id(prefix: str) -> str:
    return f"{prefix}{uuid.uuid4().hex[:8]}"


# --------------------------------------------------------------------------- #
# Cells and headers
# --------------------------------------------------------------------------- #

def parse_number(s) -> Tuple[str, Optional[float]]:
    """('num', v) | ('pct', fraction) | ('blank', None) | ('text', None)."""
    t = str(s if s is not None else "").strip()
    if t.lower() in _BLANK:
        return "blank", None
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()").replace(",", "").replace("$", "").replace(" ", "").replace("\u00a0", "")
    pct = t.endswith("%")
    t = t.rstrip("%")
    try:
        v = float(t)
    except ValueError:
        return "text", None
    if not math.isfinite(v):
        return "text", None
    v = -v if neg else v
    return ("pct", v / 100.0) if pct else ("num", v)


def parse_month(s) -> Optional[Tuple[int, Optional[int]]]:
    """A month heading -> (month 1-12, year or None), or None. 'P1' -> (period 1 as month of an Apr year)."""
    t = str(s or "").strip().lower().replace(".", "")
    if not t or len(t) > 32:
        return None
    m = re.fullmatch(r"(?:p|period|month)\s*(\d{1,2})", t)
    if m and 1 <= int(m.group(1)) <= 12:
        return ((int(m.group(1)) + 2) % 12 + 1, None)        # P1 = Apr
    m = re.fullmatch(r"([a-z]{3,9})(?:[\s\-/'_,]*(\d{2}|\d{4}))?", t)
    if m:
        w = m.group(1)
        k = next((i for i in range(12) if w == MONTHS[i].lower() or (len(w) >= 3 and _FULL[i].startswith(w))), None)
        if k is None:
            return None
        y = m.group(2)
        return (k + 1, (2000 + int(y) if y and len(y) == 2 else int(y) if y else None))
    m = re.fullmatch(r"(\d{4})-(\d{1,2})(?:-(\d{1,2}))?(?:[ t].*)?", t)
    if m and 1 <= int(m.group(2)) <= 12:
        return (int(m.group(2)), int(m.group(1)))
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{2}|\d{4})(?:\s.*)?", t)
    if m and 1 <= int(m.group(1)) <= 12:
        y = m.group(3)
        return (int(m.group(1)), 2000 + int(y) if len(y) == 2 else int(y))
    m = re.fullmatch(r"(\d{1,2})/(\d{4})", t)
    if m and 1 <= int(m.group(1)) <= 12:
        return (int(m.group(1)), int(m.group(2)))
    return None


def cell_text(v) -> str:
    """An Excel cell as text, the way it would look if copied."""
    if v is None:
        return ""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime("%b-%y")
    if isinstance(v, float):
        return repr(v) if not v.is_integer() else str(int(v))
    return str(v)


def read_table(text: str) -> List[List[str]]:
    if len(text) > MAX_CHARS:
        raise ValueError("that's a very large paste - paste one sheet at a time")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    rows = [r.split("\t") for r in text.split("\n")]
    if len(rows) > MAX_ROWS:
        raise ValueError(f"more than {MAX_ROWS:,} rows - paste a smaller block")
    if all(len(r) == 1 for r in rows) and any("," in r[0] for r in rows):
        rows = list(csv.reader(io.StringIO(text)))
    rows = [[c.strip().strip('"') for c in r] for r in rows]
    while rows and not any(rows[-1]):
        rows.pop()
    while rows and not any(rows[0]):
        rows.pop(0)
    width = max((len(r) for r in rows), default=0)
    rows = [r + [""] * (width - len(r)) for r in rows]
    keep = [c for c in range(width) if any(r[c] for r in rows)]
    return [[r[c] for c in keep] for r in rows]


def table_from_sheet(ws) -> List[List[str]]:
    rows = [[cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    return read_table("\n".join("\t".join(c.replace("\t", " ").replace("\n", " ") for c in r) for r in rows))


def _is_text(s: str) -> bool:
    return parse_number(s)[0] == "text"


def _role_of(cell: str, roles: Sequence[str]) -> Optional[str]:
    f = fold(cell)
    if not f:
        return None
    for role in roles:
        for syn in ROLES[role]:
            if f == syn or (len(syn) >= 5 and f.startswith(syn)) or (role == "value" and f.endswith(syn) and len(syn) >= 4):
                return role
    return None


def _month_blocks(cols: Dict[int, int]) -> List[List[int]]:
    """Month columns grouped into runs (e.g. Unadjusted Apr..Mar, then Adjusted Apr..Mar)."""
    out: List[List[int]] = []
    seen: set = set()
    for c in sorted(cols):
        if not out or cols[c] in seen or c - out[-1][-1] > 2:
            out.append([c]); seen = {cols[c]}
        else:
            out[-1].append(c); seen.add(cols[c])
    return out


# --------------------------------------------------------------------------- #
# Matching names
# --------------------------------------------------------------------------- #

def _tokens(s: str) -> List[str]:
    return [t for t in re.findall(r"[a-z0-9]+", (s or "").lower()) if len(t) >= 2]


def _code(s: str) -> str:
    m = re.match(r"\s*#?(\d{3,4})\b", s or "")
    return m.group(1).zfill(4) if m else ""


class Ent:
    __slots__ = ("key", "area", "kind", "id", "label", "aliases", "code", "tokens", "obj")

    def __init__(self, key, area, kind, id_, label, aliases, code="", obj=None):
        self.key, self.area, self.kind, self.id, self.label = key, area, kind, id_, label
        self.aliases = [a for a in dict.fromkeys(fold(x) for x in aliases) if a]
        self.code, self.obj = code, obj
        self.tokens = set(_tokens(" ".join(aliases)))

    def info(self, areas) -> dict:
        a = areas.get(self.area) or {}
        return {"key": self.key, "label": self.label, "area": self.area, "area_label": a.get("name", self.area)}


def plant_ents(aid: str, area: dict) -> List[Ent]:
    out = []
    for p in area.get("plants", []):
        n = p["name"]
        bare = re.sub(r"^\s*\d{3,4}\s*-\s*", "", n)
        bare2 = re.sub(r"\(.*?\)", "", bare).strip()
        no_state = re.sub(r"[\s,]+[A-Z]{2}\s*$", "", bare2)
        out.append(Ent(f"{aid}|p|{p['id']}", aid, "plant", p["id"], n, [n, bare, bare2, no_state], _code(n), p))
    return out


def supply_ents(aid: str, area: dict) -> List[Ent]:
    out = []
    for s in area.get("supply", []):
        loc, sup = s.get("location", ""), s.get("supplier", "")
        out.append(Ent(f"{aid}|s|{s['id']}", aid, "supply", s["id"], f"{loc} | {sup}" if sup else loc,
                       [loc, f"{sup} {loc}", f"{loc} {sup}"], "", s))
    return out


HUB_NAMES = {"CON": "Conway", "MTB": "Mont Belvieu", "MTB-TET": "Mont Belvieu TET", "MTB-NON-T": "Mont Belvieu non-TET"}


def hub_ents(aid: str, area: dict) -> List[Ent]:
    out = []
    for h in area.get("hubs", []):
        nm = h.get("name") or HUB_NAMES.get(h["code"].upper(), "")
        out.append(Ent(f"{aid}|h|{h['code']}", aid, "hub", h["code"], f"{h['code']} hub" + (f" ({nm})" if nm else ""),
                       [h["code"], nm, f"{h['code']} {nm}", f"{nm} hub"], "", h))
    return out


class Matcher:
    def __init__(self, ents: List[Ent]):
        self.ents = ents
        self.by_key = {e.key: e for e in ents}
        self.by_alias: Dict[str, List[Tuple[float, Ent]]] = {}
        self.by_code: Dict[str, List[Ent]] = {}
        self.by_token: Dict[str, List[Ent]] = {}
        for e in ents:
            for n, a in enumerate(e.aliases):
                self.by_alias.setdefault(a, []).append((1.0 if n == 0 else 0.97, e))
            if e.code:
                self.by_code.setdefault(e.code, []).append(e)
            for t in e.tokens:
                if len(t) >= 3:
                    self.by_token.setdefault(t, []).append(e)
        self._cache: Dict[str, List[Tuple[float, Ent]]] = {}

    def scored(self, text: str) -> List[Tuple[float, Ent]]:
        if text in self._cache:
            return self._cache[text]
        f, res = fold(text), {}

        def put(sc, e):
            if e.key not in res or res[e.key][0] < sc:
                res[e.key] = (sc, e)
        if f:
            for sc, e in self.by_alias.get(f, []):
                put(sc, e)
            code = _code(text)
            if code:
                for e in self.by_code.get(code, []):
                    put(0.96, e)
            if not res:
                toks = _tokens(text)
                cand = {e.key: e for t in toks if len(t) >= 3 for e in self.by_token.get(t, [])}
                for e in cand.values():
                    best = 0.0
                    for a in e.aliases:
                        if len(a) >= 5 and len(f) >= 5 and (a in f or f in a):
                            best = max(best, 0.88)
                        best = max(best, difflib.SequenceMatcher(None, f, a).ratio())
                    jac = len(e.tokens & set(toks)) / max(1, len(e.tokens | set(toks)))
                    if jac >= 0.5:
                        best = max(best, 0.5 + 0.45 * jac)
                    put(best, e)
        out = sorted(res.values(), key=lambda x: -x[0])
        self._cache[text] = out
        return out

    def best(self, texts: Sequence[str], restrict=None) -> dict:
        texts = [t for t in texts if t and t.strip()]
        cands = list(texts) + [f"{a} {b}" for a, b in zip(texts, texts[1:])]
        if len(texts) > 2:
            cands.append(" ".join(texts))
        best = None
        for c in cands:
            sc = [(s, e) for s, e in self.scored(c) if restrict is None or restrict(e)]
            if not sc:
                continue
            top = sc[0][0]
            tied = [e for s, e in sc if s >= top - 0.015]
            unique = len(tied) == 1
            rank = (unique and top >= AUTO_OK, top, unique, -len(tied))
            if best is None or rank > best["rank"]:
                best = {"rank": rank, "score": top, "entity": tied[0] if unique else None,
                        "options": [e.key for e in tied[:8]]}
        if best is None:
            return {"status": "unmatched", "score": 0.0, "entity": None, "options": []}
        best.pop("rank")
        if best["score"] < AUTO_CHECK:
            best.update(status="unmatched", entity=None)
        elif best["entity"] is None:
            best["status"] = "ambiguous"
        else:
            best["status"] = "matched" if best["score"] >= AUTO_OK else "check"
        return best


def area_matcher(areas: Dict[str, dict]):
    idx = {}
    for aid, a in areas.items():
        for t in (aid, a.get("name"), a.get("title"), re.sub(r"\s*\d{2}-\d{2}\s*$", "", a.get("name", ""))):
            if fold(t):
                idx.setdefault(fold(t), aid)
    return lambda text: idx.get(fold(text)) or idx.get(fold(re.sub(r"\s*\d{2}-\d{2}\s*$", "", text or "")))


# --------------------------------------------------------------------------- #
# Paths into an area
# --------------------------------------------------------------------------- #

def find(area: dict, kind: str, ident) -> Optional[dict]:
    if kind == "lanes":
        frm, to = ident
        return next((l for l in area.get("lanes", []) if l["from"] == frm and l["to"] == to), None)
    key = "code" if kind == "hubs" else "id"
    return next((x for x in area.get(kind, []) if x[key] == ident), None)


def get_path(area: dict, path: list):
    """['plants', id, field, p] / ['supply', id, field, p] / ['lanes', from, to, field, p] / ['hubs', code, 'price', p]."""
    if path[0] == "lanes":
        obj, rest = find(area, "lanes", (path[1], path[2])), path[3:]
    else:
        obj, rest = find(area, path[0], path[1]), path[2:]
    if obj is None:
        return None
    v = obj.get(rest[0])
    if len(rest) > 1:
        v = v[rest[1]] if isinstance(v, list) and rest[1] < len(v) else None
    return v


def set_path(area: dict, path: list, value) -> None:
    P = len(area["months"])
    if path[0] == "lanes":
        obj, rest = find(area, "lanes", (path[1], path[2])), path[3:]
        if obj is None:
            obj = {"from": path[1], "to": path[2], "rate": [None] * P, "min": None, "max": None}
            area.setdefault("lanes", []).append(obj)
    else:
        obj, rest = find(area, path[0], path[1]), path[2:]
        if obj is None:
            raise KeyError(f"{path[0]} {path[1]} not found")
    if len(rest) == 1:
        obj[rest[0]] = value
        return
    fld, p = rest
    cur = obj.get(fld)
    if not isinstance(cur, list):
        cur = [cur] * P if cur is not None else [None] * P
    cur = (list(cur) + [cur[-1] if cur else None] * P)[:P]
    cur[p] = value
    obj[fld] = cur


# --------------------------------------------------------------------------- #
# Parsing a table into records
# --------------------------------------------------------------------------- #

def _find_header(table, roles, need_name: str):
    """(header row index or None, {role: col}, {col: month})."""
    for r in range(min(12, len(table))):
        row = table[r]
        rmap, mcols = {}, {}
        for c, v in enumerate(row):
            pm = parse_month(v)
            if pm:
                mcols[c] = pm[0]
                continue
            role = _role_of(v, roles)
            if role and role not in rmap:
                rmap[role] = c
        named = need_name in rmap or (need_name == "lane" and ("from" in rmap or "to" in rmap))
        if (named and (mcols or len(rmap) >= 2 or "value" in rmap)) or (len(mcols) >= 2 and rmap):
            return r, rmap, mcols
    return None, {}, {}


def _records(kind: str, table, opts) -> Tuple[List[dict], dict]:
    """Rows -> records {row, area, region, names[], supplier, hub, state, frm[], to[], vals{mon|-1: (kind, v)}, min, max, open, season}."""
    k = KINDS[kind]
    ent = k["entity"]
    roles = KIND_ROLES[ent]
    name_role = {"plant": "plant", "supply": "location", "lane": "lane", "hub": "hub", "state": "state", "weather": "place"}[ent]
    hr, rmap, mcols = _find_header(table, roles, name_role)
    info: dict = {"layout": "", "header_row": hr}
    if ent == "lane" and hr is not None and "to" not in rmap:
        info["matrix"] = True
    if hr is None:
        # no header: first text column(s) = name, numeric columns = months (12) or one value
        hr = -1
        width = len(table[0]) if table else 0
        txt = [c for c in range(width) if all(_is_text(r[c]) or not r[c] for r in table)]
        num = [c for c in range(width) if c not in txt]
        if not txt or not num:
            raise ValueError("couldn't find a header row or the names - copy the header row too")
        rmap = {name_role if name_role != "lane" else "from": txt[0]}
        if ent == "lane" and len(txt) > 1:
            rmap["to"] = txt[-1]
        if len(num) == 12:
            mcols = {c: (n + 3) % 12 + 1 for n, c in enumerate(num)}      # assume Apr..Mar
            info["notes"] = ["No header row: the 12 number columns were read as Apr to Mar."]
        else:
            rmap["value"] = num[-1]
    blocks = _month_blocks(mcols) if mcols else []
    if len(blocks) > 1:
        bi = min(max(int(opts.get("block") or 0), 0), len(blocks) - 1)
        info["blocks"] = [{"index": n, "label": f"columns {col_name(b[0])}\u2013{col_name(b[-1])}"
                           + (f" ({table[hr - 1][b[0]]})" if hr > 0 and table[hr - 1][b[0]] else "")}
                          for n, b in enumerate(blocks)]
        mcols = {c: mcols[c] for c in blocks[bi]}
    info["months_found"] = sorted(set(mcols.values()), key=lambda m: SEASON_ORDER.index(MONTHS[m - 1]))
    month_role = rmap.get("month")
    records = []
    for r in range(hr + 1, len(table)):
        row = table[r]
        g = lambda role: row[rmap[role]].strip() if role in rmap else ""
        rec = {"row": r + 1, "area": g("area"), "region": g("region"), "supplier": g("supplier"),
               "hub": g("hub"), "hubname": g("hubname"), "state": g("state"), "season": g("season"),
               "open": parse_number(g("open")) if "open" in rmap else None,
               "min": parse_number(g("min")) if "min" in rmap else None,
               "max": parse_number(g("max")) if "max" in rmap else None}
        if ent == "plant":
            rec["names"] = [x for x in [g("plant")] if x]
            if not rec["names"] and "location" in rmap:
                rec["names"] = [g("location")]
        elif ent == "supply":
            rec["names"] = [x for x in [g("location"), g("supplier")] if x]
        elif ent == "lane":
            rec["frm"] = [x for x in [g("from"), g("supplier")] if x]
            rec["to"] = [x for x in [g("to")] if x]
        elif ent == "hub":
            rec["names"] = [x for x in [g("hub"), g("hubname")] if x]
        elif ent == "state":
            rec["names"] = [x for x in [g("state")] if x]
        else:
            rec["names"] = [x for x in [g("place")] if x]
        first = (rec.get("names") or rec.get("frm") or [""])
        if not any(rec.get(x) for x in ("names", "frm", "to")) or (first and _SKIP.match(first[0])):
            if not info.get("matrix"):
                continue
        vals = {}
        if mcols:
            for c, mon in mcols.items():
                vals[mon] = parse_number(row[c])
        elif month_role is not None and "value" in rmap:
            pm = parse_month(row[month_role])
            if pm:
                vals[pm[0]] = parse_number(row[rmap["value"]])
        elif "value" in rmap:
            vals[-1] = parse_number(row[rmap["value"]])
        rec["vals"] = vals
        rec["cells"] = row
        records.append(rec)
    if month_role is not None:                         # long layout: merge rows of one name
        merged: Dict[tuple, dict] = {}
        for rec in records:
            key = (rec["area"], tuple(rec.get("names", [])), tuple(rec.get("frm", [])), tuple(rec.get("to", [])), rec["season"])
            if key in merged:
                merged[key]["vals"].update(rec["vals"])
            else:
                merged[key] = rec
        records = list(merged.values())
        info["layout"] = "one row per name and month"
    elif info.get("matrix"):
        info["layout"] = "grid: supply points down, plants across"
    elif mcols:
        info["layout"] = "months across"
    else:
        info["layout"] = "one value per row"
        info["single"] = True
    info["rmap"], info["mcols"], info["hr"] = rmap, mcols, hr
    return records, info


# --------------------------------------------------------------------------- #
# Units
# --------------------------------------------------------------------------- #

def _unit(kind: str, numbers: List[float], pairs: List[Tuple[float, float]], opts, notes) -> Tuple[float, str]:
    """Factor from what was pasted to app units (gal, $/gal, fraction) and the unit's name."""
    u = KINDS[kind]["unit"]
    want = opts.get("unit") or "auto"
    if u == "pct":
        if want in ("percent", "fraction"):
            return (0.01 if want == "percent" else 1.0), want
        big = any(abs(v) > 3 for v in numbers)
        return (0.01, "percent") if big else (1.0, "fraction")
    if u == "raw" or u is None:
        return 1.0, ""
    choices = {"volume": {"gal": 1.0, "kgal": 1000.0}, "price": {"$/gal": 1.0, "$/kgal": 0.001}}[u]
    if want in choices:
        return choices[want], want
    good = [(v, c) for v, c in pairs if v and c and math.isfinite(c) and abs(c) > 1e-9]
    if len(good) >= 3:
        best = min(choices.items(), key=lambda kv: float(np.median([abs(math.log10(abs(v * kv[1] / c))) for v, c in good])))
        notes.append(f"Units: read as {best[0]} (compared with what's there now).")
        return best[1], best[0]
    nz = [abs(v) for v in numbers if v]
    med = float(np.median(nz)) if nz else 0.0
    if u == "volume":
        name = "kgal" if 0 < med < 3000 else "gal"
    else:
        name = "$/kgal" if med > 5 else "$/gal"
    notes.append(f"Units: read as {name} from the size of the numbers - change it if that's wrong.")
    return choices[name], name


# --------------------------------------------------------------------------- #
# Parse -> preview + operations
# --------------------------------------------------------------------------- #

def parse(kind: str, text: str, ctx: dict, scope: Optional[str] = None, opts: Optional[dict] = None,
          mapping: Optional[dict] = None, table: Optional[List[List[str]]] = None) -> dict:
    """
    ctx = {"areas": {id: area}, "state_pct": {...}, "months": [...], "weather": {...}}
    scope = an area id (only that area) or None (every area; rows find their area).
    """
    if kind not in KINDS:
        raise ValueError("pick what you're pasting")
    opts, mapping = dict(opts or {}), mapping or {}
    table = table if table is not None else read_table(text or "")
    if not table:
        raise ValueError("nothing to read - copy the cells in Excel (with the header row) and paste them here")
    k = KINDS[kind]
    records, info = _records(kind, table, opts)
    notes: List[str] = list(info.get("notes", []))
    areas = ctx["areas"]
    scope_ids = [scope] if scope else list(areas)
    if scope and scope not in areas:
        raise ValueError("unknown area")
    if k["entity"] in ("state", "weather"):
        return _parse_table_kind(kind, records, info, ctx, opts, notes)
    return _parse_area_kind(kind, records, info, ctx, scope_ids, scope, opts, mapping, notes, table)


def _months_for(area: dict) -> Dict[int, int]:
    from engine import month_num
    return {month_num(l, p): p for p, l in enumerate(area["months"])}


def _pick_months(opts) -> Optional[set]:
    m = opts.get("months")
    if m in (None, "", "all"):
        return None
    return {int(x) for x in m}


def _parse_area_kind(kind, records, info, ctx, scope_ids, scope, opts, mapping, notes, table):
    k = KINDS[kind]
    ent, fld = k["entity"], k["field"]
    areas = ctx["areas"]
    find_area = area_matcher({a: areas[a] for a in scope_ids})
    plants = Matcher([e for a in scope_ids for e in plant_ents(a, areas[a])])
    supplies = Matcher([e for a in scope_ids for e in supply_ents(a, areas[a])])
    hubs = Matcher([e for a in scope_ids for e in hub_ents(a, areas[a])])
    main = {"plant": plants, "supply": supplies, "hub": hubs, "lane": supplies}[ent]
    # Lists (plants, supply points, hubs) add what's new. Number uploads (budget,
    # contracts) only add new plants / supply points when asked to ("add_new"),
    # so a stray row can't quietly create a plant with no lanes.
    creatable = kind in ("plants", "supply", "hubs", "lanes") or (kind in ("budget", "contract") and bool(opts.get("add_new")))
    names: Dict[Tuple[str, str], dict] = {}
    new_areas: Dict[str, dict] = {}

    def decide(side, text, area_hint, m, matcher, allow_new):
        key = f"{area_hint}::{text}" if area_hint else text
        rec = names.get((side, key))
        if rec is None:
            status = m["status"]
            if status == "unmatched" and allow_new:
                status = "new"
            rec = names[(side, key)] = {"side": side, "key": key, "text": text, "area": area_hint, "status": status,
                                        "score": round(m["score"], 2), "suggested": m["entity"].key if m["entity"] else "",
                                        "options": m.get("options", []), "rows": 0, "choice": None}
        rec["rows"] += 1
        over = (mapping.get(side) or {}).get(key)
        if over is not None:
            rec["choice"] = over
            if over == "new":
                return "new"
            return matcher.by_key.get(over) if over else None
        if rec["status"] == "new":
            return "new"
        return m["entity"]

    # ---- area of each record ---------------------------------------------- #
    def area_of(rec) -> Optional[str]:
        if scope:
            return scope
        if rec["area"]:
            aid = find_area(rec["area"])
            if aid:
                return aid
            if creatable and kind in ("plants", "budget", "supply", "contract"):
                nid = fold(rec["area"])[:40] or "area"
                if nid not in new_areas:
                    new_areas[nid] = {"id": nid, "name": rec["area"], "region": rec["region"] or "Other",
                                      "months": list(ctx["months"]), "supply": [], "plants": [], "lanes": [], "hubs": []}
                return nid
            return "?"
        return None

    # ---- matrix lanes: the header row's plant names ----------------------- #
    matrix_cols = {}
    if info.get("matrix"):
        hr = info["hr"]
        for c, v in enumerate(table[hr]):
            if c in info["rmap"].values() or c in info["mcols"] or not v:
                continue
            matrix_cols[c] = v
        notes.append("Read as a grid: one freight rate per supply point and plant. Pick the months it applies to.")
        info["single"] = True

    ops: List[dict] = []
    creates: List[dict] = []
    new_objs: Dict[str, dict] = {}            # key -> object being created in this upload
    pairs: List[Tuple[float, float]] = []
    numbers: List[float] = []
    resolved = []                              # (rec, aid, entity-or-new-obj, other)

    for rec in records:
        aid = area_of(rec)
        if aid == "?":
            names.setdefault(("area", rec["area"]), {"side": "area", "key": rec["area"], "text": rec["area"], "area": "",
                                                    "status": "unmatched", "score": 0, "suggested": "", "options": [],
                                                    "rows": 0, "choice": None})["rows"] += 1
            continue
        restrict = (lambda e, a=aid: e.area == a) if aid else None
        if ent == "lane":
            targets = [(c, [v]) for c, v in matrix_cols.items()] if matrix_cols else [(None, rec["to"])]
            for col, to_texts in targets:
                mp = plants.best(to_texts, restrict=restrict)
                ep = decide("to", " \u00b7 ".join(to_texts), aid or "", mp, plants, False)
                if ep is None or ep == "new":
                    continue
                pa = ep.area
                ms = supplies.best(rec["frm"], restrict=lambda e, a=pa: e.area == a)
                es = decide("from", " \u00b7 ".join(rec["frm"]), pa if not scope else "", ms, supplies, False)
                if es is None or es == "new":
                    continue
                vals = {-1: parse_number(rec["cells"][col])} if col is not None else rec["vals"]
                resolved.append((rec, pa, es, ep, vals))
        else:
            text = " \u00b7 ".join(rec["names"])
            m = main.best(rec["names"], restrict=restrict)
            e = decide("name", text, aid or "", m, main, creatable and aid is not None)
            if e is None:
                continue
            if e == "new":
                if aid is None:
                    continue
                nk = f"{aid}|new|{fold(text)}"
                if nk not in new_objs:
                    new_objs[nk] = _new_obj(kind, rec, areas.get(aid) or new_areas.get(aid), ctx)
                    creates.append({"area": aid, "type": ent, "label": new_objs[nk].get("name") or new_objs[nk].get("location") or new_objs[nk].get("code")})
                resolved.append((rec, aid, new_objs[nk], None, rec["vals"]))
            else:
                resolved.append((rec, e.area, e, None, rec["vals"]))

    # ---- units ------------------------------------------------------------ #
    for rec, aid, e, ep, vals in resolved:
        area = areas.get(aid) or new_areas.get(aid)
        mp = _months_for(area)
        for mon, (kd, v) in vals.items():
            if kd == "num" and v:
                numbers.append(v)
                if isinstance(e, Ent) and fld:
                    p = mp.get(mon, 0) if mon != -1 else 0
                    cur = _current(area, ent, e, ep, fld, p)
                    if cur is not None:
                        pairs.append((v, cur))
    factor, unit_name = _unit(kind, numbers, pairs, opts, notes) if fld or kind in ("plants", "supply") else (1.0, "")
    vol_factor = 1.0
    if kind in ("plants", "supply") or ent == "lane":   # opening inventory / lane min-max are volumes
        vol_factor = 1000.0 if opts.get("vol_unit") == "kgal" else 1.0
    months_sel = _pick_months(opts)
    blanks = opts.get("blanks", "skip")

    # ---- operations -------------------------------------------------------- #
    changes, given = 0, 0
    by_area: Dict[str, int] = {}
    biggest = []
    seen_cells: Dict[tuple, int] = {}
    dup = 0
    for obj in new_objs.values():
        pass
    for nk, obj in new_objs.items():
        aid = nk.split("|", 1)[0]
        ops.append({"op": "create", "area": aid, "type": ent, "obj": obj})
    for aid, a in new_areas.items():
        ops.insert(0, {"op": "create_area", "area": aid, "obj": a})

    for rec, aid, e, ep, vals in resolved:
        area = areas.get(aid) or new_areas.get(aid)
        mp = _months_for(area)
        P = len(area["months"])
        if isinstance(e, Ent):
            ident = e.id
        else:
            ident = e.get("id") or e.get("code")
        # list kinds: attributes
        if kind == "plants" or kind == "budget":
            if rec.get("state") and isinstance(e, Ent) and (e.obj.get("state") or "") != rec["state"].upper()[:2]:
                ops.append({"op": "set", "area": aid, "path": ["plants", ident, "state"], "value": rec["state"].upper()[:2]})
                changes += 1
        if kind in ("plants", "supply") and rec.get("open") and rec["open"][0] == "num":
            nv = rec["open"][1] * vol_factor
            old = e.obj.get("open") if isinstance(e, Ent) else None
            if not isinstance(e, Ent) or not _same(old or 0.0, nv, "volume"):
                ops.append({"op": "set", "area": aid, "path": [_coll(ent), ident, "open"], "value": nv})
                changes += 1
        if kind == "supply" and isinstance(e, Ent):
            for attr in ("supplier", "hub"):
                if rec.get(attr) and (e.obj.get(attr) or "") != rec[attr]:
                    ops.append({"op": "set", "area": aid, "path": ["supply", ident, attr], "value": rec[attr]})
                    changes += 1
        if not fld:
            continue
        lane_ident = (e.id, ep.id) if ent == "lane" else None
        if ent == "lane":
            exists = find(area, "lanes", lane_ident)
            if exists is None:
                lk = (aid, lane_ident)
                if lk not in seen_cells:
                    creates.append({"area": aid, "type": "lane", "label": f"{e.label} \u2192 {ep.label}"})
                seen_cells[lk] = 1
            for lim in ("min", "max"):
                pv = rec.get(lim)
                if pv and pv[0] == "num":
                    for p in range(P):
                        if months_sel is None or _mon_of(area, p) in months_sel:
                            ops.append({"op": "set", "area": aid, "path": ["lanes", e.id, ep.id, lim, p], "value": pv[1] * vol_factor})
        for mon, (kd, v) in vals.items():
            if mon == -1:
                ps = [p for p in range(P) if months_sel is None or _mon_of(area, p) in months_sel]
            else:
                ps = [mp[mon]] if mon in mp else []
            if kd == "text":
                continue
            if kd == "blank":
                if blanks != "clear":
                    continue
                val = None if fld in ("rate",) else 0.0
            else:
                val = v * (1.0 if kd == "pct" and k["unit"] == "pct" else factor)
                if k["unit"] == "volume":
                    val = round(val, 6)
            for p in ps:
                path = (["lanes", e.id, ep.id, fld, p] if ent == "lane"
                        else ["hubs", ident, "price", p] if ent == "hub"
                        else [_coll(ent), ident, fld, p])
                ck = (aid, tuple(path))
                if ck in seen_cells and seen_cells[ck] != rec["row"]:
                    dup += 1
                seen_cells[ck] = rec["row"]
                given += 1
                cur = get_path(area, path) if isinstance(e, Ent) and (ent != "lane" or find(area, "lanes", lane_ident)) else None
                if _same(cur, val, k["unit"]):
                    continue
                ops.append({"op": "set", "area": aid, "path": path, "value": val})
                changes += 1
                by_area[aid] = by_area.get(aid, 0) + 1
                biggest.append((abs((val or 0) - (cur or 0)) if isinstance(cur, (int, float)) or cur is None else 0,
                                aid, _where(area, ent, e, ep, p), cur, val))
    if dup:
        notes.append(f"{dup:,} cell(s) were given more than once (same name on two rows); the last one is used.")
    if info.get("single") and fld and months_sel is None:
        notes.append("The paste has one value per row, so it's used for every month. Pick months below to limit it.")
    biggest.sort(key=lambda x: -x[0])
    name_list = sorted(names.values(), key=lambda r: ({"unmatched": 0, "ambiguous": 1, "other area": 2, "check": 3,
                                                       "new": 4, "matched": 5}.get(r["status"], 6), r["text"].lower()))
    ent_lists = {"name": [e.info(areas) for e in main.ents], "to": [e.info(areas) for e in plants.ents],
                 "from": [e.info(areas) for e in supplies.ents]}
    return {"kind": kind, "label": k["label"], "layout": info["layout"], "info": {x: info[x] for x in
            ("blocks", "months_found", "single", "matrix") if x in info},
            "unit": unit_name, "unit_kind": k["unit"], "rows_read": len(records), "notes": notes,
            "names": name_list, "creates": creates, "changes": changes, "cells_given": given,
            "areas": [{"area": a, "label": (areas.get(a) or new_areas.get(a, {})).get("name", a), "cells": n}
                      for a, n in sorted(by_area.items())],
            "new_areas": [a["name"] for a in new_areas.values()],
            "biggest": [{"area": a, "where": w, "old": o, "new": n} for _, a, w, o, n in biggest[:15]],
            "entities": ent_lists, "_ops": ops}


def _coll(ent: str) -> str:
    return {"plant": "plants", "supply": "supply", "hub": "hubs"}[ent]


def _mon_of(area: dict, p: int) -> int:
    from engine import month_num
    return month_num(area["months"][p], p)


def _same(a, b, unit) -> bool:
    an = a is None or (isinstance(a, float) and not math.isfinite(a))
    bn = b is None or (isinstance(b, float) and not math.isfinite(b))
    if an or bn:
        return an and bn
    tol = {"volume": 0.5, "price": 5e-6, "pct": 1e-6}.get(unit, 1e-9)
    return abs(float(a) - float(b)) <= max(tol, 1e-9 * abs(float(b)))


def _current(area, ent, e: Ent, ep, fld, p):
    if ent == "lane":
        if ep is None or not isinstance(ep, Ent):
            return None
        ln = find(area, "lanes", (e.id, ep.id))
        v = (ln or {}).get("rate")
    elif ent == "hub":
        v = e.obj.get("price")
    else:
        v = e.obj.get(fld)
    if isinstance(v, list):
        v = v[p] if p < len(v) else None
    return float(v) if isinstance(v, (int, float)) else None


def _where(area, ent, e, ep, p) -> str:
    lab = lambda x: x.label if isinstance(x, Ent) else (x.get("name") or x.get("location") or x.get("code", ""))
    m = area["months"][p] if p is not None and p < len(area["months"]) else ""
    if ent == "lane":
        return f"{lab(e)} \u2192 {lab(ep)} \u00b7 {m}"
    return f"{lab(e)} \u00b7 {m}"


def _new_obj(kind: str, rec: dict, area: dict, ctx: dict) -> dict:
    P = len(area["months"])
    ent = KINDS[kind]["entity"]
    if ent == "plant":
        n = rec["names"][0]
        st = (rec.get("state") or "").upper()[:2]
        if not st:
            m = re.findall(r"\b([A-Z]{2})\b", n)
            st = m[-1] if m else ""
        return {"id": new_id("p"), "name": n, "state": st, "budget": [0.0] * P, "open": 0.0, "cap": [0.0] * P, "inv_min": [0.0] * P}
    if ent == "supply":
        loc = rec["names"][0]
        return {"id": new_id("s"), "location": loc, "supplier": rec.get("supplier") or "", "hub": rec.get("hub") or "",
                "contract": [0.0] * P, "diff": [0.0] * P, "min_pct": [0.0] * P, "max_pct": [1.0] * P,
                "open": 0.0, "cap": [0.0] * P, "inv_min": [0.0] * P}
    if ent == "hub":
        code = rec["names"][0].upper()
        return {"code": code, "id": code, "name": rec.get("hubname") or HUB_NAMES.get(code, ""), "price": [None] * P}
    raise ValueError(kind)


def _parse_table_kind(kind, records, info, ctx, opts, notes):
    """State % and weather history: dataset-wide tables, not tied to an area."""
    k = KINDS[kind]
    numbers = [v for r in records for kd, v in r["vals"].values() if kd == "num"]
    ops, changes, given = [], 0, 0
    biggest = []
    if kind == "state_pct":
        cur = ctx.get("state_pct") or {}
        pairs = [(v, cur.get((r["names"] or [""])[0].upper(), {}).get(MONTHS[mon - 1]))
                 for r in records for mon, (kd, v) in r["vals"].items() if kd == "num" and mon != -1]
        factor, unit_name = _unit(kind, numbers, [x for x in pairs if x[1]], opts, notes)
        months_sel = _pick_months(opts)
        for r in records:
            st = (r["names"] or [""])[0].strip().upper()
            st = {"WESTTX": "W TX"}.get(fold(st).upper(), st)
            if not st:
                continue
            for mon, (kd, v) in r["vals"].items():
                if kd not in ("num", "pct"):
                    continue
                mons = [m for m in range(1, 13) if months_sel is None or m in months_sel] if mon == -1 else [mon]
                val = v if kd == "pct" else v * factor
                for m in mons:
                    given += 1
                    old = (cur.get(st) or {}).get(MONTHS[m - 1])
                    if _same(old, val, "pct"):
                        continue
                    ops.append({"op": "state_pct", "state": st, "month": MONTHS[m - 1], "value": val})
                    changes += 1
                    biggest.append((abs(val - (old or 0)), f"{st} \u00b7 {MONTHS[m - 1]}", old, val))
        biggest.sort(key=lambda x: -x[0])
        return {"kind": kind, "label": k["label"], "layout": info["layout"], "info": {}, "unit": unit_name,
                "unit_kind": "pct", "rows_read": len(records), "notes": notes, "names": [], "creates": [],
                "changes": changes, "cells_given": given, "areas": [], "new_areas": [],
                "biggest": [{"area": "", "where": w, "old": o, "new": n} for _, w, o, n in biggest[:15]],
                "entities": {}, "_ops": ops}
    # weather history
    rows = []
    for r in records:
        place = (r["names"] or [""])[0]
        season = r.get("season") or ""
        if not place or not season:
            continue
        vals = {MONTHS[mon - 1]: v for mon, (kd, v) in r["vals"].items() if kd in ("num", "pct") and mon != -1}
        if vals:
            rows.append({"key": place, "season": season, "values": vals})
    unit = opts.get("weather_unit") or ("hdd" if numbers and np.median(numbers) < 2500 and max(numbers) < 3000 else "gal")
    if not rows:
        notes.append("No rows with a place, a season and monthly values were found. The header row needs "
                     "Place (or Plant / State / Area), Season, and month columns.")
    from weather import season_label
    seasons = sorted({season_label(r["season"]) for r in rows})
    ops.append({"op": "weather", "rows": rows, "unit": unit, "replace": opts.get("replace", True)})
    return {"kind": kind, "label": k["label"], "layout": info["layout"], "info": {}, "unit": unit, "unit_kind": "raw",
            "rows_read": len(records), "notes": notes, "names": [], "creates": [],
            "changes": len(rows), "cells_given": sum(len(r["values"]) for r in rows), "areas": [], "new_areas": [],
            "weather": {"places": len({r["key"] for r in rows}), "seasons": seasons[:1] + seasons[-1:], "n_seasons": len(seasons)},
            "biggest": [], "entities": {}, "_ops": ops}


# --------------------------------------------------------------------------- #
# Applying
# --------------------------------------------------------------------------- #

def apply_ops(ctx: dict, ops: List[dict], user: str = "", ts: str = "") -> Dict[str, int]:
    """Carry out a parse's operations on ctx (mutated). Returns cells changed per area ('*' for tables)."""
    areas = ctx["areas"]
    counts: Dict[str, int] = {}
    for op in ops:
        kind = op["op"]
        if kind == "create_area":
            if op["area"] not in areas:
                areas[op["area"]] = op["obj"]
        elif kind == "create":
            a = areas[op["area"]]
            coll = {"plant": "plants", "supply": "supply", "hub": "hubs"}[op["type"]]
            a.setdefault(coll, []).append(op["obj"])
        elif kind == "set":
            set_path(areas[op["area"]], op["path"], op["value"])
            counts[op["area"]] = counts.get(op["area"], 0) + 1
        elif kind == "state_pct":
            ctx.setdefault("state_pct", {}).setdefault(op["state"], {})[op["month"]] = op["value"]
            counts["*"] = counts.get("*", 0) + 1
        elif kind == "weather":
            cur = ctx.get("weather") or {}
            rows = op["rows"] if op.get("replace", True) else (cur.get("rows") or []) + op["rows"]
            ctx["weather"] = {"rows": rows, "unit": op["unit"], "detrend": op["unit"] == "gal",
                              "by": user, "ts": ts}
            counts["*"] = counts.get("*", 0) + len(op["rows"])
    return counts


# --------------------------------------------------------------------------- #
# Excel workbooks
# --------------------------------------------------------------------------- #

SHEET_KIND = {fold(KINDS[k]["label"]): k for k in KINDS}
SHEET_KIND.update({fold("State Adjustments"): "state_pct", fold("Demand Calculations"): "budget",
                   fold("Freight"): "lanes", fold("Freight rates"): "lanes", fold("Lanes"): "lanes",
                   fold("OD pairs"): "lanes", fold("Contracts"): "contract", fold("Demand"): "budget",
                   fold("Weather"): "weather", fold("Hubs"): "hubs"})
SKIP_SHEETS = {fold("Read me"), fold("Demand For Models"), fold("Model Reference")}


def guess_kind(name: str, table) -> Optional[str]:
    f = fold(name)
    if f in SKIP_SHEETS:
        return None
    if f in SHEET_KIND:
        return SHEET_KIND[f]
    for k, v in SHEET_KIND.items():
        if k and (k in f):
            return v
    head = " ".join(fold(c) for r in table[:6] for c in r)
    if "season" in head:
        return "weather"
    if ("from" in head and "to" in head) or "freight" in head:
        return "lanes"
    if "hub" in head and "location" not in head and "supplier" not in head:
        return "hubs"
    if "contract" in head and ("location" in head or "supplier" in head):
        return "contract"
    if "plant" in head or "district" in head:
        return "budget"
    if "state" in head:
        return "state_pct"
    return None


def read_workbook(path: str) -> List[dict]:
    """[{sheet, kind (guess), table}] for every sheet that has data."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            t = table_from_sheet(ws)
            if len(t) < 2:
                continue
            out.append({"sheet": ws.title, "kind": guess_kind(ws.title, t), "table": t, "rows": len(t)})
    finally:
        wb.close()
    # Two sheets that look like the same kind (e.g. "Demand Calculations" and
    # "Demand Summary - FY27"): keep the one whose name says so, and leave the
    # others un-ticked with the guess shown, so nothing is loaded twice.
    by_kind: Dict[str, List[dict]] = {}
    for s in out:
        if s["kind"]:
            by_kind.setdefault(s["kind"], []).append(s)
    for kind, ss in by_kind.items():
        if len(ss) < 2:
            continue
        named = [s for s in ss if SHEET_KIND.get(fold(s["sheet"])) == kind]
        keep = named[0] if named else ss[0]
        for s in ss:
            if s is not keep:
                s["maybe"], s["kind"] = kind, None
    return out


def export_workbook(path: str, areas: Dict[str, dict], state_pct: dict, weather: Optional[dict],
                    kinds: Optional[Sequence[str]] = None, title: str = "") -> str:
    """Every data kind as a sheet in the layout the importer reads (gallons, $/gal, % as 75)."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    kinds = list(kinds or ORDER)
    wb = openpyxl.Workbook()
    rd = wb.active
    rd.title = "Read me"
    lines = [title or "Supply planning data", "",
             "Each sheet is one kind of data, in the layout the app reads back in.",
             "Edit offline, then upload this workbook (Data uploads > Upload Excel) - the app shows what changes before saving.",
             "Units: gallons, $ per gallon, percentages as 75 (= 75%). A blank freight rate means the lane isn't used that month.", ""]
    for k in kinds:
        lines.append(f"{KINDS[k]['label']}: {KINDS[k]['help']}")
    for r, t in enumerate(lines, start=1):
        rd.cell(row=r, column=1, value=t)
    rd["A1"].font = Font(bold=True, size=14)
    rd.column_dimensions["A"].width = 140
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    any_area = next(iter(areas.values()), None)
    months = any_area["months"] if any_area else ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]

    def sheet(name, header, rows, fmt=None):
        ws = wb.create_sheet(name[:31])
        ws.append(header)
        for row in rows:
            ws.append(row)
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(c)].width = 28 if c <= 3 else 12
        if fmt:
            for row in ws.iter_rows(min_row=2, min_col=fmt[0], max_col=len(header)):
                for cell in row:
                    cell.number_format = fmt[1]
        ws.freeze_panes = ws.cell(row=2, column=fmt[0] if fmt else 2)

    def series(v, P):
        v = v if isinstance(v, list) else [v] * P
        return [None if x is None else x for x in (list(v) + [None] * P)[:P]]

    alist = sorted(areas.values(), key=lambda a: (a.get("region", ""), a.get("name", "")))
    for k in kinds:
        lab = KINDS[k]["label"]
        if k == "plants":
            sheet(lab, ["Region", "Area", "Plant", "State", "Opening inventory (gal)"],
                  [[a.get("region", ""), a["name"], p["name"], p.get("state", ""), p.get("open") or 0] for a in alist for p in a.get("plants", [])], (5, "#,##0"))
        elif k in ("budget", "plant_cap", "plant_min"):
            f = KINDS[k]["field"]
            sheet(lab, ["Area", "Plant", "State"] + months,
                  [[a["name"], p["name"], p.get("state", "")] + series(p.get(f), len(a["months"])) for a in alist for p in a.get("plants", [])], (4, "#,##0"))
        elif k == "supply":
            sheet(lab, ["Region", "Area", "Location", "Supplier", "Hub", "Opening inventory (gal)"],
                  [[a.get("region", ""), a["name"], s["location"], s.get("supplier", ""), s.get("hub", ""), s.get("open") or 0]
                   for a in alist for s in a.get("supply", [])], (6, "#,##0"))
        elif k in ("contract", "diff", "min_pct", "max_pct", "supply_cap", "supply_min"):
            f = KINDS[k]["field"]
            pct = KINDS[k]["unit"] == "pct"
            rows = []
            for a in alist:
                for s in a.get("supply", []):
                    vals = series(s.get(f), len(a["months"]))
                    rows.append([a["name"], s["location"], s.get("supplier", "")] + [None if v is None else (v * 100 if pct else v) for v in vals])
            sheet(lab, ["Area", "Location", "Supplier"] + months, rows,
                  (4, "0.0000" if KINDS[k]["unit"] == "price" else "0.##" if pct else "#,##0"))
        elif k == "hubs":
            sheet(lab, ["Area", "Hub", "Hub name"] + months,
                  [[a["name"], h["code"], h.get("name", "")] + series(h.get("price"), len(a["months"])) for a in alist for h in a.get("hubs", [])], (4, "0.0000"))
        elif k == "lanes":
            rows = []
            for a in alist:
                S = {s["id"]: s for s in a.get("supply", [])}
                Pl = {p["id"]: p for p in a.get("plants", [])}
                for l in a.get("lanes", []):
                    s, p = S.get(l["from"]), Pl.get(l["to"])
                    if not s or not p:
                        continue
                    one = lambda v: (v[0] if isinstance(v, list) and len(set(map(str, v))) == 1 else None if isinstance(v, list) else v)
                    rows.append([a["name"], s["location"], s.get("supplier", ""), p["name"]] + series(l.get("rate"), len(a["months"]))
                                + [one(l.get("min")), one(l.get("max"))])
            sheet(lab, ["Area", "From", "Supplier", "To"] + months + ["Min gal", "Max gal"], rows, (5, "0.0000"))
        elif k == "state_pct":
            rows = [[st] + [None if (v := (state_pct.get(st) or {}).get(m)) is None else v * 100 for m in SEASON_ORDER]
                    for st in sorted(state_pct)]
            sheet(lab, ["State"] + SEASON_ORDER, rows, (2, "0.#"))
        elif k == "weather":
            rows = [[r["key"], r["season"]] + [r["values"].get(m) for m in SEASON_ORDER] for r in (weather or {}).get("rows", [])]
            sheet(lab, ["Place", "Season"] + SEASON_ORDER, rows, (3, "#,##0.##"))
    wb.save(path)
    return path
