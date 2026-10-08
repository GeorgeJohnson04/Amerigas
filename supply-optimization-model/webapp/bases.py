"""
bases.py - locked base models, one set per contract year.

How the yearly models have been made in Excel (from the team's process notes
and the workbooks themselves):

  * BASELINE for a new year = last year's WORKING model, with the NEW year's
    demand, but last year's contract volumes, differentials and freight.
  * WORKING model = that baseline plus each manager's updates: new supply
    points, the agreed basis curves, new freight rates (e.g. +3%), wholesale
    volume projections.

Every number the model needs is typed into the ModelBase sheet (the rest of
the sheet is formulas that copy or add those numbers up), so a model is fully
described by its raw inputs: contracts, hub prices, differentials, lift
limits, demand, freight, lane limits and storage. A base model here is
exactly that set of raw inputs, saved as JSON - no Excel needed to use it or
to make next year's.

    data/bases/<base id>/meta.json            what it is, who made it, how
    data/bases/<base id>/<region>__<area>.json one file per area

Bases are LOCKED: there is no way to change one through the app, and the files
are written read-only. Managers work in scenarios ("working models") that sit
on top of a base and only store their own changes.

Ways to make a base:

  import_workbooks()   read the area workbooks in the Model folder (Baseline
                       or Working files) - how the first base is made
  roll_forward()       make next year's base from a base or a manager's working
                       model: new demand from a "Demand for Models" file,
                       contracts / differentials / freight kept (optionally
                       freight +x%), months moved on a year
"""

from __future__ import annotations

import datetime as _dt
import json
import logging
import os
import re
import shutil
import stat
import threading
import uuid
from typing import Callable, Dict, List, Optional, Sequence, Tuple

import numpy as np

import model

LOG = logging.getLogger("webapp.bases")

#: Raw-input arrays that make up a model (the keys of ModelBaseReader's arrays).
ARRAY_KEYS = ("contract", "hub_price", "hub_premium", "perc_min", "perc_max", "demand",
              "trans_cost", "ship_min", "ship_max", "supply_inv_init", "supply_inv_cap",
              "supply_inv_min", "demand_inv_init", "demand_inv_cap", "demand_inv_min")

_SEASON_TOKEN = re.compile(r"(\d{2})-(\d{2})\s*$")
_MONTHS = ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]


def now_iso() -> str:
    return _dt.datetime.now().isoformat(timespec="seconds")


def _tolist(a) -> list:
    """numpy -> JSON-able nested lists, NaN as null (a blank freight cell stays blank)."""
    a = np.asarray(a, dtype=float)
    out = a.astype(object)
    out[~np.isfinite(a)] = None
    return out.tolist()


def season_of(area: str) -> str:
    """'MT-WY 26-27' -> '26-27' ('' if the name has no season)."""
    m = _SEASON_TOKEN.search(area)
    return f"{m.group(1)}-{m.group(2)}" if m else ""


def next_season(season: str) -> str:
    """'26-27' -> '27-28'."""
    m = re.fullmatch(r"(\d{2})-(\d{2})", season.strip())
    if not m:
        raise ValueError(f"season should look like 26-27, got {season!r}")
    a = (int(m.group(1)) + 1) % 100
    return f"{a:02d}-{(a + 1) % 100:02d}"


def normalise_season(text: str) -> str:
    """'2027-28', '27-28', '2027/2028' -> '27-28'."""
    nums = re.findall(r"\d+", text or "")
    if len(nums) < 1:
        raise ValueError("give the contract year, e.g. 2027-28")
    a = int(nums[0]) % 100
    return f"{a:02d}-{(a + 1) % 100:02d}"


def rename_area(area: str, season: str) -> str:
    """'MT-WY 26-27' + '27-28' -> 'MT-WY 27-28'."""
    return _SEASON_TOKEN.sub(season, area) if _SEASON_TOKEN.search(area) else f"{area} {season}"


def shift_label(label: str, months: int = 12) -> str:
    """'Apr-26' -> 'Apr-27'; anything else is left as it is."""
    try:
        d = _dt.datetime.strptime(label, "%b-%y")
    except ValueError:
        return label
    m = d.month - 1 + months
    return _dt.date(d.year + m // 12, m % 12 + 1, 1).strftime("%b-%y")


def month_slot(label: str, p: int) -> int:
    """Which Apr..Mar column (0..11) a period label is: 'Dec-26' -> 8. Falls back to position."""
    try:
        d = _dt.datetime.strptime(label, "%b-%y")
        return (d.month - 4) % 12
    except ValueError:
        return p % 12


def hub_curves(raw: dict) -> Dict[str, List[Optional[float]]]:
    """{hub code: monthly hub price} from a model's supply rows."""
    hp = np.asarray(raw["arrays"]["hub_price"], dtype=float)
    out: Dict[str, list] = {}
    for i, h in enumerate(raw["supply_hubs"]):
        if h and h not in out:
            out[h] = _tolist(hp[:, i])
    return out


def raw_from_data(d: model.ModelData) -> dict:
    """
    The raw inputs of a model that has already been built (a base plus a
    manager's edits, an uploaded workbook, ...), ready to save as a new base.
    Closed lanes keep their blank freight.
    """
    arrays = {k: np.asarray(getattr(d, k), dtype=float) for k in ARRAY_KEYS}
    arrays["trans_cost"] = d.freight_as_entered()
    return {
        "arrays": arrays, "cached_price": np.asarray(d.price, dtype=float),
        "supply_names": list(d.supply_names), "supply_groups": list(d.supply_groups),
        "supply_hubs": list(d.supply_hubs), "demand_names": list(d.demand_names),
        "period_labels": list(d.period_labels),
    }


def _pack(raw: dict, **extra) -> dict:
    out = {
        "arrays": {k: _tolist(raw["arrays"][k]) for k in ARRAY_KEYS},
        "cached_price": _tolist(raw["cached_price"]) if raw.get("cached_price") is not None else None,
        "supply_names": list(raw["supply_names"]), "supply_groups": list(raw["supply_groups"]),
        "supply_hubs": list(raw["supply_hubs"]), "demand_names": list(raw["demand_names"]),
        "period_labels": list(raw["period_labels"]),
    }
    out["hubs"] = hub_curves(raw)
    out.update(extra)
    return out


# --------------------------------------------------------------------------- #
# The "Demand for Models" workbook.
# --------------------------------------------------------------------------- #

class DemandFile:
    """
    Reads the yearly demand workbook ("Demand for Models - 26-27 ...xlsx"):

      Demand For Models   one row per plant: Lookup | | Model | Geography | Apr..Mar
                          (thousands of gallons, already state-adjusted)
      Demand Calculations district id, name, STATE, model; unadjusted and
                          adjusted demand
      State Adjustments   % of budget demand to contract for, by state and month

    Plants are matched to a model's demand points by name, then by district
    number (the "1001" in "1001-Springfield ST").
    """

    def __init__(self, path: str):
        import openpyxl
        self.path = path
        self.name = os.path.basename(path)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if "Demand For Models" not in wb.sheetnames:
                raise ValueError(f"{self.name} has no 'Demand For Models' sheet - is it the demand workbook?")
            self.rows: List[dict] = []
            for row in wb["Demand For Models"].iter_rows(min_row=2, values_only=True):
                if not row or not (row[0] or (len(row) > 3 and row[3])):
                    continue
                vals = []
                for v in row[4:16]:
                    vals.append(float(v) if isinstance(v, (int, float)) else 0.0)
                vals += [0.0] * (12 - len(vals))
                self.rows.append({"lookup": str(row[0] or "").strip(), "model": str(row[2] or "").strip(),
                                  "geo": str(row[3] or "").strip(), "months": vals})
            self.state_by_id: Dict[str, str] = {}
            self.state_by_name: Dict[str, str] = {}
            if "Demand Calculations" in wb.sheetnames:
                for row in wb["Demand Calculations"].iter_rows(min_row=4, max_col=4, values_only=True):
                    if not row or row[0] is None or not row[2]:
                        continue
                    did = str(row[0]).strip()
                    did = did.zfill(4) if did.isdigit() else did
                    st = str(row[2]).strip()
                    self.state_by_id[did] = st
                    if row[1]:
                        self.state_by_name[model.fold_name(str(row[1]))] = st
            self.state_pct: Dict[str, List[float]] = {}
            if "State Adjustments" in wb.sheetnames:
                for row in wb["State Adjustments"].iter_rows(min_row=2, max_col=15, values_only=True):
                    if not row or not row[2]:
                        continue
                    vals = [float(v) if isinstance(v, (int, float)) else None for v in row[3:15]]
                    if all(v is not None for v in vals):
                        self.state_pct[str(row[2]).strip()] = vals
        finally:
            wb.close()

    def _rows_for(self, area_code: str) -> List[dict]:
        key = model.fold_name(area_code)
        return [r for r in self.rows if model.fold_name(r["model"]) == key]

    def covers(self, area: str) -> bool:
        code = _SEASON_TOKEN.sub("", area).strip()
        return bool(self._rows_for(code))

    def state_of(self, plant: str) -> str:
        m = re.match(r"\s*(\d{3,4})\b", plant)
        if m and m.group(1).zfill(4) in self.state_by_id:
            return self.state_by_id[m.group(1).zfill(4)]
        if model.fold_name(plant) in self.state_by_name:
            return self.state_by_name[model.fold_name(plant)]
        t = re.findall(r"\b([A-Z]{2})\b", plant.upper())
        return t[-1] if t and t[-1] in self.state_pct else ""

    def match(self, area: str, demand_names: Sequence[str]) -> List[Optional[dict]]:
        """For each demand point, its row in the file (or None)."""
        code = _SEASON_TOKEN.sub("", area).strip()
        rows = self._rows_for(code)
        by_name: Dict[str, dict] = {}
        by_num: Dict[str, List[dict]] = {}
        for r in rows:
            for t in (r["lookup"], r["geo"]):
                if t:
                    by_name.setdefault(model.fold_name(t), r)
                    m = re.match(r"\s*(\d{3,4})\b", t)
                    if m:
                        by_num.setdefault(m.group(1).zfill(4), [])
                        if r not in by_num[m.group(1).zfill(4)]:
                            by_num[m.group(1).zfill(4)].append(r)
        out: List[Optional[dict]] = []
        used = set()
        for n in demand_names:
            hit = by_name.get(model.fold_name(n))
            if hit is None:
                m = re.match(r"\s*(\d{3,4})\b", n)
                cands = by_num.get(m.group(1).zfill(4), []) if m else []
                if len(cands) == 1:
                    hit = cands[0]
            if hit is not None and id(hit) in used:
                hit = None                         # never give one plant's demand to two rows
            if hit is not None:
                used.add(id(hit))
            out.append(hit)
        return out


def states_for(raw: dict, demand_file: Optional[DemandFile]) -> Tuple[List[str], Dict[str, List[float]]]:
    """The state of each demand point, and the % to contract for in each of those states."""
    names = raw["demand_names"]
    if demand_file is None:
        states = []
        for n in names:
            t = re.findall(r"\b([A-Z]{2})\b", n.upper())
            states.append(t[-1] if t and t[-1] in _US_STATES else "")
        return states, {}
    states = [demand_file.state_of(n) for n in names]
    pct = {s: demand_file.state_pct[s] for s in sorted(set(states)) if s in demand_file.state_pct}
    return states, pct


_US_STATES = set("AL AK AZ AR CA CO CT DE FL GA HI ID IL IN IA KS KY LA ME MD MA MI MN MS MO MT NE NV "
                 "NH NJ NM NY NC ND OH OK OR PA RI SC SD TN TX UT VT VA WA WV WI WY".split())


# --------------------------------------------------------------------------- #
# The library of bases.
# --------------------------------------------------------------------------- #

def _area_file(region: str, area: str) -> str:
    return re.sub(r"[^\w\-]+", "_", f"{region}__{area}").strip("_") + ".json"


class Bases:
    """All base models on disk, with an in-memory cache of the area files."""

    def __init__(self, folder: str):
        self.folder = folder
        os.makedirs(folder, exist_ok=True)
        self._lock = threading.RLock()
        self._cache: Dict[Tuple[str, str, str], dict] = {}
        self.jobs: Dict[str, dict] = {}

    # ---- reading ---------------------------------------------------------- #

    def list(self) -> List[dict]:
        out = []
        for d in sorted(os.listdir(self.folder)):
            p = os.path.join(self.folder, d, "meta.json")
            if os.path.isfile(p):
                try:
                    with open(p, encoding="utf-8") as fh:
                        out.append(json.load(fh))
                except (OSError, ValueError):
                    continue
        return sorted(out, key=lambda m: (m.get("season", ""), m.get("created", "")), reverse=True)

    def get(self, bid: str) -> Optional[dict]:
        if not re.fullmatch(r"[a-z0-9][a-z0-9\-]{0,80}", bid or ""):
            return None
        p = os.path.join(self.folder, bid, "meta.json")
        if not os.path.isfile(p):
            return None
        with open(p, encoding="utf-8") as fh:
            return json.load(fh)

    def default(self) -> Optional[dict]:
        bases = self.list()
        official = [b for b in bases if b.get("official")]
        return (official or bases or [None])[0]

    def area_raw(self, bid: str, region: str, area: str) -> Optional[dict]:
        key = (bid, region, area)
        with self._lock:
            if key in self._cache:
                return self._cache[key]
        meta = self.get(bid)
        if not meta:
            return None
        hit = next((a for a in meta.get("areas", []) if a["region"] == region and a["area"] == area), None)
        if not hit:
            return None
        with open(os.path.join(self.folder, bid, hit["file"]), encoding="utf-8") as fh:
            pkg = json.load(fh)
        # Once, not on every build: lists -> float arrays (null -> NaN = blank cell).
        pkg["arrays"] = {k: np.array(v, dtype=float) for k, v in pkg["arrays"].items()}
        if pkg.get("cached_price") is not None:
            pkg["cached_price"] = np.array(pkg["cached_price"], dtype=float)
        with self._lock:
            self._cache[key] = pkg
        return pkg

    # ---- writing (only ever a NEW base) ---------------------------------- #

    def _save(self, meta: dict, packages: List[Tuple[str, str, dict, dict]]) -> dict:
        """packages: (region, area, packed area json, per-area info). Writes, then locks."""
        bid = meta["id"]
        tmp = os.path.join(self.folder, f"_tmp_{bid}")
        os.makedirs(tmp, exist_ok=True)
        meta["areas"] = []
        for region, area, pkg, info in packages:
            fn = _area_file(region, area)
            with open(os.path.join(tmp, fn), "w", encoding="utf-8") as fh:
                json.dump(pkg, fh, separators=(",", ":"))
            meta["areas"].append({"region": region, "area": area, "file": fn, **info})
        meta["locked"] = True
        with open(os.path.join(tmp, "meta.json"), "w", encoding="utf-8") as fh:
            json.dump(meta, fh, indent=1)
        final = os.path.join(self.folder, bid)
        os.replace(tmp, final)
        for fn in os.listdir(final):                 # locked: read-only on disk too
            try:
                os.chmod(os.path.join(final, fn), stat.S_IREAD)
            except OSError:
                pass
        LOG.info("base %s saved (%d areas)", bid, len(packages))
        return meta

    def _new_meta(self, name: str, season: str, user: str, kind: str, recipe: dict) -> dict:
        slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:40] or "base"
        return {"id": f"{slug}-{uuid.uuid4().hex[:6]}", "name": name.strip()[:80], "season": season,
                "created": now_iso(), "created_by": user, "kind": kind, "recipe": recipe,
                "locked": True, "official": False, "notes": []}

    def import_workbooks(self, root: str, variant: str, user: str, name: str = "",
                         demand_file: Optional[DemandFile] = None, official: bool = False,
                         progress: Optional[Callable[[int, int, str], None]] = None) -> dict:
        """A base read straight from the area workbooks in the Model folder."""
        regions = model.discover(root, variant)
        areas = [(r, af) for r in sorted(regions) for af in regions[r]
                 if os.sep + "Test" + os.sep not in af.path]
        if not areas:
            raise ValueError(f"no {variant} workbooks found under {root}")
        seasons = sorted({season_of(af.area) for _, af in areas if season_of(af.area)})
        season = seasons[-1] if seasons else ""
        meta = self._new_meta(name or f"20{season} {variant} workbooks", season, user, "import",
                              {"from": f"{variant} workbooks in the Model folder", "root": root,
                               "variant": variant,
                               "demand_file": demand_file.name if demand_file else ""})
        meta["official"] = official
        packages = []
        for n, (region, af) in enumerate(areas, start=1):
            if progress:
                progress(n - 1, len(areas), af.area)
            reader = model.ModelBaseReader(af.path, af.area, region)
            raw = reader.raw_inputs(*reader.read_parts())
            states, pct = states_for(raw, demand_file)
            pkg = _pack(raw, region=region, area=af.area, source=os.path.basename(af.path),
                        demand_states=states, state_pct=pct)
            packages.append((region, af.area, pkg, {"source": os.path.basename(af.path)}))
        if progress:
            progress(len(areas), len(areas), "saving")
        return self._save(meta, packages)

    def roll_forward(self, sources: List[Tuple[str, str, model.ModelData]], season: str,
                     user: str, name: str, from_label: str,
                     demand_file: Optional[DemandFile] = None, freight_pct: float = 0.0,
                     progress: Optional[Callable[[int, int, str], None]] = None) -> dict:
        """
        Next year's base, the way the Excel baselines were made: the source
        model's contracts, differentials, hub prices, limits and storage, with
        the new year's demand (where the demand file has the plant) and,
        optionally, freight changed by *freight_pct* %. Every rate changes,
        $1.00/gal-plus lanes included (they are ordinary rates); lanes left
        blank stay blank (unused).
        """
        season = normalise_season(season)
        meta = self._new_meta(name or f"20{season} base", season, user, "roll-forward", {
            "from": from_label, "demand_file": demand_file.name if demand_file else "",
            "freight_pct": freight_pct})
        packages = []
        for n, (region, area, d) in enumerate(sources, start=1):
            if progress:
                progress(n - 1, len(sources), area)
            raw = raw_from_data(d)
            new_area = rename_area(area, season)
            # Months move on by however many years the season moves - none when a
            # working model (e.g. one filled in by pasting) is saved as a base
            # for its own year.
            old = season_of(area)
            years = ((int(season[:2]) - int(old[:2])) % 100) if old else 1
            years = years - 100 if years > 50 else years
            raw["period_labels"] = [shift_label(x, 12 * years) for x in raw["period_labels"]]
            info = {"source": f"{area} ({from_label})"}
            dem = np.asarray(raw["arrays"]["demand"], dtype=float)
            if demand_file is not None and demand_file.covers(area):
                hits = demand_file.match(area, raw["demand_names"])
                for j, h in enumerate(hits):
                    if h is None:
                        continue
                    for p, lab in enumerate(raw["period_labels"]):
                        dem[p, j] = h["months"][month_slot(lab, p)]
                info["demand_matched"] = sum(h is not None for h in hits)
                info["demand_kept"] = [raw["demand_names"][j] for j, h in enumerate(hits)
                                       if h is None and not model.fold_name(raw["demand_names"][j]).startswith("demand")]
            elif demand_file is not None:
                info["demand_note"] = "not in the demand file - last year's demand kept"
            raw["arrays"]["demand"] = dem
            if freight_pct:
                tc = np.asarray(raw["arrays"]["trans_cost"], dtype=float)
                raw["arrays"]["trans_cost"] = tc * (1 + freight_pct / 100.0)      # NaN (blank) stays NaN
            states, pct = states_for(raw, demand_file)
            pkg = _pack(raw, region=region, area=new_area, source=info["source"],
                        demand_states=states, state_pct=pct)
            packages.append((region, new_area, pkg, info))
        if progress:
            progress(len(sources), len(sources), "saving")
        return self._save(meta, packages)

    # ---- background jobs (imports take a while) ------------------------- #

    def start_job(self, label: str, fn: Callable[[Callable[[int, int, str], None]], dict]) -> str:
        jid = uuid.uuid4().hex[:10]
        job = {"id": jid, "label": label, "state": "running", "done": 0, "total": 0, "step": "",
               "base": None, "error": "", "started": now_iso()}
        self.jobs[jid] = job

        def progress(done, total, step):
            job.update(done=done, total=total, step=step)

        def run():
            try:
                meta = fn(progress)
                job.update(state="done", base=meta["id"])
            except Exception as exc:  # noqa: BLE001 - reported to the page
                LOG.exception("base job %s failed", label)
                job.update(state="error", error=str(exc))
                for d in os.listdir(self.folder):
                    if d.startswith("_tmp_"):
                        shutil.rmtree(os.path.join(self.folder, d), ignore_errors=True)
        threading.Thread(target=run, daemon=True, name=f"base-{jid}").start()
        return jid

    def running(self) -> List[dict]:
        return [j for j in self.jobs.values() if j["state"] == "running"]


def month_names() -> List[str]:
    return list(_MONTHS)
