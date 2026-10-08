#!/usr/bin/env python3
"""
model.py - Propane supply / transport optimization, rebuilt in Python.

WHY THIS EXISTS
---------------
Right now the optimization runs inside a family of .xlsm workbooks (one per
"area", e.g. CA-NV, AZ, MI, ...) that get solved with OpenSolver/CBC from
inside Excel. Excel is the bottleneck here: recalculating tens of thousands of
SUMPRODUCT cells and shuttling the model through the OpenSolver COM layer is
slow, and it's fragile in the way Excel automation always is.

This module reads the *same input data* straight out of those workbooks,
rebuilds the identical linear program as a sparse matrix, hands it to HiGHS to
solve, and writes the results out grouped BY REGION (folder) instead of one
file per area.

    Model/
      West/      AZ 26-27 Working.xlsm, CA-NV 26-27 Working.xlsm   -> West_Model_MM_DD_YY.xlsx
      Midwest/   IL-IN ..., MI ..., WI-IA-MN ...                   -> Midwest_Model_MM_DD_YY.xlsx
      ...

The source workbooks are only ever OPENED READ-ONLY. Nothing is ever written
back to them.

THE MODEL (transcribed from ModelBase)
--------------------------------------
Indices
    p = 1..P   periods (12 monthly buckets)
    i = 1..N   supply points
    j = 1..M   demand points

Decision variables
    ship[p,i,j]   volume moved supply i -> demand j in period p
    perc[p,i]     fraction of supply i's period-p contract volume actually lifted
    s_wd[p,i]     withdrawal from supply-point storage
    s_inj[p,i]    injection into supply-point storage
    s_end[p,i]    supply-point ending inventory
    d_wd[p,j]     withdrawal from demand-point storage
    d_inj[p,j]    injection into demand-point storage
    d_end[p,j]    demand-point ending inventory

Objective (Excel: ModelBase!D4 = SUM of per-period costs, minimize)
    min  sum_p [ sum_ij trans_cost[p,i,j] * ship[p,i,j]
               + sum_i  contract[p,i] * price[p,i] * perc[p,i] ]

    contract[] and price[] are just data here, not variables, so
    perc*contract*price stays linear even though it looks like a product of
    two unknowns. price[p,i] = hub_price[p,i] + hub_premium[p,i]
    (that's Excel col D = col R + col AF).

Constraints
    (A) supply balance, Excel col AX of each Ship block, forced = 0:
            contract[p,i]*perc[p,i] + s_wd[p,i]
              - ( sum_j ship[p,i,j] + s_inj[p,i] ) = 0

    (B) supply inventory chain, Excel col BB (end inv), bounded by BC/BD:
            s_end[p,i] = s_begin[p,i] - s_wd[p,i] + s_inj[p,i]
            s_begin[1,i] = supply initial inventory
            s_begin[p,i] = s_end[p-1,i]                for p > 1
            min_inv[p,i] <= s_end[p,i] <= supply storage capacity[p,i]

    (C) demand balance, Excel "Propane Demand Balance" block, forced = 0:
            sum_i ship[p,i,j] + d_wd[p,j] - ( demand[p,j] + d_inj[p,j] ) = 0

    (D) demand inventory chain, same shape as (B):
            d_end[p,j] = d_begin[p,j] - d_wd[p,j] + d_inj[p,j]
            min_inv[p,j] <= d_end[p,j] <= demand storage capacity[p,j]

Bounds
    min_perc[p,i] <= perc[p,i] <= max_perc[p,i]          (Excel D/R/AF % block)
    min_ship[p,i,j] <= ship[p,i,j] <= max_ship[p,i,j]    (Excel Min/Max Transport)
    all withdraw / inject >= 0

UNITS (per the workbooks' user documentation)
    Volumes are in THOUSANDS of gallons - 300,000 gal gets entered as 300.
    Prices and costs are $ per thousand gallons (mills per gallon), so a
    differential of MTB + $0.045 shows up as 45.0 in the model.
    Volume times price therefore lands in plain DOLLARS, which is how cost
    ends up getting reported.

WHY BLOCKS ARE LOCATED BY LABEL, NOT BY ARITHMETIC
--------------------------------------------------
Every workbook comes out of the same generator (MixedFormulator), but the row
offsets are NOT a clean function of N and M. For example, between the demand
table and the "Inventory at Demand Points" table, CA-NV has 3 spacer rows
while Alaska has 5 - close enough to tempt you into hardcoding an offset, and
just different enough to break it. So instead, anchors come from (a) the
workbook's own defined names (Period<p>_p_Contract / _Price / _HubPrice /
_HubPremium / _Perc / _Trans_Cost / _Ship) and (b) scanning for the block
captions in columns A-C. Both get verified on load, and a clear error is
raised the moment a sheet doesn't match what's expected.

REQUIREMENTS
------------
    python -m pip install openpyxl scipy numpy

    scipy is what provides HiGHS, via scipy.optimize.linprog. If scipy isn't
    available the module falls back to PuLP + CBC (pip install pulp) - the
    same solver the Excel version used, just driven directly instead of
    through OpenSolver.

USAGE
-----
    # every region found under Model/
    python model.py

    # what is available
    python model.py --list-regions
    python model.py --list-groups

Picking regions
    Managers generally only care about their own patch, so region selection
    is deliberately forgiving about how you type things. All of these are
    equivalent:

        python model.py --region Rail --region Rockies
        python model.py --region Rail,Rockies
        python model.py -r "Rail; Rockies"

    Matching ignores case, spaces, hyphens and underscores, and if that still
    doesn't land a hit it'll try a unique prefix or substring. So every one of
    these finds the Northeast folder: Northeast, northeast, "North East",
    north_east, "north east". Exact names always win over prefixes, so
    --region South picks South and never gets confused with Southeast.

        python model.py -r "North East"        # -> Northeast
        python model.py -r "new england"       # -> New England
        python model.py -r pnw                 # -> PNW
        python model.py -r rock                # -> Rockies  (unique prefix)

    Or skip regions instead of spelling out the ones you want:

        python model.py --exclude Rail
        python model.py -x Rail,Alaska

    Saved groups live in groups.json beside this script, so a recurring set
    of regions turns into one short command:

        python model.py --group jdoe
        python model.py --group east --region Rail

Other options
    # use the Baseline workbooks instead of the Working ones
    python model.py --variant Baseline

    # diagnose an infeasible area instead of just reporting "infeasible"
    python model.py --region South --allow-shortfall

    # inspect what would run, without solving
    python model.py --region West --list

    # point at a different tree / output folder
    python model.py --root "..\\Model" --out-dir "runs"

Output location
    Workbooks land in an ``Output`` folder next to this script by default -
    never into the folders holding the source workbooks.

Windows users
    ``Run Model.bat`` in this folder gives you a menu-driven version of
    everything above. It also takes a region directly on the command line,
    so a per-manager shortcut is just the batch file plus an argument:
    ``"Run Model.bat" Rail``
"""

from __future__ import annotations

import argparse
import datetime as _dt
import difflib
import logging
import math
import os
import re
import sys
import traceback
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence, Tuple

# --------------------------------------------------------------------------- #
# Third-party imports. If one is missing, tell the user how to fix it rather
# than letting them puzzle out a bare ImportError.
# --------------------------------------------------------------------------- #

try:
    import numpy as np
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "numpy is required.  Install the dependencies with:\n"
        "    python -m pip install openpyxl scipy numpy"
    ) from exc

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "openpyxl is required.  Install the dependencies with:\n"
        "    python -m pip install openpyxl scipy numpy"
    ) from exc


LOG = logging.getLogger("model")

# --------------------------------------------------------------------------- #
# Constants that describe the workbook template.
# --------------------------------------------------------------------------- #

MODEL_SHEET = "ModelBase"
N_PERIODS = 12

#: Captions used to find the data blocks. Matching is case-insensitive and
#: whitespace-normalised, because the workbooks are full of trailing spaces
#: someone left behind. Each entry is a tuple so alternative spellings can be
#: bolted on later without touching the lookup code itself.
CAPTION_SUPPLY_INVENTORY = ("inventory at supply points",)
CAPTION_DEMAND = ("demand",)
CAPTION_DEMAND_GEOGRAPHY = ("geography",)
CAPTION_DEMAND_INVENTORY = ("inventory at demand points",)
CAPTION_INITIAL_INVENTORY = ("initial inventory",)
CAPTION_STORAGE_CAPACITY = ("available storage capacity",)
CAPTION_TRANSPORT_COSTS = ("transport costs",)
CAPTION_MIN_TRANSPORT = ("min transport quantities", "minimum transport quantities")
CAPTION_MAX_TRANSPORT = ("max transport quantities", "maximum transport quantities")

#: Columns scanned when looking for block captions (A, B, C).
CAPTION_SEARCH_COLS = (1, 2, 3)

#: Captions on the row just above the first Ship block, which is where the
#: supply-side inventory columns live. Unlike the demand side these are single
#: columns, not 12-wide blocks: each period's Ship block occupies its own rows,
#: so one column covers every period.
CAPTION_SUPPLY_BALANCE = ("propane supply balance",)
CAPTION_SUPPLY_INIT_INV = ("initial inv", "initial inventory")
CAPTION_SUPPLY_END_INV = ("end inv", "ending inventory")
CAPTION_SUPPLY_MIN_INV = ("min inv", "min inventory", "minimum inv",
                          "minimum inventory")
CAPTION_SUPPLY_MAX_INV = ("max inv", "max inventory", "maximum inv",
                          "maximum inventory")

#: Captions on the "Propane Demand Balance" header row. These ARE 12-wide
#: blocks, with data starting one column right of the caption.
CAPTION_DEMAND_BALANCE = ("propane demand balance",)
CAPTION_DEMAND_MIN_INV = ("min inventory", "min inv", "minimum inventory")
CAPTION_DEMAND_MAX_INV = ("max inventory", "max inv", "maximum inventory")

#: Fallback used when a workbook doesn't spell out an explicit per-lane
#: maximum.
DEFAULT_MAX_SHIP = 999999.0

#: Freight at or above this ($/kgal, i.e. $1.00/gal) is "high freight": it gets
#: a WARNING, and nothing else. It is an ordinary price - the LP uses the lane
#: whenever it is the cheapest way to supply a plant, exactly as Excel did.
#: (Planners used to type 1000 to mean "do not use", but real freight on some
#: lanes is now above $1.00/gal, so a rate can no longer mean "unused".)
#: The ONLY way a lane is unused is a blank freight cell - see closed lanes.
HIGH_FREIGHT_RATE = 1000.0
HIGH_FREIGHT_TOL = 1e-6
# Older names, kept so existing code and decks still read.
BLOCKED_LANE_RATE = HIGH_FREIGHT_RATE
BLOCKED_LANE_TOL = HIGH_FREIGHT_TOL

#: Don't raise the high-freight warning below this volume. A simplex solution
#: can leave a few gallons on a lane as a rounding artifact, and warning about
#: "0.0 kgal" trains people to ignore the warning that matters.
HIGH_FREIGHT_REPORT_MIN = 0.1     # kgal, i.e. 100 gallons
BLOCKED_LANE_REPORT_MIN = HIGH_FREIGHT_REPORT_MIN

#: Defined-name pattern for the generator's optional second transport route,
#: and the markers for its optional fuel-blending build. Neither is modelled
#: here; both are refused outright, because a workbook built with them has a
#: whole second block of shipment variables that this reader would silently
#: ignore - producing a clean-looking but wrong answer.
ALT_ROUTE_NAME_RE = re.compile(r"Period\d+_p_Alt(Ship|Trans_Cost)", re.IGNORECASE)
MIXED_FUEL_SHEET = "output mix"
MIXED_FUEL_NAME_PREFIX = "m_"

#: Defined name holding the top-left cell of the annual cost block on the
#: 'Output Propane' sheet. The generator stores it as a quoted STRING (e.g.
#: '"$C$33"') rather than a real range reference, and without a sheet name.
SAVED_COST_NAME = "p_TopAnnualCommCost"
SAVED_COST_SHEET = "Output Propane"

#: Row offsets below the anchor, and the captions to confirm them by. Column C
#: holds the current/saved figure and column D the previous baseline.
SAVED_COST_ROWS = (
    ("commodity", 0, ("commodity",)),
    ("premium", 1, ("premium",)),
    ("transport", 2, ("transportation", "transport")),
    ("total", 3, ("total",)),
)
SAVED_COST_COL_CURRENT = 3      # column C
SAVED_COST_COL_BASELINE = 4     # column D

#: Cost charged against unmet demand when the shortfall relaxation is active.
#: It needs to be big enough that the solver never prefers leaving demand unmet
#: over making a real shipment, but not so big it wrecks the LP's numerics.
#: Real freight/commodity rates are in the hundreds to low thousands per kgal,
#: so 1e7 is about four orders of magnitude clear of anything legitimate.
SHORTFALL_PENALTY = 1.0e7

#: Cost charged against contracted volume that gets lifted but has nowhere to
#: go. This is the mirror image of a shortfall: it shows up when contract
#: MINIMUMS force more volume than demand can absorb and there is no storage
#: to park it in ("Min Too High" on the spreadsheet's own feasibility check).
SURPLUS_PENALTY = 1.0e7

#: Penalty on the artificial slacks used by the last-resort elastic solve.
#: Larger than the other two so the solver exhausts the meaningful escape
#: valves (unmet demand, surplus) before resorting to bending a balance
#: equation, which keeps the diagnosis pointed at the real culprit.
ELASTIC_PENALTY = 1.0e9

#: Constraint families, in the order equality_triplets() emits them. Used to
#: turn a bare row number back into something a human can act on.
FAM_SUPPLY_BALANCE = "supply balance"
FAM_SUPPLY_INVENTORY = "supply inventory"
FAM_DEMAND_BALANCE = "demand balance"
FAM_DEMAND_INVENTORY = "demand inventory"

#: Default name for the manual-entry deck, looked for beside this script.
DEFAULT_OVERRIDES_NAME = "overrides.txt"

#: Exit code used when a deck cannot be parsed or one of its lines resolves to
#: nothing. Distinct from 1/2/3/4 so the launcher can point the user at CHECK.
EXIT_BAD_OVERRIDES = 5

#: Folders under the model root that aren't regions - just housekeeping
#: directories we should skip over while scanning.
NON_REGION_DIRS = {"test", "output", "archive", "old", "_inspect", "__pycache__"}

#: Workbook name patterns that should never be treated as a model file.
SKIP_FILE_PATTERNS = (
    re.compile(r"^~\$"),          # Excel's own lock files
    re.compile(r"^copy of ", re.I),
    re.compile(r"^mixedformulator", re.I),
)

_VARIANT_RE = re.compile(
    r"""^(?P<stem>.*?)          # area name, e.g. "CA-NV 26-27"
        \s*
        (?:v(?P<vers>\d+)\s*)?  # optional "v2"
        (?P<variant>baseline|working)
        \s*(?P<seq>\d+)?        # optional trailing "1" / " 2"
        $""",
    re.IGNORECASE | re.VERBOSE,
)


# --------------------------------------------------------------------------- #
# Small helpers for dealing with Excel's A1-style cell references.
# --------------------------------------------------------------------------- #

_RANGE_RE = re.compile(
    r"^(?:'(?P<q>[^']+)'|(?P<u>[^'!]+))!"
    r"\$?(?P<c1>[A-Za-z]{1,3})\$?(?P<r1>\d+)"
    r"(?::\$?(?P<c2>[A-Za-z]{1,3})\$?(?P<r2>\d+))?$"
)


def col_to_index(letters: str) -> int:
    """'A' -> 1, 'AA' -> 27."""
    n = 0
    for ch in letters.upper():
        if not ("A" <= ch <= "Z"):
            raise ValueError(f"bad column reference {letters!r}")
        n = n * 26 + (ord(ch) - 64)
    return n


def index_to_col(idx: int) -> str:
    """1 -> 'A', 27 -> 'AA'."""
    if idx < 1:
        raise ValueError(f"column index must be >= 1, got {idx}")
    out = ""
    while idx:
        idx, rem = divmod(idx - 1, 26)
        out = chr(65 + rem) + out
    return out


@dataclass(frozen=True)
class CellRange:
    """A rectangular range on one sheet, 1-based and inclusive."""

    sheet: str
    row1: int
    col1: int
    row2: int
    col2: int

    @property
    def n_rows(self) -> int:
        return self.row2 - self.row1 + 1

    @property
    def n_cols(self) -> int:
        return self.col2 - self.col1 + 1

    def __str__(self) -> str:
        return (
            f"{self.sheet}!{index_to_col(self.col1)}{self.row1}"
            f":{index_to_col(self.col2)}{self.row2}"
        )


def parse_range(ref: str) -> CellRange:
    """
    Parse a defined-name target such as ``ModelBase!$R$7:$R$36``.

    A sheet-qualified single cell is fine too - it just becomes a 1x1 range.
    External references (``[1]ModelBase!...``) and multi-area references get
    rejected outright; those only ever turn up on the sheet-local copies of
    the names, which this module deliberately ignores in favour of the
    workbook-global ones.
    """
    text = (ref or "").strip()
    if text.startswith("="):
        text = text[1:].strip()
    if not text:
        raise ValueError("empty range reference")
    if "," in text or ";" in text:
        raise ValueError(f"multi-area reference not supported: {ref!r}")
    if text.startswith("[") or "]" in text:
        raise ValueError(f"external reference not supported: {ref!r}")

    m = _RANGE_RE.match(text)
    if not m:
        raise ValueError(f"unrecognised range reference: {ref!r}")

    sheet = m.group("q") or m.group("u")
    r1, c1 = int(m.group("r1")), col_to_index(m.group("c1"))
    if m.group("r2"):
        r2, c2 = int(m.group("r2")), col_to_index(m.group("c2"))
    else:
        r2, c2 = r1, c1
    return CellRange(sheet, min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2))


def normalise_caption(value: object) -> str:
    """Lower-case, collapse whitespace; '' for anything non-textual."""
    if value is None:
        return ""
    if not isinstance(value, str):
        return ""
    return re.sub(r"\s+", " ", value).strip().lower()


def as_float(value: object, default: float = 0.0) -> float:
    """
    Coerce a cell value to a float.

    Blank cells, text, Excel error strings and NaN all just collapse down to
    *default*. That's not an arbitrary choice - it mirrors how the spreadsheet
    itself behaves, since SUMPRODUCT treats text and blanks inside a numeric
    block as zero.
    """
    if value is None or value is True or value is False:
        return default
    if isinstance(value, (int, float)):
        f = float(value)
        return default if math.isnan(f) or math.isinf(f) else f
    if isinstance(value, str):
        text = value.strip().replace(",", "").replace("$", "")
        if not text or text.startswith("#"):
            return default
        if text.endswith("%"):
            try:
                return float(text[:-1]) / 100.0
            except ValueError:
                return default
        try:
            return float(text)
        except ValueError:
            return default
    return default


def as_text(value: object) -> str:
    """Cell value as a trimmed display string."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, (_dt.datetime, _dt.date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


# --------------------------------------------------------------------------- #
# A cached, random-access view of one worksheet.
# --------------------------------------------------------------------------- #


class SheetGrid:
    """
    The whole worksheet slurped into a dict keyed by ``(row, col)``.

    openpyxl's read-only mode streams rows top to bottom and doesn't support
    jumping around to arbitrary cells, but resolving this model means hopping
    all over the sheet constantly. Reading everything into a sparse dict once,
    up front, is both simpler to write and faster than re-scanning the sheet
    every time we need a value.
    """

    __slots__ = ("name", "cells", "max_row", "max_col")

    def __init__(self, name: str):
        self.name = name
        self.cells: Dict[Tuple[int, int], object] = {}
        self.max_row = 0
        self.max_col = 0

    @classmethod
    def from_worksheet(cls, ws) -> "SheetGrid":
        grid = cls(ws.title)
        cells = grid.cells
        max_row = max_col = 0
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if row is None:
                continue
            for c, value in enumerate(row, start=1):
                if value is None:
                    continue
                if isinstance(value, str):
                    if not value.strip():
                        continue
                cells[(r, c)] = value
                if c > max_col:
                    max_col = c
            if row and any(v is not None for v in row):
                max_row = r
        grid.max_row = max_row
        grid.max_col = max_col
        return grid

    def get(self, row: int, col: int) -> object:
        return self.cells.get((row, col))

    def num(self, row: int, col: int, default: float = 0.0) -> float:
        return as_float(self.cells.get((row, col)), default)

    def text(self, row: int, col: int) -> str:
        return as_text(self.cells.get((row, col)))

    def matrix(self, rng: CellRange, default: float = 0.0) -> np.ndarray:
        """Numeric contents of *rng* as a 2-D float array."""
        out = np.full((rng.n_rows, rng.n_cols), default, dtype=float)
        for dr in range(rng.n_rows):
            row = rng.row1 + dr
            for dc in range(rng.n_cols):
                v = self.cells.get((row, rng.col1 + dc))
                if v is not None:
                    out[dr, dc] = as_float(v, default)
        return out

    def column(self, row1: int, row2: int, col: int, default: float = 0.0) -> np.ndarray:
        return np.array(
            [as_float(self.cells.get((r, col)), default) for r in range(row1, row2 + 1)],
            dtype=float,
        )

    def find_caption(
        self,
        captions: Sequence[str],
        row_lo: int = 1,
        row_hi: Optional[int] = None,
        cols: Sequence[int] = CAPTION_SEARCH_COLS,
        exact: bool = True,
    ) -> Optional[Tuple[int, int]]:
        """
        Find the first ``(row, col)`` in the scan window whose text matches
        *captions*.

        Rows are scanned top-down, so whatever comes back is the topmost
        occurrence.
        """
        wanted = {c.lower() for c in captions}
        hi = self.max_row if row_hi is None else min(row_hi, self.max_row)
        for r in range(max(1, row_lo), hi + 1):
            for c in cols:
                got = normalise_caption(self.cells.get((r, c)))
                if not got:
                    continue
                hit = (got in wanted) if exact else any(w in got for w in wanted)
                if hit:
                    return (r, c)
        return None

    def find_caption_in_row(
        self, row: int, captions: Sequence[str], col_lo: int = 1, col_hi: Optional[int] = None
    ) -> Optional[int]:
        """Column of the first cell in *row* matching *captions*."""
        wanted = {c.lower() for c in captions}
        hi = self.max_col if col_hi is None else min(col_hi, self.max_col)
        for c in range(max(1, col_lo), hi + 1):
            if normalise_caption(self.cells.get((row, c))) in wanted:
                return c
        return None


# --------------------------------------------------------------------------- #
# Resolved geometry, plus all the data pulled out of one workbook.
# --------------------------------------------------------------------------- #


class ModelStructureError(RuntimeError):
    """Raised when a workbook doesn't match the expected ModelBase template."""


@dataclass
class ModelData:
    """Everything the LP needs from one area workbook."""

    area: str
    region: str
    source_path: str

    n_supply: int
    n_demand: int
    n_periods: int

    period_labels: List[str]
    supply_names: List[str]
    supply_groups: List[str]      # column B - the counterparty / supplier group
    supply_hubs: List[str]        # column D - market hub (MTB / CON / ...)
    demand_names: List[str]

    contract: np.ndarray          # (P, N) contracted volume
    price: np.ndarray             # (P, N) delivered price = hub + premium
    hub_price: np.ndarray         # (P, N)
    hub_premium: np.ndarray       # (P, N)
    perc_min: np.ndarray          # (P, N)
    perc_max: np.ndarray          # (P, N)

    demand: np.ndarray            # (P, M)

    trans_cost: np.ndarray        # (P, N, M)
    ship_min: np.ndarray          # (P, N, M)
    ship_max: np.ndarray          # (P, N, M)

    supply_inv_init: np.ndarray   # (N,)
    supply_inv_cap: np.ndarray    # (P, N)  Max Inv beside the Ship block
    supply_inv_min: np.ndarray    # (P, N)  Min Inv beside the Ship block
    demand_inv_init: np.ndarray   # (M,)
    demand_inv_cap: np.ndarray    # (P, M)
    demand_inv_min: np.ndarray    # (P, M)

    notes: List[str] = field(default_factory=list)

    #: Input contradictions we had to paper over to make the LP solvable at
    #: all (e.g. a minimum % entered above the maximum %). These are real
    #: data errors in the workbook, not warnings, and they get surfaced as
    #: Errors on the Diagnostics sheet so somebody goes and fixes the source.
    repairs: List[str] = field(default_factory=list)

    #: Supply->demand pairs with freight of $1.00/gal or more in at least one
    #: period (high freight - a warning, not a block).
    blocked_lane_pairs: int = 0

    #: (P, N, M) True where the freight cell was left blank: the lane is not
    #: used at all (the LP holds it at 0). None means no lane is closed.
    lane_closed: Optional[np.ndarray] = None

    #: Annual cost block already stored on the workbook's 'Output Propane'
    #: sheet, when --compare-saved asked for it. Keys: commodity, premium,
    #: transport, total, and the same four with a "_baseline" suffix.
    saved_cost: Optional[Dict[str, float]] = None

    #: One entry per override actually applied to this area, for the audit
    #: trail. Empty when no deck is in play.
    override_log: List["OverrideRecord"] = field(default_factory=list)

    # -- convenience ------------------------------------------------------- #

    @property
    def total_demand(self) -> float:
        return float(self.demand.sum())

    def high_freight_mask(self) -> np.ndarray:
        """(P, N, M) True where freight is $1.00/gal or more (closed lanes excluded)."""
        return (self.trans_cost >= HIGH_FREIGHT_RATE - HIGH_FREIGHT_TOL) & ~self.closed_lane_mask()

    # Older name.
    blocked_lane_mask = high_freight_mask

    def closed_lane_mask(self) -> np.ndarray:
        """(P, N, M) True where the lane is closed (blank freight)."""
        if self.lane_closed is None or self.lane_closed.shape != self.trans_cost.shape:
            return np.zeros(self.trans_cost.shape, dtype=bool)
        return self.lane_closed

    def freight_as_entered(self) -> np.ndarray:
        """Freight with closed lanes back as NaN (blank), for showing/comparing inputs."""
        return np.where(self.closed_lane_mask(), np.nan, self.trans_cost)

    def max_available_supply(self) -> np.ndarray:
        """Upper bound on liftable volume per period, ignoring storage."""
        return (self.contract * self.perc_max).sum(axis=1)

    def min_forced_supply(self) -> np.ndarray:
        """Volume that MUST be lifted per period because of contract minimums."""
        return (self.contract * self.perc_min).sum(axis=1)

    def feasibility_warnings(self) -> List[str]:
        """
        Reproduce the "Min Too High" / "Max Too Low" checks from ModelBase
        row 2.

        These are cheap, aggregate-volume sanity checks - they don't prove the
        LP is feasible (lane bounds and storage limits can still trip it up
        even when these pass), but they do catch the mistakes people actually
        tend to make when filling in the sheet.
        """
        out: List[str] = []
        per_period_demand = self.demand.sum(axis=1)
        lo = self.min_forced_supply()
        hi = self.max_available_supply()
        for p in range(self.n_periods):
            label = self.period_labels[p]
            if lo[p] > per_period_demand[p] + 1e-6:
                out.append(
                    f"{label}: contract minimums ({lo[p]:,.1f}) exceed demand "
                    f"({per_period_demand[p]:,.1f}) - 'Min Too High'"
                )
            if hi[p] + 1e-6 < per_period_demand[p]:
                out.append(
                    f"{label}: contract maximums ({hi[p]:,.1f}) cannot cover demand "
                    f"({per_period_demand[p]:,.1f}) - 'Max Too Low'"
                )
        return out


# --------------------------------------------------------------------------- #
# Manual entry: the override deck
#
# The area supply team needs to change inputs without touching the source
# workbooks (which stay strictly read-only) and without opening Excel at all.
# Per the SOP their real loop is: run wide open, then vary contract minimums
# and maximums, then vary differentials, then re-run and compare. The deck is
# shaped around exactly that.
#
# The single biggest risk in this whole feature is a silent no-op: a typo'd
# location name that quietly changes nothing means somebody runs a scenario,
# reads a number, and makes a purchasing decision on inputs they believe they
# changed and didn't. So every override MUST resolve to something, and failing
# to resolve is a hard error with the line number and the nearest candidates.
# --------------------------------------------------------------------------- #

#: Field name -> (array in the reader's `arrays` dict, how it is indexed).
#: "period" fields accept a monthly series; "scalar" fields are a single
#: opening balance that is not per-period.
_SUPPLY_FIELDS: Dict[str, Tuple[str, str]] = {
    "contract": ("contract", "period"),
    "diff": ("hub_premium", "period"),
    "hub_price": ("hub_price", "period"),
    "min_pct": ("perc_min", "period"),
    "max_pct": ("perc_max", "period"),
    "store_max": ("supply_inv_cap", "period"),
    "store_min": ("supply_inv_min", "period"),
    "store_init": ("supply_inv_init", "scalar"),
}
_DEMAND_FIELDS: Dict[str, Tuple[str, str]] = {
    "demand": ("demand", "period"),
    "store_max": ("demand_inv_cap", "period"),
    "store_min": ("demand_inv_min", "period"),
    "store_init": ("demand_inv_init", "scalar"),
}
_LANE_FIELDS: Dict[str, Tuple[str, str]] = {
    "freight": ("trans_cost", "period"),
    "min_ship": ("ship_min", "period"),
    "max_ship": ("ship_max", "period"),
}
_HUB_FIELDS: Dict[str, Tuple[str, str]] = {
    "price": ("hub_price", "period"),
}
#: Bare keys allowed directly under [all] or [area: X] - blanket sweeps across
#: every supply point, which is the common case when running wide open.
_BARE_FIELDS: Dict[str, Tuple[str, str]] = {
    "min_pct": ("perc_min", "period"),
    "max_pct": ("perc_max", "period"),
}

#: Overriding either of these makes the workbook's cached delivered price
#: stale, so it has to be rebuilt as hub_price + hub_premium.
_PRICE_COMPONENTS = ("hub_price", "hub_premium")

_TARGET_FIELDS = {
    "supply": _SUPPLY_FIELDS,
    "demand": _DEMAND_FIELDS,
    "lane": _LANE_FIELDS,
    "hub": _HUB_FIELDS,
}


def fold_name(text: str) -> str:
    """Fold a location / supplier / field name for loose comparison."""
    return re.sub(r"[^a-z0-9]+", "", (text or "").lower())


class OverrideError(RuntimeError):
    """
    A deck could not be parsed, or one of its lines resolved to nothing.

    Carries the deck path and line number so the message can point straight at
    what to fix.
    """

    def __init__(self, message: str, path: str = "", line_no: int = 0):
        self.path = path
        self.line_no = line_no
        if path and line_no:
            full = f"{os.path.basename(path)} line {line_no}: {message}"
        elif path:
            full = f"{os.path.basename(path)}: {message}"
        else:
            full = message
        super().__init__(full)


@dataclass
class OverrideRecord:
    """One applied override, for the audit trail."""

    area: str
    target: str            # human-readable, e.g. 'supply Terminal A, ST | Acme'
    field: str
    before: str
    after: str
    line_no: int
    scope: str = ""        # "[all]" or "[area: X]"


@dataclass
class OverrideRule:
    """One parsed instruction, not yet resolved against a workbook."""

    line_no: int
    scope: str                       # "all" or the area text from [area: X]
    target: str                      # supply | demand | lane | hub | ""
    selector: Tuple[str, ...]        # names, already unquoted
    field: str                       # canonical field name
    values: List[float]


@dataclass
class OverrideDeck:
    """A parsed deck: a scenario name plus an ordered list of rules."""

    path: str
    scenario: str = ""
    rules: List[OverrideRule] = field(default_factory=list)

    @property
    def area_scopes(self) -> List[str]:
        """The distinct [area: X] names the deck mentions."""
        seen: List[str] = []
        for r in self.rules:
            if r.scope != "all" and r.scope not in seen:
                seen.append(r.scope)
        return seen

    def rules_for(self, area: str) -> List[OverrideRule]:
        """
        Rules that apply to *area*, in precedence order.

        Workbook value -> [all] -> [area: X], last wins. Returning the [all]
        rules first and the area-specific ones after means simply applying the
        list in order gives exactly that precedence.
        """
        general = [r for r in self.rules if r.scope == "all"]
        specific = [r for r in self.rules
                    if r.scope != "all" and _area_matches(r.scope, area)]
        return general + specific

    def validate_area_scopes(self, all_areas: Sequence[str]) -> None:
        """
        Check every [area: X] section names an area that actually exists.

        A section naming an area that is simply not in this run's region
        selection is fine and silently inert; one naming an area that exists
        nowhere is a typo and must be caught.
        """
        for scope in self.area_scopes:
            if not any(_area_matches(scope, a) for a in all_areas):
                line = next((r.line_no for r in self.rules if r.scope == scope), 0)
                raise OverrideError(
                    f"section [area: {scope}] does not match any area in the model. "
                    + _suggest(scope, list(all_areas)),
                    self.path, line,
                )


def _area_matches(scope: str, area: str) -> bool:
    """
    Does a deck's [area: X] label refer to this area?

    Deck sections are written the way people speak - ``[area: AZ]`` for the
    workbook whose area is ``AZ 26-27`` - so a folded prefix match is what is
    wanted here.
    """
    s, a = fold_name(scope), fold_name(area)
    if not s or not a:
        return False
    return a == s or a.startswith(s) or s in a


def _suggest(needle: str, candidates: Sequence[str], n: int = 3) -> str:
    """'did you mean' text listing the closest candidates."""
    if not candidates:
        return ""
    folded = {}
    for c in candidates:
        folded.setdefault(fold_name(c), c)
    close = difflib.get_close_matches(fold_name(needle), list(folded), n=n, cutoff=0.0)
    if not close:
        return ""
    pretty = [folded[k] for k in close]
    pad = "\n" + " " * 17
    return "did you mean:  " + pad.join(pretty)


def _split_outside_quotes(text: str, sep: str) -> List[str]:
    """Split on *sep*, ignoring occurrences inside double quotes."""
    parts: List[str] = []
    buf: List[str] = []
    in_q = False
    i = 0
    while i < len(text):
        ch = text[i]
        if ch == '"':
            in_q = not in_q
            buf.append(ch)
            i += 1
            continue
        if not in_q and text.startswith(sep, i):
            parts.append("".join(buf))
            buf = []
            i += len(sep)
            continue
        buf.append(ch)
        i += 1
    parts.append("".join(buf))
    return parts


def _strip_comment(line: str) -> str:
    """Remove a trailing # comment, respecting quoted names."""
    out: List[str] = []
    in_q = False
    for ch in line:
        if ch == '"':
            in_q = not in_q
        if ch == "#" and not in_q:
            break
        out.append(ch)
    return "".join(out)


def _unquote(text: str) -> str:
    text = text.strip()
    if len(text) >= 2 and text[0] == '"' and text[-1] == '"':
        return text[1:-1].strip()
    return text


def _parse_values(raw: str, path: str, line_no: int, field: str) -> List[float]:
    """Parse a value list: separated by spaces and/or commas."""
    tokens = [t for t in re.split(r"[,\s]+", raw.strip()) if t]
    if not tokens:
        raise OverrideError(f"{field}: no value given", path, line_no)
    vals: List[float] = []
    for t in tokens:
        cleaned = t.rstrip("%").replace("$", "")
        try:
            vals.append(float(cleaned))
        except ValueError:
            raise OverrideError(
                f"{field}: {t!r} is not a number", path, line_no
            ) from None
    return vals


def _canonical_field(raw: str, allowed: Dict[str, Tuple[str, str]],
                     path: str, line_no: int, target: str) -> str:
    folded = fold_name(raw)
    for key in allowed:
        if fold_name(key) == folded:
            return key
    raise OverrideError(
        f"{target or 'this section'} has no field {raw.strip()!r}. "
        f"Valid fields here: {', '.join(sorted(allowed))}",
        path, line_no,
    )


def parse_override_deck(path: str) -> OverrideDeck:
    """
    Parse an override deck.

    Deliberately forgiving about layout (blank lines, comments, spacing, case)
    and deliberately strict about meaning (unknown fields, bad numbers and wrong
    series lengths are all errors, named with their line number).
    """
    if not os.path.isfile(path):
        raise OverrideError(f"override file not found: {path}")

    deck = OverrideDeck(path=path)
    scope = "all"          # anything before the first section behaves as [all]
    in_scenario = False

    with open(path, "r", encoding="utf-8-sig", errors="replace") as fh:
        for line_no, raw_line in enumerate(fh, start=1):
            line = _strip_comment(raw_line).strip()
            if not line:
                continue

            # ---- section header -------------------------------------------
            if line.startswith("[") and line.endswith("]"):
                inner = line[1:-1].strip()
                low = inner.lower()
                in_scenario = False
                if low == "scenario":
                    in_scenario = True
                    scope = "scenario"
                elif low == "all":
                    scope = "all"
                elif low.startswith("area"):
                    _, _, rest = inner.partition(":")
                    name = rest.strip()
                    if not name:
                        raise OverrideError(
                            "an [area: ...] section needs an area name",
                            path, line_no,
                        )
                    scope = name
                else:
                    raise OverrideError(
                        f"unknown section [{inner}]. Expected [scenario], [all] "
                        f"or [area: <name>]",
                        path, line_no,
                    )
                continue

            # ---- inside [scenario] ----------------------------------------
            if in_scenario:
                key, _, val = line.partition("=")
                if fold_name(key) == "name":
                    deck.scenario = val.strip()
                else:
                    raise OverrideError(
                        f"[scenario] only understands 'name = ...', got "
                        f"{key.strip()!r}",
                        path, line_no,
                    )
                continue

            # ---- targeted rule:  <kw> <selector> : <field> = <value> ------
            head, sep, tail = _partition_outside_quotes(line, ":")
            if sep:
                words = head.strip().split(None, 1)
                keyword = words[0].lower() if words else ""
                if keyword in _TARGET_FIELDS:
                    selector_text = words[1] if len(words) > 1 else ""
                    fld_raw, eq, val_raw = tail.partition("=")
                    if not eq:
                        raise OverrideError(
                            "expected '<field> = <value>' after the colon",
                            path, line_no,
                        )
                    fields = _TARGET_FIELDS[keyword]
                    fld = _canonical_field(fld_raw, fields, path, line_no, keyword)
                    selector = _parse_selector(
                        keyword, selector_text, path, line_no
                    )
                    deck.rules.append(OverrideRule(
                        line_no=line_no, scope=scope, target=keyword,
                        selector=selector, field=fld,
                        values=_parse_values(val_raw, path, line_no, fld),
                    ))
                    continue
                # A colon but not a known keyword - fall through so the bare
                # key branch can report something helpful.

            # ---- bare key (blanket sweep) ---------------------------------
            key, eq, val = line.partition("=")
            if not eq:
                raise OverrideError(
                    f"cannot make sense of {line!r}. Expected either "
                    f"'<field> = <value>' or "
                    f"'supply|demand|lane|hub <name> : <field> = <value>'",
                    path, line_no,
                )
            if scope == "scenario":
                raise OverrideError(
                    "values must sit under [all] or [area: <name>]",
                    path, line_no,
                )
            fld = _canonical_field(key, _BARE_FIELDS, path, line_no,
                                   "a blanket sweep")
            deck.rules.append(OverrideRule(
                line_no=line_no, scope=scope, target="", selector=(),
                field=fld, values=_parse_values(val, path, line_no, fld),
            ))

    return deck


def _partition_outside_quotes(text: str, sep: str) -> Tuple[str, str, str]:
    """Like str.partition but ignores separators inside double quotes."""
    in_q = False
    for i, ch in enumerate(text):
        if ch == '"':
            in_q = not in_q
            continue
        if not in_q and text.startswith(sep, i):
            return text[:i], sep, text[i + len(sep):]
    return text, "", ""


def _parse_selector(keyword: str, text: str, path: str,
                    line_no: int) -> Tuple[str, ...]:
    """Pull the names out of a rule's selector."""
    text = text.strip()
    if keyword == "lane":
        parts = _split_outside_quotes(text, "->")
        if len(parts) != 2:
            raise OverrideError(
                'a lane needs two locations: lane "From" -> "To" : field = value',
                path, line_no,
            )
        frm, to = _unquote(parts[0]), _unquote(parts[1])
        if not frm or not to:
            raise OverrideError("a lane needs both a from and a to location",
                                path, line_no)
        return (frm, to)

    if keyword == "supply":
        parts = _split_outside_quotes(text, "|")
        name = _unquote(parts[0])
        supplier = _unquote(parts[1]) if len(parts) > 1 else ""
        if not name:
            raise OverrideError("supply needs a location name", path, line_no)
        return (name, supplier)

    name = _unquote(text)
    if not name:
        raise OverrideError(f"{keyword} needs a name", path, line_no)
    return (name,)


def _resolve_one(needle: str, candidates: Sequence[str]) -> List[int]:
    """
    Indices of *candidates* matching *needle*, most specific match wins.

    Exact folded match first, then unique prefix, then substring - the same
    ladder the region matcher uses, so "terminal a st" finds "Terminal A, ST".
    """
    key = fold_name(needle)
    if not key:
        return []
    folded = [fold_name(c) for c in candidates]

    exact = [i for i, f in enumerate(folded) if f == key]
    if exact:
        return exact
    pref = [i for i, f in enumerate(folded) if f.startswith(key)]
    if pref:
        return pref
    return [i for i, f in enumerate(folded) if key in f]


def _supply_label(i: int, names: Sequence[str], groups: Sequence[str]) -> str:
    group = groups[i] if i < len(groups) and groups[i] else ""
    return f"{names[i]} | {group}" if group else names[i]


def _resolve_supply(rule: OverrideRule, area: str, names: Sequence[str],
                    groups: Sequence[str], path: str,
                    allow_many: bool = False) -> List[int]:
    """
    Resolve a supply selector to a row, or raise.

    By default an ambiguous location name is an error, because a ``supply``
    rule changes commercial terms and guessing which of three suppliers was
    meant is not acceptable - the ``| <supplier>`` qualifier exists for that.

    ``allow_many`` is used for lane rules, where the name identifies one end of
    a ROUTE rather than a contract. "Terminal A, ST -> Springfield" naturally means
    every Terminal A lane into Springfield, and each one resolved is logged
    separately so the audit trail stays exact.
    """
    wanted_name, wanted_supplier = rule.selector[0], rule.selector[1]
    hits = _resolve_one(wanted_name, names)

    labels = [_supply_label(i, names, groups) for i in range(len(names))]
    if not hits:
        raise OverrideError(
            f"no supply point in area {area} matches {wanted_name!r}\n  "
            + _suggest(wanted_name, labels),
            path, rule.line_no,
        )

    if wanted_supplier:
        narrowed = [i for i in hits
                    if fold_name(groups[i]) == fold_name(wanted_supplier)]
        if not narrowed:
            narrowed = [i for i in hits
                        if fold_name(wanted_supplier) in fold_name(groups[i])]
        if not narrowed:
            raise OverrideError(
                f"supply point {wanted_name!r} exists in area {area} but not "
                f"with supplier {wanted_supplier!r}\n  "
                + _suggest(wanted_supplier, [labels[i] for i in hits]),
                path, rule.line_no,
            )
        hits = narrowed

    if len(hits) > 1 and not allow_many:
        options = "\n" + "\n".join(f"                 {labels[i]}" for i in hits)
        raise OverrideError(
            f"{wanted_name!r} matches {len(hits)} supply points in area {area}. "
            f"Add a supplier to say which, e.g. "
            f'supply "{names[hits[0]]}" | {groups[hits[0]] or "<supplier>"}'
            f"{options}",
            path, rule.line_no,
        )
    return hits


def _resolve_simple(rule: OverrideRule, needle: str, area: str,
                    candidates: Sequence[str], what: str,
                    path: str) -> List[int]:
    """Resolve a demand point or hub name to exactly one entry, or raise."""
    hits = _resolve_one(needle, candidates)
    if not hits:
        raise OverrideError(
            f"no {what} in area {area} matches {needle!r}\n  "
            + _suggest(needle, list(dict.fromkeys(candidates))),
            path, rule.line_no,
        )
    if len(hits) > 1:
        options = "\n" + "\n".join(
            f"                 {candidates[i]}" for i in hits
        )
        raise OverrideError(
            f"{needle!r} matches {len(hits)} {what}s in area {area}; "
            f"be more specific{options}",
            path, rule.line_no,
        )
    return hits


def _fmt_series(values: np.ndarray) -> str:
    """Render a before/after value for the audit sheet."""
    arr = np.atleast_1d(np.asarray(values, dtype=float))
    if arr.size == 0:
        return ""
    if np.allclose(arr, arr.flat[0], rtol=0.0, atol=1e-9):
        v = float(arr.flat[0])
        return f"{v:g}"
    return " ".join(f"{float(v):g}" for v in arr.ravel())


def _expand(rule: OverrideRule, kind: str, n_periods: int,
            path: str) -> List[float]:
    """Broadcast or validate a value list against the period count."""
    vals = rule.values
    if kind == "scalar":
        if len(vals) != 1:
            raise OverrideError(
                f"{rule.field} is a single opening balance, not a monthly "
                f"series, but {len(vals)} values were given",
                path, rule.line_no,
            )
        return vals
    if len(vals) == 1:
        return vals * n_periods
    if len(vals) == n_periods:
        return vals
    raise OverrideError(
        f"{rule.field} needs either 1 value (applied to every month) or "
        f"exactly {n_periods}, but {len(vals)} were given",
        path, rule.line_no,
    )


def apply_overrides(
    deck: OverrideDeck,
    area: str,
    arrays: Dict[str, np.ndarray],
    supply_names: Sequence[str],
    supply_groups: Sequence[str],
    supply_hubs: Sequence[str],
    demand_names: Sequence[str],
    period_labels: Sequence[str],
    price_touched: np.ndarray,
) -> List[OverrideRecord]:
    """
    Apply the deck's rules for one area, in place, returning an audit record
    for every change actually made.

    Raises :class:`OverrideError` the moment a rule fails to resolve. That
    strictness is the whole point: a silent no-op here would let somebody make
    a purchasing decision on inputs they only think they changed.
    """
    path = deck.path
    n_periods = len(period_labels)
    log: List[OverrideRecord] = []

    def record(rule: OverrideRule, target: str, before, after) -> None:
        log.append(OverrideRecord(
            area=area, target=target, field=rule.field,
            before=_fmt_series(before), after=_fmt_series(after),
            line_no=rule.line_no,
            scope="[all]" if rule.scope == "all" else f"[area: {rule.scope}]",
        ))

    for rule in deck.rules_for(area):
        # ---- blanket sweep over every supply point --------------------
        if not rule.target:
            arr_key, kind = _BARE_FIELDS[rule.field]
            arr = arrays[arr_key]
            series = np.asarray(_expand(rule, kind, n_periods, path))
            for i in range(len(supply_names)):
                before = arr[:, i].copy()
                arr[:, i] = series
                # Every touched point is recorded, including ones that were
                # already at the requested value: the audit trail documents
                # what the deck asserted, not merely what happened to move.
                record(rule,
                       f"supply {_supply_label(i, supply_names, supply_groups)}",
                       before, arr[:, i])
                if arr_key in _PRICE_COMPONENTS:
                    price_touched[:, i] = True
            continue

        fields = _TARGET_FIELDS[rule.target]
        arr_key, kind = fields[rule.field]
        arr = arrays[arr_key]
        series = np.asarray(_expand(rule, kind, n_periods, path))

        if rule.target == "supply":
            idx = _resolve_supply(rule, area, supply_names, supply_groups, path)
            for i in idx:
                label = f"supply {_supply_label(i, supply_names, supply_groups)}"
                if kind == "scalar":
                    before = np.asarray([arr[i]])
                    arr[i] = series[0]
                    after = np.asarray([arr[i]])
                else:
                    before = arr[:, i].copy()
                    arr[:, i] = series
                    after = arr[:, i]
                    if arr_key in _PRICE_COMPONENTS:
                        price_touched[:, i] = True
                record(rule, label, before, after)

        elif rule.target == "demand":
            idx = _resolve_simple(rule, rule.selector[0], area, demand_names,
                                  "demand point", path)
            for j in idx:
                label = f"demand {demand_names[j]}"
                if kind == "scalar":
                    before = np.asarray([arr[j]])
                    arr[j] = series[0]
                    after = np.asarray([arr[j]])
                else:
                    before = arr[:, j].copy()
                    arr[:, j] = series
                    after = arr[:, j]
                record(rule, label, before, after)

        elif rule.target == "lane":
            frm, to = rule.selector
            # A lane name may carry a supplier qualifier too, e.g.
            #   lane "Terminal A, ST | Acme" -> "1001-Springfield ST"
            frm_parts = _split_outside_quotes(frm, "|")
            si = _resolve_supply(
                OverrideRule(rule.line_no, rule.scope, "supply",
                             (_unquote(frm_parts[0]),
                              _unquote(frm_parts[1]) if len(frm_parts) > 1 else ""),
                             rule.field, rule.values),
                area, supply_names, supply_groups, path, allow_many=True,
            )
            dj = _resolve_simple(rule, to, area, demand_names,
                                 "demand point", path)
            for i in si:
                for j in dj:
                    before = arr[:, i, j].copy()
                    arr[:, i, j] = series
                    # Qualify the supply end, otherwise three same-location rows
                    # produce three identical-looking audit lines.
                    record(rule,
                           f"lane {_supply_label(i, supply_names, supply_groups)}"
                           f" -> {demand_names[j]}",
                           before, arr[:, i, j])

        elif rule.target == "hub":
            hub = rule.selector[0]
            idx = [i for i, h in enumerate(supply_hubs)
                   if fold_name(h) == fold_name(hub)]
            if not idx:
                idx = [i for i, h in enumerate(supply_hubs)
                       if fold_name(hub) and fold_name(hub) in fold_name(h)]
            if not idx:
                hubs_present = sorted({h for h in supply_hubs if h})
                raise OverrideError(
                    f"no supply point in area {area} uses hub {hub!r}\n  "
                    + _suggest(hub, hubs_present),
                    path, rule.line_no,
                )
            for i in idx:
                before = arr[:, i].copy()
                arr[:, i] = series
                if arr_key in _PRICE_COMPONENTS:
                    price_touched[:, i] = True
                record(rule,
                       f"hub {hub} -> {_supply_label(i, supply_names, supply_groups)}",
                       before, arr[:, i])

    return log


# --------------------------------------------------------------------------- #
# Workbook reader
# --------------------------------------------------------------------------- #


#: Inputs the web app may edit, and what each index position means. The key
#: is the name in ModelBaseReader._build()'s ``arrays`` dict; the axes are the
#: array's dimensions in order (p = period, i = supply point, j = demand point).
CELL_EDIT_FIELDS: Dict[str, Tuple[str, ...]] = {
    "contract": ("p", "i"),
    "hub_price": ("p", "i"),
    "hub_premium": ("p", "i"),
    "perc_min": ("p", "i"),
    "perc_max": ("p", "i"),
    "demand": ("p", "j"),
    "trans_cost": ("p", "i", "j"),
    "ship_min": ("p", "i", "j"),
    "ship_max": ("p", "i", "j"),
    "supply_inv_init": ("i",),
    "supply_inv_cap": ("p", "i"),
    "supply_inv_min": ("p", "i"),
    "demand_inv_init": ("j",),
    "demand_inv_cap": ("p", "j"),
    "demand_inv_min": ("p", "j"),
}

#: Inputs where an edit may be blank (value None) - a blank freight rate
#: closes the lane, exactly like a blank cell in the workbook.
BLANKABLE_EDIT_FIELDS = frozenset({"trans_cost"})


def apply_cell_edits(
    edits: Sequence[Dict[str, object]],
    area: str,
    arrays: Dict[str, np.ndarray],
    supply_names: Sequence[str],
    supply_groups: Sequence[str],
    demand_names: Sequence[str],
    period_labels: Sequence[str],
    price_touched: np.ndarray,
) -> List[OverrideRecord]:
    """
    Apply position-based edits (from the web app) in place, and audit them.

    Each edit is ``{"field": <CELL_EDIT_FIELDS key>, "index": [ints],
    "value": float}`` plus optional ``"user"``. Positions rather than names
    are used deliberately: several areas have two supply rows with the same
    location name, so a name can be ambiguous where a row number never is.

    Anything that doesn't fit - unknown field, wrong number of indices, out of
    range, non-numeric value - raises OverrideError. Like the deck, an edit
    must never be silently dropped.
    """
    log: List[OverrideRecord] = []
    for n, e in enumerate(edits, start=1):
        fld = str(e.get("field", ""))
        axes = CELL_EDIT_FIELDS.get(fld)
        if axes is None:
            raise OverrideError(f"edit {n}: {fld!r} is not an editable input")
        arr = arrays[fld]
        idx = e.get("index")
        if not isinstance(idx, (list, tuple)) or len(idx) != len(axes):
            raise OverrideError(f"edit {n} ({fld}): expected {len(axes)} index "
                                f"value(s) ({', '.join(axes)}), got {idx!r}")
        raw = e.get("value")
        # A blank freight rate is a value in its own right: "lane not used".
        blank = fld in BLANKABLE_EDIT_FIELDS and (raw is None or raw == "")
        try:
            pos = tuple(int(k) for k in idx)
            value = float("nan") if blank else float(raw)
        except (TypeError, ValueError):
            raise OverrideError(f"edit {n} ({fld}): index and value must be numbers")
        if not blank and not math.isfinite(value):
            raise OverrideError(f"edit {n} ({fld}): value must be a finite number")
        if any(k < 0 or k >= s for k, s in zip(pos, arr.shape)):
            raise OverrideError(f"edit {n} ({fld}): position {list(pos)} is outside "
                                f"this area's {list(arr.shape)} table")

        before = float(arr[pos])
        arr[pos] = value
        if fld in _PRICE_COMPONENTS:
            price_touched[pos] = True

        where = dict(zip(axes, pos))
        parts = []
        if "i" in where and "j" in where:
            parts.append(f"lane {_supply_label(where['i'], supply_names, supply_groups)}"
                         f" -> {demand_names[where['j']]}")
        elif "i" in where:
            parts.append(f"supply {_supply_label(where['i'], supply_names, supply_groups)}")
        elif "j" in where:
            parts.append(f"demand {demand_names[where['j']]}")
        if "p" in where:
            parts.append(f"[{period_labels[where['p']]}]")
        who = str(e.get("user") or "").strip()
        log.append(OverrideRecord(
            area=area, target=" ".join(parts), field=fld,
            before="blank" if math.isnan(before) else _fmt_series(np.asarray([before])),
            after="blank (lane not used)" if blank else _fmt_series(np.asarray([value])),
            line_no=n, scope=f"[web{': ' + who if who else ''}]",
        ))
    return log


#: Suffix added to the supplier of every supply point created in the model
#: builder, so it is obvious on every sheet which rows did not come from the
#: workbook.
BUILDER_TAG = "(new)"

#: Arrays whose first axis is the period. Used when the builder changes the
#: number of time periods.
_PERIOD_ARRAYS = ("contract", "hub_price", "hub_premium", "perc_min", "perc_max", "demand",
                  "trans_cost", "ship_min", "ship_max", "supply_inv_cap", "supply_inv_min",
                  "demand_inv_cap", "demand_inv_min")
#: When periods are ADDED, these start empty (nothing contracted, nothing
#: demanded, nothing forced onto a lane); every other array repeats its last
#: month so prices, freight, limits and storage carry forward sensibly.
_PERIOD_START_EMPTY = ("contract", "demand", "ship_min")

#: How the builder names the price hubs a new supply point can buy at.
#: Workbooks carry one code per hub (column D); Mont Belvieu is split into
#: TET and non-TET product. Codes not listed here are offered as they are.
HUB_VARIANTS: Dict[str, Tuple[str, ...]] = {"MTB": ("MTB-Tet", "MTB-Non-T")}
HUB_DESCRIPTIONS: Dict[str, str] = {
    "CON": "Conway", "MTB": "Mont Belvieu",
    "MTB-Tet": "Mont Belvieu, TET", "MTB-Non-T": "Mont Belvieu, non-TET",
}


def builder_hub_choices(workbook_hubs: Sequence[str]) -> List[Dict[str, str]]:
    """
    The price-hub choices for a new supply point, in order: each workbook hub
    code, with MTB replaced by MTB-Tet and MTB-Non-T. ``value`` is what gets
    saved; ``label`` is what the dropdown shows ("CON - Conway").
    """
    out: List[Dict[str, str]] = []
    for h in workbook_hubs:
        for v in HUB_VARIANTS.get(h, (h,)):
            desc = HUB_DESCRIPTIONS.get(v, "")
            out.append({"value": v, "label": f"{v} - {desc}" if desc else v})
    return out


#: Largest sizes the builder accepts, to keep a slip of the slider from
#: building something absurd.
BUILDER_MAX_PERIODS = 36
BUILDER_MAX_ADDED = 60


def _next_label(label: str, step: int) -> str:
    """'Mar-27' + 1 -> 'Apr-27'. Anything unparseable becomes 'P13', 'P14', ..."""
    try:
        d = _dt.datetime.strptime(label, "%b-%y")
    except ValueError:
        return ""
    m = d.month - 1 + step
    return _dt.date(d.year + m // 12, m % 12 + 1, 1).strftime("%b-%y")


def extend_structure(
    structure: Dict[str, object],
    area: str,
    arrays: Dict[str, np.ndarray],
    supply_names: List[str],
    supply_groups: List[str],
    supply_hubs: List[str],
    demand_names: List[str],
    period_labels: List[str],
    cached_price: np.ndarray,
    price_touched: np.ndarray,
) -> Tuple[List[OverrideRecord], np.ndarray, np.ndarray]:
    """
    The web app's model builder: resize an area's model without rebuilding it.

    This is the MixedFormulator "Build Optimization Model" panel (Time Periods,
    Supply Points, Demand Points, Price Hubs) done on the fly. In Excel a
    different size means generating a new workbook; here it just means wider
    or longer arrays before the LP is built. *structure* is::

        periods   int                  1..36; fewer truncates, more repeats the
                                       last month's prices/freight/limits and
                                       starts volumes and demand at zero
        hubs      [{name, price}]      extra price hubs; price is 1 or P values
        supply    [{location, supplier, hub}]
                                       extra supply points (new contracts). They
                                       start with no volume and every freight
                                       cell blank (lane not used), so they change
                                       nothing until someone fills them in
        demand    [{name}]             extra plants, starting at zero demand

    Rows are only ever APPENDED after the workbook's own, so every saved cell
    edit keeps pointing at the same cell. Workbook rows can't be removed here.

    The supply-/period-shaped entries of *arrays* are replaced by resized
    copies, and the name/label lists are changed in place. Returns
    (audit records, resized cached_price, resized price_touched).
    """
    log: List[OverrideRecord] = []
    P0, N0 = arrays["contract"].shape

    def record(target, after):
        log.append(OverrideRecord(area=area, target=target, field="model builder",
                                  before="-", after=after, line_no=len(log) + 1, scope="[web]"))

    # ---- time periods ------------------------------------------------------ #
    P = int(structure.get("periods") or P0)
    if not 1 <= P <= BUILDER_MAX_PERIODS:
        raise OverrideError(f"time periods must be between 1 and {BUILDER_MAX_PERIODS}, got {P}")
    if P < P0:
        for k in _PERIOD_ARRAYS:
            arrays[k] = arrays[k][:P].copy()
        cached_price, price_touched = cached_price[:P].copy(), price_touched[:P].copy()
        dropped = period_labels[P:]
        del period_labels[P:]
        record("time periods", f"{P0} -> {P} (dropped {', '.join(dropped)})")
    elif P > P0:
        extra = P - P0
        for k in _PERIOD_ARRAYS:
            a = arrays[k]
            tail = np.zeros((extra,) + a.shape[1:]) if k in _PERIOD_START_EMPTY \
                else np.repeat(a[-1:], extra, axis=0)
            arrays[k] = np.concatenate([a, tail], axis=0)
        cached_price = np.concatenate([cached_price, np.full((extra, N0), np.nan)], axis=0)
        price_touched = np.concatenate([price_touched, np.ones((extra, N0), dtype=bool)], axis=0)
        last = period_labels[-1] if period_labels else ""
        for s in range(1, extra + 1):
            period_labels.append(_next_label(last, s) or f"P{P0 + s}")
        record("time periods", f"{P0} -> {P} (added {period_labels[P0]} to {period_labels[-1]}; "
                               f"volumes and demand start at 0, prices and freight repeat {last})")

    # ---- price hubs -------------------------------------------------------- #
    hub_price: Dict[str, np.ndarray] = {}
    for i in range(N0):
        h = supply_hubs[i]
        if h and fold_name(h) not in hub_price:
            hub_price[fold_name(h)] = arrays["hub_price"][:, i].copy()
    hub_label = {fold_name(h): h for h in supply_hubs[:N0] if h}
    new_hubs = list(structure.get("hubs") or [])
    if len(new_hubs) > BUILDER_MAX_ADDED:
        raise OverrideError(f"at most {BUILDER_MAX_ADDED} new price hubs")
    for n, h in enumerate(new_hubs, start=1):
        name = str(h.get("name") or f"Hub {len(hub_label) + 1}").strip()
        if fold_name(name) in hub_label:
            raise OverrideError(f"new price hub {n}: {name!r} already exists in {area}")
        raw = h.get("price")
        vals = list(raw) if isinstance(raw, (list, tuple)) else [raw if raw not in (None, "") else 0.0]
        try:
            vals = [float(v) if v not in (None, "") else 0.0 for v in vals]
        except (TypeError, ValueError):
            raise OverrideError(f"new price hub {name}: prices must be numbers")
        if len(vals) == 1:
            vals = vals * P
        if len(vals) < P:                       # periods were added after the prices were typed
            vals = vals + [vals[-1]] * (P - len(vals))
        vals = vals[:P]
        if not all(math.isfinite(v) for v in vals):
            raise OverrideError(f"new price hub {name}: prices must be finite")
        hub_price[fold_name(name)] = np.asarray(vals)
        hub_label[fold_name(name)] = name
        record(f"hub {name}", f"new price hub, average ${np.mean(vals):,.2f}/kgal")

    # Hub names the builder offers beyond the workbook's own codes - Mont
    # Belvieu as TET and non-TET. They start on the workbook hub's price
    # (planners set the contract's differential as usual); a new price hub
    # created above with the same name and its own prices takes precedence.
    for base, variants in HUB_VARIANTS.items():
        if fold_name(base) in hub_price:
            for v in variants:
                if fold_name(v) not in hub_price:
                    hub_price[fold_name(v)] = hub_price[fold_name(base)].copy()
                    hub_label[fold_name(v)] = v

    M0 = arrays["demand"].shape[1]

    # ---- supply points (new contracts) ------------------------------------ #
    new_sup = list(structure.get("supply") or [])
    if len(new_sup) > BUILDER_MAX_ADDED:
        raise OverrideError(f"at most {BUILDER_MAX_ADDED} new supply points")
    if new_sup:
        default_hub = next(iter(hub_label), "")
        K = len(new_sup)
        hp = np.zeros((P, K))
        for k, s in enumerate(new_sup):
            loc = str(s.get("location") or "").strip() or f"New supply point {k + 1}"
            sup = str(s.get("supplier") or "").strip()
            hub = fold_name(str(s.get("hub") or "")) or default_hub
            if hub and hub not in hub_price:
                raise OverrideError(f"new supply point {loc}: price hub {s.get('hub')!r} "
                                    f"isn't one of this area's hubs")
            if hub:
                hp[:, k] = hub_price[hub]
            supply_names.append(loc)
            supply_groups.append(f"{sup} {BUILDER_TAG}".strip())
            supply_hubs.append(hub_label.get(hub, ""))
            record(f"supply {loc}", f"new supply point, supplier {sup or '-'}, hub {hub_label.get(hub, '-')}")
        grow = {"contract": np.zeros((P, K)), "hub_price": hp, "hub_premium": np.zeros((P, K)),
                "perc_min": np.zeros((P, K)), "perc_max": np.ones((P, K)),
                "supply_inv_cap": np.zeros((P, K)), "supply_inv_min": np.zeros((P, K))}
        for key, add in grow.items():
            arrays[key] = np.concatenate([arrays[key], add], axis=1)
        arrays["supply_inv_init"] = np.concatenate([arrays["supply_inv_init"], np.zeros(K)])
        lane_fill = {"trans_cost": np.nan, "ship_min": 0.0, "ship_max": DEFAULT_MAX_SHIP}
        for key, fill in lane_fill.items():
            arrays[key] = np.concatenate([arrays[key], np.full((P, K, M0), fill)], axis=1)
        cached_price = np.concatenate([cached_price, np.full((P, K), np.nan)], axis=1)
        price_touched = np.concatenate([price_touched, np.ones((P, K), dtype=bool)], axis=1)

    # ---- demand points (new plants) --------------------------------------- #
    new_dem = list(structure.get("demand") or [])
    if len(new_dem) > BUILDER_MAX_ADDED:
        raise OverrideError(f"at most {BUILDER_MAX_ADDED} new demand points")
    if new_dem:
        L = len(new_dem)
        N = arrays["contract"].shape[1]
        for k, d in enumerate(new_dem):
            name = str(d.get("name") or "").strip() or f"New plant {k + 1}"
            demand_names.append(name)
            record(f"demand {name}", "new demand point")
        for key in ("demand", "demand_inv_cap", "demand_inv_min"):
            arrays[key] = np.concatenate([arrays[key], np.zeros((P, L))], axis=1)
        arrays["demand_inv_init"] = np.concatenate([arrays["demand_inv_init"], np.zeros(L)])
        lane_fill = {"trans_cost": np.nan, "ship_min": 0.0, "ship_max": DEFAULT_MAX_SHIP}
        for key, fill in lane_fill.items():
            arrays[key] = np.concatenate([arrays[key], np.full((P, N, L), fill)], axis=2)

    return log, cached_price, price_touched


class ModelBaseReader:
    """Pulls a :class:`ModelData` out of one ``*.xlsm`` workbook."""

    def __init__(self, path: str, area: str, region: str,
                 n_periods: int = N_PERIODS,
                 deck: Optional["OverrideDeck"] = None,
                 want_saved_cost: bool = False,
                 cell_edits: Optional[Sequence[Dict[str, object]]] = None,
                 structure: Optional[Dict[str, object]] = None,
                 demand_scale: Optional[np.ndarray] = None,
                 demand_scale_note: str = ""):
        #: Position-based edits from the web app - see apply_cell_edits().
        self.cell_edits = list(cell_edits or [])
        #: (P, M) multipliers on demand (the web app's "% of demand by state"),
        #: applied after the builder and before cell edits, so a demand typed
        #: in by hand is still taken exactly as typed. Rows/plants outside the
        #: array (added by the builder) are left alone.
        self.demand_scale = None if demand_scale is None else np.asarray(demand_scale, dtype=float)
        self.demand_scale_note = demand_scale_note
        #: Model-builder resize from the web app - see extend_structure().
        self.structure = dict(structure or {})
        self.path = path
        self.area = area
        self.region = region
        self.n_periods = n_periods
        self.notes: List[str] = []
        self.repairs: List[str] = []
        #: Hand-entered overrides to apply on top of the workbook, or None.
        self.deck = deck
        #: Filled in by read() when --compare-saved is active.
        self.want_saved_cost = want_saved_cost
        self.saved_cost: Optional[Dict[str, float]] = None
        #: Audit trail of every override actually applied to this area.
        self.override_log: List["OverrideRecord"] = []

    # -- public ------------------------------------------------------------ #

    def read(self) -> ModelData:
        grid, names = self.read_parts()
        return self._build(grid, names)

    def build_from(self, grid: "SheetGrid", names: Dict[str, str]) -> ModelData:
        """
        Build a ModelData from a sheet already loaded by read_parts().

        Opening the workbook is the slow part (seconds for the big areas);
        building from the loaded grid is milliseconds. The web app loads each
        workbook once and rebuilds from the cached grid after every edit.
        """
        return self._build(grid, names)

    def read_parts(self) -> Tuple["SheetGrid", Dict[str, str]]:
        """Open the workbook (read-only) and return the ModelBase grid and names."""
        LOG.debug("opening %s", self.path)
        wb = openpyxl.load_workbook(
            self.path, data_only=True, read_only=True, keep_links=False
        )
        try:
            if MODEL_SHEET not in wb.sheetnames:
                raise ModelStructureError(
                    f"{os.path.basename(self.path)}: no '{MODEL_SHEET}' worksheet "
                    f"(found: {', '.join(wb.sheetnames)})"
                )
            # Refuse before doing any real work, so an unsupported build can
            # never reach the solver and come back with a plausible number.
            self._refuse_unsupported_features(wb)
            names = self._global_defined_names(wb)
            grid = SheetGrid.from_worksheet(wb[MODEL_SHEET])
            if self.want_saved_cost:
                self.saved_cost = self._read_saved_cost(wb, names)
        finally:
            wb.close()

        return grid, names

    # -- unsupported-feature guard ----------------------------------------- #

    def _refuse_unsupported_features(self, wb) -> None:
        """
        Reject workbooks built with the generator's optional second transport
        route or its fuel-blending mode.

        Neither is modelled here, and neither can be safely ignored. An
        ``AltSupplyRoute`` build carries a second complete block of shipment
        variables; this reader would parse only the first, then solve happily
        and report a cost that understates both freight and the routing options
        actually available. A blunt refusal the team can act on is far better
        than a wrong answer that looks clean.
        """
        found: List[str] = []

        # Scan EVERY defined name, global and sheet-local. The marker may well
        # sit on a sheet-local name, and for a guard we want maximum reach.
        for name in self._all_defined_name_keys(wb):
            bare = name.rsplit("!", 1)[-1]
            if ALT_ROUTE_NAME_RE.search(bare):
                found.append(
                    f"second transport route (alt shipping) - defined name {bare!r}"
                )
                break

        for sheet in wb.sheetnames:
            if sheet.strip().lower() == MIXED_FUEL_SHEET:
                found.append(f"fuel blending - worksheet {sheet!r} present")
                break
        else:
            for name in self._all_defined_name_keys(wb):
                bare = name.rsplit("!", 1)[-1]
                if bare.lower().startswith(MIXED_FUEL_NAME_PREFIX):
                    found.append(
                        f"fuel blending - defined name {bare!r}"
                    )
                    break

        if found:
            raise ModelStructureError(
                f"{os.path.basename(self.path)}: this workbook uses generator "
                f"feature(s) this model does not implement: "
                + "; ".join(found)
                + ". Refusing to solve rather than silently modelling only part "
                  "of it, which would understate cost and routing. Re-build the "
                  "area as a single-route, propane-only model, or extend model.py."
            )

    @staticmethod
    def _all_defined_name_keys(wb) -> List[str]:
        """Names of every defined name in the workbook, global and local."""
        keys: List[str] = []
        container = getattr(wb, "defined_names", None)
        if container is not None:
            try:
                keys.extend(str(k) for k in container.keys())
            except AttributeError:
                keys.extend(str(dn.name)
                            for dn in getattr(container, "definedName", []))
        for ws in wb.worksheets:
            local = getattr(ws, "defined_names", None)
            if not local:
                continue
            try:
                keys.extend(f"{ws.title}!{k}" for k in local.keys())
            except AttributeError:
                continue
        return keys

    # -- saved cost (the regression oracle) -------------------------------- #

    def _read_saved_cost(self, wb, names: Dict[str, str]) -> Optional[Dict[str, float]]:
        """
        Read the annual cost block the generator already wrote onto
        'Output Propane'.

        This is the only figure in the workbook that can be compared against a
        fresh solve, which makes it the one real check on transcription
        fidelity. The anchor is awkward: the generator stores
        ``p_TopAnnualCommCost`` as a quoted string like ``'"$C$33"'`` rather
        than a proper range, and without a sheet name, so it needs unpicking by
        hand rather than through parse_range().

        Layout below the anchor, confirmed against every uploaded workbook::

            C                 D
            Commodity         current      previous baseline
            Premium           ...          ...
            Transportation    ...          ...
            Total             ...          ...

        Row offsets are verified against the caption in column B where
        possible, and only fall back to the fixed offsets if the captions have
        been edited.
        """
        raw = names.get(SAVED_COST_NAME)
        if not raw:
            self.notes.append(
                f"--compare-saved: no {SAVED_COST_NAME!r} defined name, so the "
                f"saved plan cost could not be located."
            )
            return None
        if SAVED_COST_SHEET not in wb.sheetnames:
            self.notes.append(
                f"--compare-saved: no {SAVED_COST_SHEET!r} sheet, so the saved "
                f"plan cost could not be read."
            )
            return None

        m = re.search(r"\$?([A-Za-z]{1,3})\$?(\d+)", str(raw))
        if not m:
            self.notes.append(
                f"--compare-saved: could not make sense of {SAVED_COST_NAME}="
                f"{raw!r}."
            )
            return None
        anchor_row = int(m.group(2))

        ws = wb[SAVED_COST_SHEET]
        grid = SheetGrid.from_worksheet(ws)

        out: Dict[str, float] = {}
        for key, offset, captions in SAVED_COST_ROWS:
            row = anchor_row + offset
            label = normalise_caption(grid.get(row, 2))
            if label and not any(c in label for c in captions):
                # Captions have drifted; search a small window for the right one.
                for probe in range(anchor_row - 2, anchor_row + 8):
                    probe_label = normalise_caption(grid.get(probe, 2))
                    if probe_label and any(c in probe_label for c in captions):
                        row = probe
                        break
            out[key] = grid.num(row, SAVED_COST_COL_CURRENT)
            out[f"{key}_baseline"] = grid.num(row, SAVED_COST_COL_BASELINE)

        if all(abs(v) < 1e-9 for v in out.values()):
            self.notes.append(
                "--compare-saved: the saved cost block on 'Output Propane' is "
                "empty, so there is nothing to compare against."
            )
            return None
        return out

    # -- defined names ----------------------------------------------------- #

    @staticmethod
    def _global_defined_names(wb) -> Dict[str, str]:
        """
        Workbook-global defined names as ``{name: target}``.

        Sheet-local names are deliberately left out here: in these workbooks
        the sheet-local copies point at *other* workbooks entirely, via
        external links (``[1]ModelBase!...``) left over from the Scenarios
        sheet. Only the global set actually describes the workbook we're
        holding.
        """
        out: Dict[str, str] = {}
        container = getattr(wb, "defined_names", None)
        if container is None:
            return out

        # openpyxl 3.1+ exposes this as a mapping of name -> DefinedName.
        try:
            items = list(container.items())
        except AttributeError:
            # Older openpyxl (3.0) keeps .definedName as a flat list instead,
            # with localSheetId marking which ones are sheet-local.
            items = [
                (dn.name, dn)
                for dn in getattr(container, "definedName", [])
                if getattr(dn, "localSheetId", None) is None
            ]

        for name, dn in items:
            if getattr(dn, "localSheetId", None) is not None:
                continue
            target = getattr(dn, "attr_text", None) or getattr(dn, "value", None)
            if isinstance(target, str) and target.strip():
                out[str(name)] = target.strip()
        return out

    def _need_range(self, names: Dict[str, str], key: str) -> CellRange:
        if key not in names:
            raise ModelStructureError(
                f"{os.path.basename(self.path)}: missing workbook-level defined "
                f"name {key!r}. This file does not look like a ModelBase "
                f"optimization workbook."
            )
        try:
            rng = parse_range(names[key])
        except ValueError as exc:
            raise ModelStructureError(
                f"{os.path.basename(self.path)}: defined name {key!r} -> "
                f"{names[key]!r} could not be parsed ({exc})."
            ) from exc
        if rng.sheet != MODEL_SHEET:
            raise ModelStructureError(
                f"{os.path.basename(self.path)}: {key!r} points at sheet "
                f"{rng.sheet!r}, expected {MODEL_SHEET!r}."
            )
        return rng

    # -- assembly ---------------------------------------------------------- #

    def _build(self, grid: SheetGrid, names: Dict[str, str]) -> ModelData:
        return self.build_from_inputs(self.raw_inputs(grid, names))

    def raw_inputs(self, grid: SheetGrid, names: Dict[str, str]) -> Dict[str, object]:
        """
        The model's typed-in inputs, read off the ModelBase sheet: every input
        array (in kgal and $/kgal, before any override or edit), the cached
        delivered price, and the supply / demand / period labels. This is
        all a model needs - a base model stores exactly this.
        """
        P = self.n_periods

        contract_rng = [self._need_range(names, f"Period{p}_p_Contract") for p in range(1, P + 1)]
        price_rng = [self._need_range(names, f"Period{p}_p_Price") for p in range(1, P + 1)]
        hubp_rng = [self._need_range(names, f"Period{p}_p_HubPrice") for p in range(1, P + 1)]
        prem_rng = [self._need_range(names, f"Period{p}_p_HubPremium") for p in range(1, P + 1)]
        perc_rng = [self._need_range(names, f"Period{p}_p_Perc") for p in range(1, P + 1)]
        tc_rng = [self._need_range(names, f"Period{p}_p_Trans_Cost") for p in range(1, P + 1)]
        # The Ship blocks themselves are all decision variables, so their
        # CONTENTS are irrelevant here - but their row spans locate the
        # supply-side inventory columns that sit alongside them.
        ship_rng = [self._need_range(names, f"Period{p}_p_Ship") for p in range(1, P + 1)]

        N = contract_rng[0].n_rows
        M = tc_rng[0].n_cols
        if N < 1 or M < 1:
            raise ModelStructureError(
                f"{os.path.basename(self.path)}: degenerate size N={N}, M={M}."
            )

        # -- make sure the named ranges actually agree with each other ----- #
        for p in range(P):
            for label, rng, want_rows in (
                (f"Period{p+1}_p_Contract", contract_rng[p], N),
                (f"Period{p+1}_p_Price", price_rng[p], N),
                (f"Period{p+1}_p_HubPrice", hubp_rng[p], N),
                (f"Period{p+1}_p_HubPremium", prem_rng[p], N),
                (f"Period{p+1}_p_Perc", perc_rng[p], N),
                (f"Period{p+1}_p_Trans_Cost", tc_rng[p], N),
            ):
                if rng.n_rows != want_rows:
                    raise ModelStructureError(
                        f"{os.path.basename(self.path)}: {label} covers "
                        f"{rng.n_rows} rows, expected {want_rows}."
                    )
            if tc_rng[p].n_cols != M:
                raise ModelStructureError(
                    f"{os.path.basename(self.path)}: Period{p+1}_p_Trans_Cost has "
                    f"{tc_rng[p].n_cols} columns, expected {M}."
                )

        supply_row1 = contract_rng[0].row1
        supply_row2 = contract_rng[0].row2
        perc_row1 = perc_rng[0].row1
        perc_row2 = perc_rng[0].row2

        # Supply point names are read up front so the data-quality messages
        # below can say "Northfield, ST" instead of "row 12".
        supply_names_raw = [grid.text(r, 3) for r in range(supply_row1, supply_row2 + 1)]

        # -- per-period scalar / vector data ------------------------------- #
        contract = np.zeros((P, N))
        price = np.zeros((P, N))
        hub_price = np.zeros((P, N))
        hub_premium = np.zeros((P, N))
        perc_min = np.zeros((P, N))
        perc_max = np.zeros((P, N))
        cached_price = np.full((P, N), float("nan"))

        for p in range(P):
            contract[p] = grid.column(supply_row1, supply_row2, contract_rng[p].col1)
            hub_price[p] = grid.column(price_rng[p].row1, price_rng[p].row2, hubp_rng[p].col1)
            hub_premium[p] = grid.column(price_rng[p].row1, price_rng[p].row2, prem_rng[p].col1)

            # In theory Excel col D = col R + col AF, but people sometimes
            # hard-code column D by hand. Keep the cached value aside for now:
            # price is DERIVED further down, after any hand-entered overrides
            # have had a chance to change hub price or premium. Deriving it
            # here would leave it stale the moment somebody overrides a diff.
            cached_price[p] = grid.column(
                price_rng[p].row1, price_rng[p].row2, price_rng[p].col1,
                default=float("nan"),
            )

            # The percentage block shares the same rows as Perc: the minimum
            # lives in the Price columns and the maximum in the HubPremium
            # columns. Slightly confusing reuse of layout, but that's how the
            # sheet is built.
            perc_min[p] = grid.column(perc_row1, perc_row2, price_rng[p].col1)
            perc_max[p] = grid.column(perc_row1, perc_row2, prem_rng[p].col1)

        # -- labels --------------------------------------------------------- #
        # Column B holds the supplier / counterparty, C the supply point
        # itself, D the market hub.
        supply_names = [name or f"Supply {k + 1}"
                        for k, name in enumerate(supply_names_raw)]
        supply_groups = [grid.text(r, 2) for r in range(supply_row1, supply_row2 + 1)]
        supply_hubs = [grid.text(r, 4) for r in range(supply_row1, supply_row2 + 1)]

        period_labels = self._period_labels(grid, price_rng)

        # -- blocks located by caption -------------------------------------- #
        sup_inv = self._inventory_block(
            grid, CAPTION_SUPPLY_INVENTORY, row_lo=perc_row2, n_rows=N, what="supply"
        )
        dem_block = self._demand_block(grid, row_lo=sup_inv["row2"], n_rows=M)
        dem_inv = self._inventory_block(
            grid, CAPTION_DEMAND_INVENTORY, row_lo=dem_block["row2"], n_rows=M, what="demand"
        )

        demand_names = [
            grid.text(r, dem_block["name_col"]) or f"Demand {r - dem_block['row1'] + 1}"
            for r in range(dem_block["row1"], dem_block["row2"] + 1)
        ]

        demand = np.zeros((P, M))
        for p in range(P):
            col = dem_block["period_col1"] + p
            demand[p] = grid.column(dem_block["row1"], dem_block["row2"], col)

        supply_inv_init = grid.column(sup_inv["row1"], sup_inv["row2"], sup_inv["init_col"])
        supply_inv_cap = np.zeros((P, N))
        for p in range(P):
            supply_inv_cap[p] = grid.column(
                sup_inv["row1"], sup_inv["row2"], sup_inv["cap_col1"] + p
            )

        demand_inv_init = grid.column(dem_inv["row1"], dem_inv["row2"], dem_inv["init_col"])
        demand_inv_cap = np.zeros((P, M))
        for p in range(P):
            demand_inv_cap[p] = grid.column(
                dem_inv["row1"], dem_inv["row2"], dem_inv["cap_col1"] + p
            )

        # Minimum inventory. The generator writes an editable Min Inv cell on
        # BOTH sides, defaulting to 0, and LoadSolverUp constrains ending
        # inventory to be >= it at both supply and demand points. Reading only
        # the demand side (as this used to) meant a planner who typed a minimum
        # working inventory at a supply point had it honoured in Excel and
        # silently dropped here.
        supply_inv_min = self._supply_min_inventory(grid, ship_rng, N, P)
        demand_inv_min = self._demand_min_inventory(grid, dem_block, dem_inv, M, P)

        # -- transport cost + lane bounds ----------------------------------- #
        trans_cost = np.zeros((P, N, M))
        ship_min = np.zeros((P, N, M))
        ship_max = np.full((P, N, M), DEFAULT_MAX_SHIP)

        # The three matrices sit at the same columns in every period block, so
        # resolve the layout once from period 1 and reuse it.
        min_col1, max_col1 = self._transport_side_blocks(grid, tc_rng[0], M)
        for p in range(P):
            block = tc_rng[p]
            # NaN, not 0, for a blank freight cell: blank means "this lane
            # can't be used" (see "closed lanes" below), which a 0 - a free
            # lane - would turn into the opposite.
            trans_cost[p] = grid.matrix(block, default=float("nan"))
            ship_min[p] = grid.matrix(
                CellRange(MODEL_SHEET, block.row1, min_col1,
                          block.row2, min_col1 + M - 1)
            )
            raw = grid.matrix(
                CellRange(MODEL_SHEET, block.row1, max_col1,
                          block.row2, max_col1 + M - 1),
                default=float("nan"),
            )
            # Same idea as before: a blank maximum means "unbounded", not
            # "zero" - so fall back to the default ceiling rather than 0.
            ship_max[p] = np.where(np.isnan(raw), DEFAULT_MAX_SHIP, raw)

        # ------------------------------------------------------------------ #
        # Hand-entered overrides, then derived values, then validation.
        #
        # This ordering is deliberate and matters:
        #   1. overrides land on the RAW workbook inputs
        #   2. price is derived from the (possibly overridden) hub price + diff
        #   3. every validation and repair runs over the FINAL numbers
        # An override can introduce a contradiction just as easily as a
        # workbook can, so the checks have to come last, not before step 1.
        # ------------------------------------------------------------------ #
        arrays = {
            "contract": contract,
            "hub_price": hub_price,
            "hub_premium": hub_premium,
            "perc_min": perc_min,
            "perc_max": perc_max,
            "demand": demand,
            "trans_cost": trans_cost,
            "ship_min": ship_min,
            "ship_max": ship_max,
            "supply_inv_init": supply_inv_init,
            "supply_inv_cap": supply_inv_cap,
            "supply_inv_min": supply_inv_min,
            "demand_inv_init": demand_inv_init,
            "demand_inv_cap": demand_inv_cap,
            "demand_inv_min": demand_inv_min,
        }
        return {
            "arrays": arrays, "cached_price": cached_price,
            "supply_names": supply_names, "supply_groups": supply_groups,
            "supply_hubs": supply_hubs, "demand_names": demand_names,
            "period_labels": period_labels,
        }

    def build_from_inputs(self, raw: Dict[str, object]) -> ModelData:
        """
        Build a ModelData from raw inputs - the typed-in numbers of a model,
        as returned by :meth:`raw_inputs` or stored in a base model
        (basemodels.py) - without any workbook.

        Deck overrides, model-builder structure, cell edits, price derivation
        and every validation run exactly as for a workbook. *raw* is never
        modified: every array is copied first.
        """
        arrays = {k: np.array(v, dtype=float, copy=True) for k, v in raw["arrays"].items()}
        cached_price = np.array(raw.get("cached_price"), dtype=float, copy=True) \
            if raw.get("cached_price") is not None else np.full(arrays["contract"].shape, np.nan)
        supply_names = list(raw["supply_names"])
        supply_groups = list(raw["supply_groups"])
        supply_hubs = list(raw["supply_hubs"])
        demand_names = list(raw["demand_names"])
        period_labels = list(raw["period_labels"])
        contract, hub_price, hub_premium = arrays["contract"], arrays["hub_price"], arrays["hub_premium"]
        perc_min, perc_max, demand = arrays["perc_min"], arrays["perc_max"], arrays["demand"]
        trans_cost, ship_min, ship_max = arrays["trans_cost"], arrays["ship_min"], arrays["ship_max"]
        supply_inv_init, supply_inv_cap = arrays["supply_inv_init"], arrays["supply_inv_cap"]
        supply_inv_min, demand_inv_init = arrays["supply_inv_min"], arrays["demand_inv_init"]
        demand_inv_cap, demand_inv_min = arrays["demand_inv_cap"], arrays["demand_inv_min"]
        P, N = contract.shape
        M = demand.shape[1]

        price_touched = np.zeros((P, N), dtype=bool)
        if self.deck is not None:
            self.override_log = apply_overrides(
                deck=self.deck,
                area=self.area,
                arrays=arrays,
                supply_names=supply_names,
                supply_groups=supply_groups,
                supply_hubs=supply_hubs,
                demand_names=demand_names,
                period_labels=period_labels,
                price_touched=price_touched,
            )
        # The model builder (web app): extra periods, hubs, supply points and
        # plants. Applied BEFORE the cell edits so the new rows and months can be
        # filled in through the ordinary grids like any other cell.
        if self.structure:
            built, cached_price, price_touched = extend_structure(
                self.structure, self.area, arrays,
                supply_names=supply_names, supply_groups=supply_groups,
                supply_hubs=supply_hubs, demand_names=demand_names,
                period_labels=period_labels, cached_price=cached_price,
                price_touched=price_touched,
            )
            self.override_log.extend(built)
            # The arrays were replaced by resized ones - rebind every local.
            contract, hub_price, hub_premium = arrays["contract"], arrays["hub_price"], arrays["hub_premium"]
            perc_min, perc_max, demand = arrays["perc_min"], arrays["perc_max"], arrays["demand"]
            trans_cost, ship_min, ship_max = arrays["trans_cost"], arrays["ship_min"], arrays["ship_max"]
            supply_inv_init, supply_inv_cap = arrays["supply_inv_init"], arrays["supply_inv_cap"]
            supply_inv_min, demand_inv_init = arrays["supply_inv_min"], arrays["demand_inv_init"]
            demand_inv_cap, demand_inv_min = arrays["demand_inv_cap"], arrays["demand_inv_min"]
            P, N = contract.shape
            M = demand.shape[1]

        # Demand scaling (the web app's "% of demand by state").
        if self.demand_scale is not None and self.demand_scale.size:
            sp = min(P, self.demand_scale.shape[0])
            sm = min(M, self.demand_scale.shape[1])
            fac = self.demand_scale[:sp, :sm]
            if not np.allclose(fac, 1.0):
                before = float(demand.sum())
                demand[:sp, :sm] = demand[:sp, :sm] * fac
                self.override_log.append(OverrideRecord(
                    area=self.area, target="demand by state", field="demand %",
                    before=f"{before:,.1f} kgal", after=f"{float(demand.sum()):,.1f} kgal",
                    line_no=len(self.override_log) + 1,
                    scope=f"[web{': ' + self.demand_scale_note if self.demand_scale_note else ''}]"))

        # Cell-level edits (the web app). Same point in the pipeline as the
        # deck, applied after it, so price derivation and every validation
        # below still run over the final numbers.
        if self.cell_edits:
            self.override_log.extend(apply_cell_edits(
                self.cell_edits, self.area, arrays,
                supply_names=supply_names, supply_groups=supply_groups,
                demand_names=demand_names, period_labels=period_labels,
                price_touched=price_touched,
            ))

        # -- derived: delivered price --------------------------------------- #
        # Trust the workbook's own cached column D where it exists, EXCEPT for
        # cells whose hub price or differential an override has just changed -
        # there the cache is stale by definition and must be rebuilt.
        rebuilt = hub_price + hub_premium
        price = np.where(np.isnan(cached_price) | price_touched, rebuilt, cached_price)

        # -- validation and repair (over the final numbers) ----------------- #
        # A blank maximum means "nobody entered a ceiling", not "take nothing".
        # Excel would read that blank as 0 and quietly shut the supplier off
        # without telling anyone, so flag it loudly here instead of silently
        # producing a different answer than someone expects.
        blank_max = (perc_max <= 0) & (contract > 0)
        if blank_max.any():
            self.notes.append(
                f"{int(blank_max.sum())} supply/period cell(s) have a contracted "
                f"volume but a maximum lift of 0%; those suppliers cannot be used."
            )

        # A minimum above the maximum is simply a typo - there is no value that
        # satisfies both, so Excel's solver would fail here too. Rather than
        # refuse to produce anything, trust the ceiling (the more conservative
        # purchase), pull the floor down to meet it, and record exactly what
        # changed so it can be corrected at source.
        bad = perc_min > perc_max + 1e-9
        if bad.any():
            for p, i in np.argwhere(bad):
                self.repairs.append(
                    f"{supply_names[int(i)]} in {period_labels[int(p)]}: minimum "
                    f"lift {perc_min[p, i]:.4g} is above maximum "
                    f"{perc_max[p, i]:.4g}; minimum lowered to "
                    f"{perc_max[p, i]:.4g} so the model can run."
                )
            perc_min = np.minimum(perc_min, perc_max)

        # Negative demand is never meaningful - it would require a location to
        # ship product back out, which nothing in the model can do, so the LP
        # becomes infeasible with a very unhelpful error. Usually a demand
        # formula that went negative after an adjustment.
        neg_demand = demand < -1e-9
        if neg_demand.any():
            for p, j in np.argwhere(neg_demand):
                self.repairs.append(
                    f"{demand_names[int(j)]} has NEGATIVE demand of "
                    f"{demand[p, j]:,.1f} kgal in {period_labels[int(p)]}; treated "
                    f"as zero. Check the demand formula for this location."
                )
            demand = np.maximum(demand, 0.0)

        # -- closed lanes ---------------------------------------------------- #
        # A freight cell left blank (or holding text such as "n/a") means
        # "do not use this lane", and the LP holds that lane at 0. This is a
        # hard block, unlike the older habit of typing $1/gal (1000 $/kgal),
        # which is only a price - and stops deterring anything where real
        # freight runs above $1/gal, as some remote lanes can. Freight becomes 0
        # here only so the arithmetic stays clean; nothing can move on it.
        lane_closed = ~np.isfinite(trans_cost)
        if lane_closed.any():
            trans_cost = np.where(lane_closed, 0.0, trans_cost)
            forced = lane_closed & (ship_min > 1e-9)
            if forced.any():
                self.notes.append(
                    f"{int(forced.sum())} lane/month cell(s) have a minimum shipment "
                    f"but no freight rate (blank = lane not used); the minimum was "
                    f"dropped. Enter a freight rate to use the lane."
                )
                ship_min = np.where(lane_closed, 0.0, ship_min)

        np.clip(ship_min, 0.0, None, out=ship_min)
        crossed = (ship_min > ship_max + 1e-9) & ~lane_closed
        if crossed.any():
            self.notes.append(
                f"{int(crossed.sum())} lane(s) have a minimum shipment above the "
                f"maximum; minimum lowered to the maximum to keep the LP solvable."
            )
            ship_min = np.minimum(ship_min, ship_max)

        # Storage minimum above storage capacity, on either side. Same
        # treatment as the other crossed bounds: trust the capacity, clamp the
        # minimum down to it, and say so.
        for label, lo_arr, hi_arr, names_list in (
            ("supply", supply_inv_min, supply_inv_cap, supply_names),
            ("demand", demand_inv_min, demand_inv_cap, demand_names),
        ):
            bad_store = lo_arr > hi_arr + 1e-9
            if bad_store.any():
                for p, k in np.argwhere(bad_store):
                    self.repairs.append(
                        f"{names_list[int(k)]} in {period_labels[int(p)]}: minimum "
                        f"{label} inventory {lo_arr[p, k]:,.1f} kgal exceeds the "
                        f"storage capacity {hi_arr[p, k]:,.1f} kgal; minimum "
                        f"lowered to the capacity."
                    )
                np.minimum(lo_arr, hi_arr, out=lo_arr)

        # -- high-freight census (a warning only) ---------------------------- #
        high = (trans_cost >= HIGH_FREIGHT_RATE - HIGH_FREIGHT_TOL) & ~lane_closed
        blocked_pairs = int(high.any(axis=0).sum())

        data = ModelData(
            area=self.area,
            region=self.region,
            source_path=self.path,
            n_supply=N,
            n_demand=M,
            n_periods=P,
            period_labels=period_labels,
            supply_names=supply_names,
            supply_groups=supply_groups,
            supply_hubs=supply_hubs,
            demand_names=demand_names,
            contract=contract,
            price=price,
            hub_price=hub_price,
            hub_premium=hub_premium,
            perc_min=perc_min,
            perc_max=perc_max,
            demand=demand,
            trans_cost=trans_cost,
            ship_min=ship_min,
            ship_max=ship_max,
            supply_inv_init=supply_inv_init,
            supply_inv_cap=supply_inv_cap,
            supply_inv_min=supply_inv_min,
            demand_inv_init=demand_inv_init,
            demand_inv_cap=demand_inv_cap,
            demand_inv_min=demand_inv_min,
            notes=list(self.notes),
            repairs=list(self.repairs),
            blocked_lane_pairs=blocked_pairs,
            lane_closed=lane_closed if lane_closed.any() else None,
            saved_cost=self.saved_cost,
            override_log=list(self.override_log),
        )
        LOG.info(
            "%-16s %2d supply x %3d demand x %2d periods, demand %s kgal",
            self.area, N, M, P, f"{data.total_demand:,.0f}",
        )
        return data

    # -- block locators ---------------------------------------------------- #

    def _period_labels(self, grid: SheetGrid, price_rng: Sequence[CellRange]) -> List[str]:
        """
        Pull period captions from the header row above the tables
        (ModelBase row 2).

        Those cells usually hold real dates; if they're blank we just fall
        back to ``P1..P12`` rather than leaving something empty.
        """
        labels: List[str] = []
        for p, rng in enumerate(price_rng, start=1):
            label = ""
            for header_row in (2, 1, 3):
                raw = grid.get(header_row, rng.col1)
                if raw is None:
                    continue
                if isinstance(raw, (_dt.datetime, _dt.date)):
                    label = raw.strftime("%b-%y")
                    break
                if isinstance(raw, (int, float)) and 20000 < float(raw) < 80000:
                    # Looks like an Excel serial date (1900 epoch, with the
                    # usual 2-day fudge factor Excel carries for compatibility).
                    label = (
                        _dt.date(1899, 12, 30) + _dt.timedelta(days=int(raw))
                    ).strftime("%b-%y")
                    break
                text = as_text(raw)
                if text:
                    label = text
                    break
            labels.append(label or f"P{p}")

        if len(set(labels)) != len(labels):  # duplicate captions aren't useful to anyone
            labels = [f"P{i}" for i in range(1, len(labels) + 1)]
        return labels

    def _inventory_block(
        self, grid: SheetGrid, captions: Sequence[str], row_lo: int, n_rows: int, what: str
    ) -> Dict[str, int]:
        """
        Locate an "Inventory at ... Points" table.

        The layout, checked against both a 10x10 and a 30x45 workbook::

            <caption>  | Initial Inventory | ... | Available Storage Capacity | p1 | p2 ...
                       | <value>           |     | <name>                    | cap| cap ...

        In every workbook we've seen, the data column sits immediately to the
        right of its caption.
        """
        hit = grid.find_caption(captions, row_lo=row_lo)
        if hit is None:
            raise ModelStructureError(
                f"{os.path.basename(self.path)}: could not find the "
                f"'{captions[0]}' block below row {row_lo}."
            )
        header_row, _ = hit
        row1 = header_row + 1
        row2 = row1 + n_rows - 1

        init_col_caption = grid.find_caption_in_row(header_row, CAPTION_INITIAL_INVENTORY)
        cap_col_caption = grid.find_caption_in_row(header_row, CAPTION_STORAGE_CAPACITY)
        if init_col_caption is None or cap_col_caption is None:
            raise ModelStructureError(
                f"{os.path.basename(self.path)}: the {what} inventory block at row "
                f"{header_row} is missing the 'Initial Inventory' and/or "
                f"'Available Storage Capacity' heading."
            )
        return {
            "header_row": header_row,
            "row1": row1,
            "row2": row2,
            "init_col": init_col_caption + 1,
            "cap_col1": cap_col_caption + 1,
        }

    def _demand_block(self, grid: SheetGrid, row_lo: int, n_rows: int) -> Dict[str, int]:
        """
        Locate the demand table.

        ``Demand`` sits alone in column B, with ``Geography`` directly under
        it in column C, and the twelve monthly columns starting one to the
        right of that.
        """
        search_from = row_lo
        while True:
            hit = grid.find_caption(CAPTION_DEMAND, row_lo=search_from)
            if hit is None:
                raise ModelStructureError(
                    f"{os.path.basename(self.path)}: could not find the 'Demand' "
                    f"block below row {row_lo}."
                )
            header_row, _ = hit
            geo_col = grid.find_caption_in_row(header_row + 1, CAPTION_DEMAND_GEOGRAPHY)
            if geo_col is not None:
                row1 = header_row + 2
                return {
                    "header_row": header_row,
                    "row1": row1,
                    "row2": row1 + n_rows - 1,
                    "name_col": geo_col,
                    "period_col1": geo_col + 1,
                }
            # That was just a stray "Demand" caption somewhere else on the
            # sheet - keep scanning further down for the real table.
            search_from = header_row + 1

    def _supply_min_inventory(
        self, grid: SheetGrid, ship_rng: Sequence[CellRange], n_supply: int, n_periods: int
    ) -> np.ndarray:
        """
        Read the supply-side Min Inv column that sits beside the Ship blocks.

        Unlike the demand side this is a SINGLE column, not a 12-wide block.
        Each period's Ship block occupies its own rows, so one column covers
        every period::

            <- Ship matrix ->   Propane Supply Balance | Initial Inv | End Inv | Min Inv | Max Inv
            period 1 rows       ...                      ...           ...       ...       ...
            period 2 rows       ...                      ...           ...       ...       ...

        Returns (P, N) of zeros if the caption cannot be found, which is the
        generator's own default and therefore a safe fallback.
        """
        out = np.zeros((n_periods, n_supply))
        caption_row = ship_rng[0].row1 - 1
        if caption_row < 1:
            return out

        col = grid.find_caption_in_row(caption_row, CAPTION_SUPPLY_MIN_INV)
        if col is None:
            self.notes.append(
                f"No supply-side 'Min Inv' heading found on row {caption_row}; "
                f"assuming a minimum inventory of 0 at every supply point."
            )
            return out

        for p in range(n_periods):
            block = ship_rng[p]
            for i in range(n_supply):
                out[p, i] = grid.num(block.row1 + i, col)
        return out

    def _demand_min_inventory(
        self, grid: SheetGrid, dem_block: Dict[str, int], dem_inv: Dict[str, int],
        n_demand: int, n_periods: int
    ) -> np.ndarray:
        """
        Read the demand-side Min Inventory block from the Propane Demand
        Balance section.

        This one IS a 12-wide block, laid out like the other demand blocks with
        the caption one column left of the first period. Zeros are the
        generator's default, so a missing caption is not an error.
        """
        out = np.zeros((n_periods, n_demand))
        hit = grid.find_caption(CAPTION_DEMAND_BALANCE, row_lo=dem_inv["row2"])
        if hit is None:
            return out
        header_row, _ = hit

        col = grid.find_caption_in_row(header_row, CAPTION_DEMAND_MIN_INV)
        if col is None:
            return out

        # Data rows line up one-for-one with the demand points, starting on the
        # row below the caption.
        row1 = header_row + 1
        for p in range(n_periods):
            out[p] = grid.column(row1, row1 + n_demand - 1, col + 1 + p)
        return out

    def _transport_side_blocks(
        self, grid: SheetGrid, tc_block: CellRange, n_demand: int
    ) -> Tuple[int, int]:
        """
        Find the first data column of the Min / Max Transport Quantities
        matrices.

        Each period's transport section is really three matrices sitting side
        by side, sharing the same rows as the cost matrix::

            <-- n_demand --> gap <-- n_demand --> gap <-- n_demand -->
            Transport Costs      Min Transport Qty    Max Transport Qty

        Captions get read from the section header rows. If a caption is
        missing for some reason, we fall back to the template's fixed
        3-column gutter and leave a note behind so the substitution shows up
        on the Diagnostics sheet instead of just silently happening.
        """
        # tc_block.row1 is the first data row; the caption itself sits 2 rows
        # above that, and the per-column demand-point names 1 row above.
        caption_row = tc_block.row1 - 2
        rows_to_try = [r for r in (caption_row, caption_row + 1, caption_row - 1) if r >= 1]

        if not any(
            grid.find_caption_in_row(r, CAPTION_TRANSPORT_COSTS) is not None
            for r in rows_to_try
        ):
            self.notes.append(
                f"No 'Transport Costs' heading near row {caption_row}; the "
                f"transport block layout could not be confirmed."
            )

        min_col: Optional[int] = None
        max_col: Optional[int] = None
        for row in rows_to_try:
            if min_col is None:
                c = grid.find_caption_in_row(row, CAPTION_MIN_TRANSPORT)
                if c is not None:
                    min_col = c + 1
            if max_col is None:
                c = grid.find_caption_in_row(row, CAPTION_MAX_TRANSPORT)
                if c is not None:
                    max_col = c + 1
            if min_col is not None and max_col is not None:
                break

        if min_col is None:
            min_col = tc_block.col2 + 3
            self.notes.append(
                "Min Transport Quantities heading not found; assumed it starts at "
                f"column {index_to_col(min_col)}."
            )
        if max_col is None:
            max_col = min_col + n_demand + 2
            self.notes.append(
                "Max Transport Quantities heading not found; assumed it starts at "
                f"column {index_to_col(max_col)}."
            )
        return min_col, max_col


# --------------------------------------------------------------------------- #
# Building the LP and solving it.
# --------------------------------------------------------------------------- #


@dataclass
class Solution:
    """The solved (or failed) result for one area."""

    data: ModelData
    status: str                     # "optimal" | "infeasible" | "unbounded" | "error"
    message: str = ""
    objective: float = float("nan")

    ship: Optional[np.ndarray] = None          # (P, N, M)
    perc: Optional[np.ndarray] = None          # (P, N)
    supply_wd: Optional[np.ndarray] = None     # (P, N)
    supply_inj: Optional[np.ndarray] = None    # (P, N)
    supply_end: Optional[np.ndarray] = None    # (P, N)
    demand_wd: Optional[np.ndarray] = None     # (P, M)
    demand_inj: Optional[np.ndarray] = None    # (P, M)
    demand_end: Optional[np.ndarray] = None    # (P, M)
    shortfall: Optional[np.ndarray] = None     # (P, M) demand we could not serve
    excess: Optional[np.ndarray] = None        # (P, M) forced in beyond what a point can use
    surplus: Optional[np.ndarray] = None       # (P, N) volume lifted with nowhere to go
    unsourced: Optional[np.ndarray] = None     # (P, N) shipped without contract cover

    solver: str = ""
    solve_seconds: float = 0.0

    #: True when the faithful LP was infeasible and we had to relax something
    #: to get an answer. The plan is still useful, but it is NOT a plan the
    #: spreadsheet could have produced, so it must be labelled as such.
    relaxed: bool = False
    relax_note: str = ""

    #: Plain-language findings from the elastic (phase-1) solve, naming the
    #: exact constraints the workbook's data cannot satisfy.
    violations: List[str] = field(default_factory=list)

    #: HiGHS dual values of the final LP ("upper": per-variable upper-bound
    #: marginals, "eq": per-row equality marginals), and the builder that knows
    #: the variable layout. Used for "what is one more gallon worth".
    duals: Optional[Dict[str, np.ndarray]] = field(default=None, repr=False)
    lp: Optional["LPBuilder"] = field(default=None, repr=False)

    @property
    def ok(self) -> bool:
        return self.status == "optimal"

    @property
    def unmet_total(self) -> float:
        return 0.0 if self.shortfall is None else float(self.shortfall.sum())

    @property
    def surplus_total(self) -> float:
        return 0.0 if self.surplus is None else float(self.surplus.sum())

    @property
    def excess_total(self) -> float:
        return 0.0 if self.excess is None else float(self.excess.sum())

    @property
    def unsourced_total(self) -> float:
        return 0.0 if self.unsourced is None else float(self.unsourced.sum())

    @property
    def other_slack_total(self) -> float:
        """Relaxations beyond the two headline ones, for the Summary sheet."""
        return self.excess_total + self.unsourced_total

    # -- high-freight usage (warning only) ------------------------------- #

    def high_freight_volume(self) -> float:
        """
        Volume moved on lanes with freight of $1.00/gal or more. Those lanes
        are used like any other when they're the cheapest way; this only makes
        the expensive ones easy to spot so the rates get checked.
        """
        if self.ship is None:
            return 0.0
        return float(self.ship[self.data.high_freight_mask()].sum())

    def high_freight_cost(self) -> float:
        """Freight dollars spent on lanes at $1.00/gal or more."""
        if self.ship is None:
            return 0.0
        mask = self.data.high_freight_mask()
        return float((self.ship * self.data.trans_cost)[mask].sum())

    # Older names.
    blocked_lane_volume = high_freight_volume
    blocked_lane_cost = high_freight_cost

    # -- derived reporting figures ----------------------------------------- #

    def purchased(self) -> np.ndarray:
        """(P, N) volume lifted from each supply point."""
        return self.data.contract * self.perc

    def commodity_cost(self) -> np.ndarray:
        """(P, N) dollars spent on product."""
        return self.purchased() * self.data.price

    def freight_cost(self) -> np.ndarray:
        """(P, N, M) dollars spent moving product."""
        return self.ship * self.data.trans_cost

    def delivered(self) -> np.ndarray:
        """(P, M) volume arriving at each demand point."""
        return self.ship.sum(axis=1)


class LPBuilder:
    """
    Turns a :class:`ModelData` into ``min c'x  s.t.  A_eq x = b_eq,  lo <= x <= hi``.

    Every variable is continuous and every constraint is an equality; storage
    limits and lane limits live entirely in the variable bounds instead of as
    separate rows. That keeps the constraint matrix small and lets HiGHS's
    presolve do most of the heavy lifting.
    """

    def __init__(self, data: ModelData, allow_shortfall: bool = False,
                 allow_surplus: bool = False, allow_excess: bool = False,
                 allow_unsourced: bool = False, allow_elastic: bool = False):
        self.d = data
        self.allow_shortfall = allow_shortfall
        self.allow_surplus = allow_surplus
        self.allow_excess = allow_excess
        self.allow_unsourced = allow_unsourced
        self.allow_elastic = allow_elastic

        P, N, M = data.n_periods, data.n_supply, data.n_demand
        self.P, self.N, self.M = P, N, M

        # Row layout of the equality system, needed to decode elastic slacks
        # back into "which constraint, which supply point, which period".
        self.n_rows = P * (2 * N + 2 * M)
        self.row_supply_balance = 0
        self.row_supply_inventory = P * N
        self.row_demand_balance = 2 * P * N
        self.row_demand_inventory = 2 * P * N + P * M

        # Every kind of variable gets its own contiguous block, one after
        # another, so a single flat array can hold the whole problem.
        self.n_ship = P * N * M
        self.off_ship = 0
        self.off_perc = self.off_ship + self.n_ship
        self.off_s_wd = self.off_perc + P * N
        self.off_s_inj = self.off_s_wd + P * N
        self.off_s_end = self.off_s_inj + P * N
        self.off_d_wd = self.off_s_end + P * N
        self.off_d_inj = self.off_d_wd + P * M
        self.off_d_end = self.off_d_inj + P * M
        self.off_short = self.off_d_end + P * M
        self.off_excess = self.off_short + (P * M if allow_shortfall else 0)
        self.off_surplus = self.off_excess + (P * M if allow_excess else 0)
        self.off_unsourced = self.off_surplus + (P * N if allow_surplus else 0)
        self.off_epos = self.off_unsourced + (P * N if allow_unsourced else 0)
        self.off_eneg = self.off_epos + (self.n_rows if allow_elastic else 0)
        self.n_vars = self.off_eneg + (self.n_rows if allow_elastic else 0)

    # -- index helpers ----------------------------------------------------- #

    def iship(self, p: int, i: int, j: int) -> int:
        return self.off_ship + (p * self.N + i) * self.M + j

    def iperc(self, p: int, i: int) -> int:
        return self.off_perc + p * self.N + i

    def is_wd(self, p: int, i: int) -> int:
        return self.off_s_wd + p * self.N + i

    def is_inj(self, p: int, i: int) -> int:
        return self.off_s_inj + p * self.N + i

    def is_end(self, p: int, i: int) -> int:
        return self.off_s_end + p * self.N + i

    def id_wd(self, p: int, j: int) -> int:
        return self.off_d_wd + p * self.M + j

    def id_inj(self, p: int, j: int) -> int:
        return self.off_d_inj + p * self.M + j

    def id_end(self, p: int, j: int) -> int:
        return self.off_d_end + p * self.M + j

    def ishort(self, p: int, j: int) -> int:
        return self.off_short + p * self.M + j

    def iexcess(self, p: int, j: int) -> int:
        return self.off_excess + p * self.M + j

    def isurplus(self, p: int, i: int) -> int:
        return self.off_surplus + p * self.N + i

    def iunsourced(self, p: int, i: int) -> int:
        return self.off_unsourced + p * self.N + i

    # -- assembly ---------------------------------------------------------- #

    def objective(self) -> np.ndarray:
        c = np.zeros(self.n_vars)
        c[self.off_ship:self.off_ship + self.n_ship] = self.d.trans_cost.reshape(-1)
        c[self.off_perc:self.off_perc + self.P * self.N] = (
            self.d.contract * self.d.price
        ).reshape(-1)
        if self.allow_shortfall:
            c[self.off_short:self.off_short + self.P * self.M] = SHORTFALL_PENALTY
        if self.allow_excess:
            c[self.off_excess:self.off_excess + self.P * self.M] = SHORTFALL_PENALTY
        if self.allow_surplus:
            c[self.off_surplus:self.off_surplus + self.P * self.N] = SURPLUS_PENALTY
        if self.allow_unsourced:
            c[self.off_unsourced:self.off_unsourced + self.P * self.N] = SURPLUS_PENALTY
        if self.allow_elastic:
            c[self.off_epos:self.off_epos + self.n_rows] = ELASTIC_PENALTY
            c[self.off_eneg:self.off_eneg + self.n_rows] = ELASTIC_PENALTY
        return c

    def bounds(self) -> Tuple[np.ndarray, np.ndarray]:
        lo = np.zeros(self.n_vars)
        hi = np.full(self.n_vars, np.inf)

        lo[self.off_ship:self.off_ship + self.n_ship] = self.d.ship_min.reshape(-1)
        # Closed lanes (blank freight) can carry nothing.
        hi[self.off_ship:self.off_ship + self.n_ship] = np.where(
            self.d.closed_lane_mask(), 0.0, self.d.ship_max).reshape(-1)

        PN = self.P * self.N
        lo[self.off_perc:self.off_perc + PN] = self.d.perc_min.reshape(-1)
        hi[self.off_perc:self.off_perc + PN] = self.d.perc_max.reshape(-1)

        # Ending inventory sits between the planner's minimum working
        # inventory and whatever storage is available that period. The
        # generator constrains BOTH ends at supply points, same as it does at
        # demand points.
        lo[self.off_s_end:self.off_s_end + PN] = self.d.supply_inv_min.reshape(-1)
        hi[self.off_s_end:self.off_s_end + PN] = self.d.supply_inv_cap.reshape(-1)

        PM = self.P * self.M
        lo[self.off_d_end:self.off_d_end + PM] = self.d.demand_inv_min.reshape(-1)
        hi[self.off_d_end:self.off_d_end + PM] = self.d.demand_inv_cap.reshape(-1)

        # Deliberately NOT clamped to lo <= hi here. A crossed bound means the
        # input data itself is contradictory, and that's infeasible in Excel
        # too - quietly widening it would just produce a plan the spreadsheet
        # would have rejected. solve() catches this condition and reports it
        # properly instead.
        return lo, hi

    def crossed_bounds(self) -> List[str]:
        """
        Describe, in plain English, any variable whose lower bound exceeds
        its upper bound. A non-empty result means the LP is infeasible by
        construction, before the solver even gets involved.
        """
        lo, hi = self.bounds()
        bad = np.flatnonzero(lo > hi + 1e-9)
        out: List[str] = []
        for k in bad[:20]:
            out.append(f"{self.describe_var(int(k))}: min {lo[k]:,.4g} > max {hi[k]:,.4g}")
        if bad.size > 20:
            out.append(f"...and {bad.size - 20} more")
        return out

    def describe_var(self, k: int) -> str:
        """Name variable *k* in plain business terms, for error messages."""
        d = self.d
        P, N, M = self.P, self.N, self.M

        def period(p: int) -> str:
            return d.period_labels[p] if 0 <= p < len(d.period_labels) else f"P{p+1}"

        if k < self.off_perc:
            r = k - self.off_ship
            p, rem = divmod(r, N * M)
            i, j = divmod(rem, M)
            return f"shipment {d.supply_names[i]} -> {d.demand_names[j]} in {period(p)}"
        if k < self.off_s_wd:
            p, i = divmod(k - self.off_perc, N)
            return f"contract % for {d.supply_names[i]} in {period(p)}"
        if k < self.off_s_inj:
            p, i = divmod(k - self.off_s_wd, N)
            return f"storage withdrawal at {d.supply_names[i]} in {period(p)}"
        if k < self.off_s_end:
            p, i = divmod(k - self.off_s_inj, N)
            return f"storage injection at {d.supply_names[i]} in {period(p)}"
        if k < self.off_d_wd:
            p, i = divmod(k - self.off_s_end, N)
            return f"ending inventory at {d.supply_names[i]} in {period(p)}"
        if k < self.off_d_inj:
            p, j = divmod(k - self.off_d_wd, M)
            return f"storage withdrawal at {d.demand_names[j]} in {period(p)}"
        if k < self.off_d_end:
            p, j = divmod(k - self.off_d_inj, M)
            return f"storage injection at {d.demand_names[j]} in {period(p)}"
        if k < self.off_short:
            p, j = divmod(k - self.off_d_end, M)
            return f"ending inventory at {d.demand_names[j]} in {period(p)}"
        if k < self.off_surplus:
            p, j = divmod(k - self.off_short, M)
            return f"unmet demand at {d.demand_names[j]} in {period(p)}"
        p, i = divmod(k - self.off_surplus, N)
        return f"unplaceable volume at {d.supply_names[i]} in {period(p)}"

    def equality_triplets(
        self,
    ) -> Tuple[List[int], List[int], List[float], List[float], int]:
        """
        Build the equality system in plain-Python COO form.

        Returns ``(rows, cols, vals, b, n_rows)``. This is kept scipy-free on
        purpose, so the PuLP/CBC fallback still works on machines that only
        have PuLP installed. Triplets come out in non-decreasing row order,
        which lets whoever consumes them walk row by row without needing to
        sort anything first.
        """
        P, N, M = self.P, self.N, self.M
        rows: List[int] = []
        cols: List[int] = []
        vals: List[float] = []
        b: List[float] = []
        row = 0

        def put(c: int, v: float) -> None:
            # The float()/int() calls actually matter here: multiplying a
            # PuLP variable by a raw numpy scalar routes through numpy's
            # dispatch machinery and is noticeably slower than using plain
            # Python floats.
            rows.append(row)
            cols.append(int(c))
            vals.append(float(v))

        def elastic() -> None:
            # A +/- slack pair on the current row. Heavily penalised, so the
            # solver only touches it when the row genuinely cannot be met, at
            # which point the non-zero slack tells us exactly where the
            # workbook's data contradicts itself.
            if self.allow_elastic:
                put(self.off_epos + row, 1.0)
                put(self.off_eneg + row, -1.0)

        # (A) supply balance - what leaves a supply point has to equal what
        # arrived there: contract*perc + wd - sum_j ship - inj = 0
        for p in range(P):
            for i in range(N):
                put(self.iperc(p, i), self.d.contract[p, i])
                put(self.is_wd(p, i), 1.0)
                put(self.is_inj(p, i), -1.0)
                for j in range(M):
                    put(self.iship(p, i, j), -1.0)
                if self.allow_surplus:
                    # A sink for volume a contract minimum forces us to lift
                    # when demand and storage cannot absorb it.
                    put(self.isurplus(p, i), -1.0)
                if self.allow_unsourced:
                    # A source for the opposite case: per-lane minimum
                    # shipments demanding more than the contract ceiling allows.
                    put(self.iunsourced(p, i), 1.0)
                elastic()
                b.append(0.0)
                row += 1

        # (B) supply-side inventory chain, linking each period's ending
        # inventory to the one before it: end[p] - end[p-1] + wd[p] - inj[p] = 0
        for p in range(P):
            for i in range(N):
                put(self.is_end(p, i), 1.0)
                put(self.is_wd(p, i), 1.0)
                put(self.is_inj(p, i), -1.0)
                elastic()
                if p == 0:
                    b.append(float(self.d.supply_inv_init[i]))
                else:
                    put(self.is_end(p - 1, i), -1.0)
                    b.append(0.0)
                row += 1

        # (C) demand balance - what arrives at a demand point (plus whatever
        # comes out of storage there) has to cover demand:
        # sum_i ship + wd - inj (+ shortfall) = demand
        for p in range(P):
            for j in range(M):
                for i in range(N):
                    put(self.iship(p, i, j), 1.0)
                put(self.id_wd(p, j), 1.0)
                put(self.id_inj(p, j), -1.0)
                if self.allow_shortfall:
                    put(self.ishort(p, j), 1.0)
                if self.allow_excess:
                    # A sink for volume that per-lane minimum shipments force
                    # into this demand point beyond what it can use or store.
                    put(self.iexcess(p, j), -1.0)
                elastic()
                b.append(float(self.d.demand[p, j]))
                row += 1

        # (D) demand-side inventory chain, same shape as (B) but at the
        # demand points: end[p] - end[p-1] + wd[p] - inj[p] = 0
        for p in range(P):
            for j in range(M):
                put(self.id_end(p, j), 1.0)
                put(self.id_wd(p, j), 1.0)
                put(self.id_inj(p, j), -1.0)
                elastic()
                if p == 0:
                    b.append(float(self.d.demand_inv_init[j]))
                else:
                    put(self.id_end(p - 1, j), -1.0)
                    b.append(0.0)
                row += 1

        return rows, cols, vals, b, row

    def equalities(self):
        """``(A_eq, b_eq)`` with ``A_eq`` as a scipy CSR matrix."""
        from scipy import sparse

        rows, cols, vals, b, n_rows = self.equality_triplets()
        A = sparse.coo_matrix(
            (vals, (rows, cols)), shape=(n_rows, self.n_vars), dtype=float
        ).tocsr()
        return A, np.asarray(b, dtype=float)

    # -- solve ------------------------------------------------------------- #

    def solve(self, time_limit: Optional[float] = None) -> Solution:
        import time

        t0 = time.perf_counter()

        # A crossed bound is infeasible by construction, so say exactly why up
        # front rather than letting the solver blow up with an opaque
        # ValueError later.
        crossed = self.crossed_bounds()
        if crossed:
            return Solution(
                data=self.d,
                status="infeasible",
                message=(
                    "Contradictory input bounds (minimum above maximum):\n  "
                    + "\n  ".join(crossed)
                ),
                solver="pre-solve check",
                solve_seconds=time.perf_counter() - t0,
            )

        try:
            from scipy.optimize import linprog
        except ImportError:
            return self._solve_with_pulp(t0, time_limit)

        c = self.objective()
        lo, hi = self.bounds()
        A_eq, b_eq = self.equalities()

        options = {"presolve": True}
        if time_limit:
            options["time_limit"] = float(time_limit)

        res = linprog(
            c,
            A_eq=A_eq,
            b_eq=b_eq,
            bounds=np.column_stack([lo, hi]),
            method="highs",
            options=options,
        )
        elapsed = time.perf_counter() - t0

        status_map = {
            0: "optimal",
            1: "iteration/time limit",
            2: "infeasible",
            3: "unbounded",
            4: "numerical failure",
        }
        status = status_map.get(res.status, f"status {res.status}")
        if res.status != 0 or res.x is None:
            return Solution(
                data=self.d,
                status="infeasible" if res.status == 2 else status,
                message=str(res.message or status),
                solver="scipy/HiGHS",
                solve_seconds=elapsed,
            )

        sol = self._unpack(res.x, float(res.fun), "scipy/HiGHS", elapsed)
        try:
            sol.duals = {"upper": np.asarray(res.upper.marginals, dtype=float),
                         "eq": np.asarray(res.eqlin.marginals, dtype=float)}
            sol.lp = self
        except (AttributeError, TypeError):
            pass                                   # older scipy: no duals, feature just hides
        return sol

    def _solve_with_pulp(self, t0: float, time_limit: Optional[float]) -> Solution:
        """
        Fallback path using PuLP + CBC.

        Slower than HiGHS, but it's the same solver OpenSolver was using
        under the hood, which makes it a handy way to cross-check a result
        that looks suspicious.
        """
        import time

        try:
            import pulp
        except ImportError:
            return Solution(
                data=self.d,
                status="error",
                message=(
                    "No LP solver available. Install one of:\n"
                    "    python -m pip install scipy      (preferred, HiGHS)\n"
                    "    python -m pip install pulp       (CBC)"
                ),
                solve_seconds=time.perf_counter() - t0,
            )

        c = self.objective()
        lo, hi = self.bounds()
        rows, cols, vals, b_eq, n_rows = self.equality_triplets()

        prob = pulp.LpProblem(f"{self.d.area}_supply", pulp.LpMinimize)
        xs = [
            pulp.LpVariable(
                f"x{k}",
                lowBound=float(lo[k]),
                upBound=None if math.isinf(hi[k]) else float(hi[k]),
            )
            for k in range(self.n_vars)
        ]
        prob += pulp.lpSum(float(c[k]) * xs[k] for k in np.nonzero(c)[0])

        # Triplets arrive already sorted by row, so we can just accumulate
        # each row's terms and flush them into a constraint as we go.
        start = 0
        total = len(rows)
        for r in range(n_rows):
            end = start
            while end < total and rows[end] == r:
                end += 1
            prob += (
                pulp.lpSum(vals[k] * xs[cols[k]] for k in range(start, end))
                == float(b_eq[r])
            )
            start = end

        solver = pulp.PULP_CBC_CMD(msg=False, timeLimit=time_limit)
        prob.solve(solver)
        elapsed = time.perf_counter() - t0

        state = pulp.LpStatus[prob.status]
        if state != "Optimal":
            return Solution(
                data=self.d,
                status="infeasible" if state == "Infeasible" else state.lower(),
                message=f"CBC returned {state}",
                solver="pulp/CBC",
                solve_seconds=elapsed,
            )

        x = np.array([v.value() or 0.0 for v in xs], dtype=float)
        return self._unpack(x, float(pulp.value(prob.objective)), "pulp/CBC", elapsed)

    def _unpack(self, x: np.ndarray, obj: float, solver: str, elapsed: float) -> Solution:
        P, N, M = self.P, self.N, self.M

        def block(off: int, shape: Tuple[int, ...]) -> np.ndarray:
            size = int(np.prod(shape))
            return x[off:off + size].reshape(shape)

        # Mop up the -1e-13-ish noise an interior-point or simplex solver
        # tends to leave behind, so nobody stares at "-0.0000000001" and
        # wonders if something's wrong.
        def clean(a: np.ndarray) -> np.ndarray:
            a = np.asarray(a, dtype=float).copy()
            a[np.abs(a) < 1e-9] = 0.0
            return a

        sol = Solution(
            data=self.d,
            status="optimal",
            objective=obj,
            ship=clean(block(self.off_ship, (P, N, M))),
            perc=clean(block(self.off_perc, (P, N))),
            supply_wd=clean(block(self.off_s_wd, (P, N))),
            supply_inj=clean(block(self.off_s_inj, (P, N))),
            supply_end=clean(block(self.off_s_end, (P, N))),
            demand_wd=clean(block(self.off_d_wd, (P, M))),
            demand_inj=clean(block(self.off_d_inj, (P, M))),
            demand_end=clean(block(self.off_d_end, (P, M))),
            solver=solver,
            solve_seconds=elapsed,
        )
        if self.allow_shortfall:
            sol.shortfall = clean(block(self.off_short, (P, M)))
        if self.allow_excess:
            sol.excess = clean(block(self.off_excess, (P, M)))
        if self.allow_surplus:
            sol.surplus = clean(block(self.off_surplus, (P, N)))
        if self.allow_unsourced:
            sol.unsourced = clean(block(self.off_unsourced, (P, N)))
        if self.allow_elastic:
            epos = clean(x[self.off_epos:self.off_epos + self.n_rows])
            eneg = clean(x[self.off_eneg:self.off_eneg + self.n_rows])
            sol.violations = self.decode_violations(epos - eneg)
        return sol

    def decode_violations(self, viol: np.ndarray, limit: int = 25) -> List[str]:
        """
        Turn non-zero elastic slacks into plain-language findings.

        A non-zero slack means that row could not be satisfied by any
        combination of the allowed decisions, so it names the exact constraint,
        location and period that the workbook's data breaks.
        """
        P, N, M = self.P, self.N, self.M
        d = self.d
        found: List[Tuple[float, str]] = []

        def period(p: int) -> str:
            return d.period_labels[p] if 0 <= p < len(d.period_labels) else f"P{p+1}"

        for r in np.flatnonzero(np.abs(viol) > 1e-6):
            v = float(viol[r])
            if r < self.row_supply_inventory:
                p, i = divmod(int(r) - self.row_supply_balance, N)
                what = ("had to invent" if v > 0 else "had to discard")
                found.append((abs(v), (
                    f"{FAM_SUPPLY_BALANCE}: {what} {abs(v):,.1f} kgal at "
                    f"{d.supply_names[i]} in {period(p)}"
                )))
            elif r < self.row_demand_balance:
                p, i = divmod(int(r) - self.row_supply_inventory, N)
                found.append((abs(v), (
                    f"{FAM_SUPPLY_INVENTORY}: storage balance off by {v:,.1f} kgal "
                    f"at {d.supply_names[i]} in {period(p)}"
                )))
            elif r < self.row_demand_inventory:
                p, j = divmod(int(r) - self.row_demand_balance, M)
                what = ("unmet demand of" if v > 0
                        else "forced over-delivery of")
                found.append((abs(v), (
                    f"{FAM_DEMAND_BALANCE}: {what} {abs(v):,.1f} kgal at "
                    f"{d.demand_names[j]} in {period(p)}"
                )))
            else:
                p, j = divmod(int(r) - self.row_demand_inventory, M)
                found.append((abs(v), (
                    f"{FAM_DEMAND_INVENTORY}: storage balance off by {v:,.1f} kgal "
                    f"at {d.demand_names[j]} in {period(p)}"
                )))

        found.sort(key=lambda t: -t[0])
        out = [msg for _, msg in found[:limit]]
        if len(found) > limit:
            out.append(f"...and {len(found) - limit} more")
        return out


def solve_area(data: ModelData, allow_shortfall: bool = False,
               time_limit: Optional[float] = None,
               auto_relax: bool = True) -> Solution:
    """
    Build and solve the LP for one area.

    The first attempt is completely faithful to the spreadsheet: no slack of
    any kind, so if it solves, the answer is one Excel could have produced.

    If that is infeasible, the cause is almost always the workbook's own input
    data rather than anything structural, and the spreadsheet's built-in
    "Feasible Check" cells usually already say so:

      * "Max Too Low"  - contract volume x maximum lift % cannot cover demand,
                         so some demand is simply unservable.
      * "Min Too High" - contract MINIMUMS force more volume than demand can
                         absorb, and with no storage capacity there is nowhere
                         to put the excess.

    Excel's solver fails on these too, which is why the saved workbooks often
    contain a stale solution. Returning nothing at all is the least useful
    response though, so the second attempt re-solves with two penalised escape
    valves - unmet demand and unplaceable volume - which makes the model always
    solvable and, more importantly, quantifies exactly how far off the inputs
    are and where. The result gets flagged as relaxed so nobody mistakes it for
    a clean plan.
    """
    builder = LPBuilder(data, allow_shortfall=allow_shortfall)
    LOG.debug(
        "%s: %d variables, %d equality rows",
        data.area, builder.n_vars,
        builder.P * (2 * builder.N + 2 * builder.M),
    )
    sol = builder.solve(time_limit=time_limit)

    if sol.ok:
        extra = ""
        if sol.unmet_total > 1e-6:
            extra = f"  UNMET {sol.unmet_total:,.1f} kgal"
        LOG.info(
            "%-16s optimal  cost $%s%s  (%s, %.2fs)",
            data.area, f"{sol.objective:,.0f}", extra, sol.solver, sol.solve_seconds,
        )
        return sol

    if not auto_relax:
        LOG.warning("%-16s %s: %s", data.area, sol.status, sol.message)
        return sol

    LOG.warning(
        "%-16s %s with the workbook's own limits; retrying with unmet-demand "
        "and surplus allowances to find out why", data.area, sol.status,
    )
    strict_message = sol.message

    # Stage 2 makes both balance equations elastic in all four directions, but
    # leaves the inventory chains rigid - those are accounting identities and
    # bending them would produce nonsense rather than insight.
    relaxed = LPBuilder(
        data, allow_shortfall=True, allow_surplus=True,
        allow_excess=True, allow_unsourced=True,
    ).solve(time_limit=time_limit)

    if not relaxed.ok:
        # Unmet demand and surplus were not enough, so the contradiction is
        # somewhere other than aggregate volume - most often per-lane minimum
        # shipments that force more into a demand point than it can take. Fall
        # back to a fully elastic (phase-1) solve, which is always feasible and
        # names the offending constraints outright.
        LOG.warning(
            "%-16s still infeasible; running an elastic solve to localise it",
            data.area,
        )
        elastic = LPBuilder(
            data, allow_shortfall=True, allow_surplus=True,
            allow_excess=True, allow_unsourced=True, allow_elastic=True,
        ).solve(time_limit=time_limit)

        if not elastic.ok:
            LOG.error("%-16s elastic solve also failed: %s",
                      data.area, elastic.message)
            sol.message = (
                f"{strict_message}\nRelaxed retry failed: {relaxed.message}"
                f"\nElastic retry failed: {elastic.message}"
            )
            return sol

        elastic.relaxed = True
        elastic.relax_note = (
            "This area could only be solved by bending the balance equations "
            "themselves, which means the inputs are contradictory beyond simply "
            "being short of supply. The specific constraints that could not be "
            "satisfied are listed below; treat the plan as indicative only."
        )
        LOG.warning("%-16s ELASTIC: %d constraint(s) could not be satisfied",
                    data.area, len(elastic.violations))
        for v in elastic.violations[:5]:
            LOG.warning("%-16s   %s", data.area, v)
        return elastic

    relaxed.relaxed = True
    bits = []
    if relaxed.unmet_total > 1e-6:
        bits.append(f"{relaxed.unmet_total:,.1f} kgal of demand cannot be served "
                    f"from the contracted supply ('Max Too Low')")
    if relaxed.surplus_total > 1e-6:
        bits.append(f"{relaxed.surplus_total:,.1f} kgal is forced up by contract "
                    f"minimums with no demand or storage to absorb it "
                    f"('Min Too High')")
    if relaxed.excess_total > 1e-6:
        bits.append(f"{relaxed.excess_total:,.1f} kgal is forced into demand points "
                    f"by minimum transport quantities beyond what they can use "
                    f"or store")
    if relaxed.unsourced_total > 1e-6:
        bits.append(f"{relaxed.unsourced_total:,.1f} kgal would have to be shipped "
                    f"without any contracted volume behind it")
    relaxed.relax_note = (
        "The workbook's limits as saved have no feasible solution. "
        + ("; ".join(bits) if bits else "Cause not isolated.")
        + ". Figures below are the best plan achievable once those gaps are "
          "allowed for - correct the workbook and re-run for a clean plan."
    )
    LOG.warning("%-16s RELAXED: %s", data.area, "; ".join(bits) or "cause unclear")
    return relaxed


# --------------------------------------------------------------------------- #
# Reading more out of a solved model.
# --------------------------------------------------------------------------- #

#: Dual values above this ($/kgal) come from the penalty on unmet demand or
#: surplus in a relaxed solve, not from real prices - flagged as such.
PENALTY_DUAL = 20000.0


def marginal_values(sol: Solution, top: int = 25) -> Dict[str, List[Dict[str, object]]]:
    """
    What one more gallon is worth, read from the LP's dual values.

    Values are $/kgal (the model's units; the app shows $/gal). Negative =
    the season gets cheaper by that much per extra kgal.

      contracts  raise a contract's maximum lift by 1 kgal in a month
      lanes      allow 1 more kgal on a lane that is at its maximum - including
                 lanes closed with a blank freight rate ("opening it would save")
      storage    1 more kgal of storage at a supply point or plant
      demand     what 1 more kgal of demand at a plant costs (its marginal
                 delivered cost)

    Each entry lists the months where it applies and the average value over
    those months. Only meaningful for an optimal solve; a relaxed solve's
    values can include the unmet-demand penalty and are flagged "penalty".
    """
    out: Dict[str, List[Dict[str, object]]] = {"contracts": [], "lanes": [], "storage": [], "demand": []}
    if not sol.ok or not sol.duals or sol.lp is None:
        return out
    lp, d = sol.lp, sol.data
    up, eq = sol.duals["upper"], sol.duals["eq"]
    P, N, M = lp.P, lp.N, lp.M
    tol = 1e-6

    def entry(kind, months, vals, **kw):
        v = float(np.mean(vals))
        return {"kind": kind, "months": [int(p) for p in months], "value": v,
                "best": float(np.min(vals)), "penalty": bool(abs(v) > PENALTY_DUAL), **kw}

    mu = up[lp.off_perc:lp.off_perc + P * N].reshape(P, N)
    with np.errstate(divide="ignore", invalid="ignore"):
        per_kgal = np.where(d.contract > 1e-9, mu / np.where(d.contract > 1e-9, d.contract, 1.0), 0.0)
    for i in range(N):
        months = [p for p in range(P) if per_kgal[p, i] < -tol]
        if months:
            out["contracts"].append(entry("contract", months, per_kgal[months, i], i=i))

    ship_mu = up[lp.off_ship:lp.off_ship + lp.n_ship].reshape(P, N, M)
    closed = d.closed_lane_mask()
    for i, j in zip(*np.nonzero((ship_mu < -tol).any(axis=0))):
        months = [p for p in range(P) if ship_mu[p, i, j] < -tol]
        out["lanes"].append(entry("lane", months, ship_mu[months, i, j], i=int(i), j=int(j),
                                  closed=bool(closed[months, i, j].any())))

    for kind, off, n in (("supply storage", lp.off_s_end, N), ("plant storage", lp.off_d_end, M)):
        smu = up[off:off + P * n].reshape(P, n)
        for k in range(n):
            months = [p for p in range(P) if smu[p, k] < -tol]
            if months:
                key = "i" if kind == "supply storage" else "j"
                out["storage"].append(entry(kind, months, smu[months, k], **{key: k}))

    lam = eq[lp.row_demand_balance:lp.row_demand_balance + P * M].reshape(P, M)
    for j in range(M):
        months = [p for p in range(P) if d.demand[p, j] > 1e-9]
        if months:
            e = entry("demand", months, lam[months, j], j=j)
            e["best"] = float(np.max(lam[months, j]))          # dearest month
            e["by_month"] = [float(x) for x in lam[:, j]]
            out["demand"].append(e)

    for k in ("contracts", "lanes", "storage"):
        out[k] = sorted(out[k], key=lambda e: e["value"])[:top]
    out["demand"] = sorted(out["demand"], key=lambda e: -e["value"])[:top]
    return out


def contract_use(sol: Solution) -> List[Dict[str, object]]:
    """
    How hard each contract is worked: season volume against contract, and how
    many months it sits at its minimum, at its maximum, or in between. A
    contract at its maximum all winter is a candidate for more volume; one at
    its minimum all season is a candidate for less.
    """
    if not sol.ok or sol.perc is None:
        return []
    d = sol.data
    out = []
    for i in range(d.n_supply):
        live = d.contract[:, i] > 1e-9
        if not live.any():
            continue
        perc = sol.perc[:, i]
        at_max = live & (perc >= d.perc_max[:, i] - 1e-6)
        at_min = live & ~at_max & (perc <= d.perc_min[:, i] + 1e-6)
        state = ["" if not live[p] else "max" if at_max[p] else "min" if at_min[p] else "mid"
                 for p in range(d.n_periods)]
        bought = float((d.contract[:, i] * perc).sum())
        contracted = float(d.contract[:, i].sum())
        out.append({"i": i, "bought": bought, "contract": contracted,
                    "share": bought / contracted if contracted > 1e-9 else 0.0,
                    "at_max": int(at_max.sum()), "at_min": int(at_min.sum()),
                    "mid": int((live & ~at_max & ~at_min).sum()), "months": state,
                    "lift": [float(x) for x in perc]})
    return out


# --------------------------------------------------------------------------- #
# Discovery - figuring out what regions, areas and workbooks actually exist.
# --------------------------------------------------------------------------- #


@dataclass
class AreaFile:
    region: str
    area: str
    variant: str          # "Working" | "Baseline"
    path: str
    candidates: List[str] = field(default_factory=list)


def _classify(filename: str) -> Tuple[str, str, int, int]:
    """
    Split a workbook name into ``(area, variant, version, sequence)``.

    ``"S LA 26-27 v2 Working.xlsm"`` -> ``("S LA 26-27", "Working", 2, 0)``
    ``"West Rail District 26-27 Working1.xlsm"`` -> ``(..., "Working", 0, 1)``
    """
    stem = os.path.splitext(filename)[0]
    m = _VARIANT_RE.match(stem)
    if not m:
        return stem.strip(), "", 0, 0
    return (
        m.group("stem").strip(),
        m.group("variant").capitalize(),
        int(m.group("vers") or 0),
        int(m.group("seq") or 0),
    )


def normalise_region(name: str) -> str:
    """
    Fold a region name down to a comparison key: lower-case, letters and
    digits only.

    This is what lets ``"North East"``, ``"north-east"`` and ``"Northeast"``
    all land on the ``Northeast`` folder, and ``"New England"`` land on
    ``New England`` regardless of how it was typed.
    """
    return re.sub(r"[^a-z0-9]+", "", (name or "").lower())


#: Extra spellings accepted for a region, keyed by the folded folder name.
#: Only needed where the actual words differ - plain spacing, case and
#: punctuation differences are already handled by normalise_region() above.
REGION_ALIASES: Dict[str, Tuple[str, ...]] = {
    "pnw": ("pacificnorthwest", "pacnw", "northwest"),
    "rockies": ("rocky", "rockymountain", "rockymountains", "rockys"),
    "alaska": ("ak",),
    "newengland": ("neweng",),
    "midwest": ("mw",),
}


def _alias_lookup() -> Dict[str, str]:
    """``{alias key: canonical folded name}``."""
    out: Dict[str, str] = {}
    for canonical, aliases in REGION_ALIASES.items():
        for a in aliases:
            out[normalise_region(a)] = canonical
    return out


def resolve_region(name: str, available: Sequence[str]) -> Tuple[Optional[str], str]:
    """
    Match one user-supplied region name against the folders that actually
    exist.

    Returns ``(folder_name, "")`` on success, or ``(None, reason)`` when it
    can't figure out what was meant.

    It tries these in order, most specific first:
      1. exact match after folding      -- "north east"  -> Northeast
      2. known alias                    -- "pacific nw"  -> PNW
      3. unique prefix                  -- "rock"        -> Rockies
      4. unique substring               -- "england"     -> New England

    Exact matches always win over prefixes, so ``"South"`` picks South and
    never turns ambiguous against Southeast.
    """
    key = normalise_region(name)
    if not key:
        return None, "empty region name"

    folded = {normalise_region(r): r for r in available}

    if key in folded:
        return folded[key], ""

    alias = _alias_lookup().get(key)
    if alias and alias in folded:
        return folded[alias], ""

    for predicate, label in (
        (lambda k: k.startswith(key), "prefix"),
        (lambda k: key in k, "substring"),
    ):
        hits = [orig for k, orig in folded.items() if predicate(k)]
        if len(hits) == 1:
            return hits[0], ""
        if len(hits) > 1:
            return None, (
                f"{name!r} is ambiguous as a {label} - matches "
                f"{', '.join(sorted(hits))}"
            )

    return None, f"{name!r} does not match any region"


def expand_selector(
    name: str, available: Sequence[str], groups: Dict[str, List[str]]
) -> Tuple[Optional[List[str]], str]:
    """
    Resolve one name the user typed into a list of region folders.

    A name can be either a region or a saved group, so managers can type
    whichever they think in without having to remember which flag it belongs
    to. Precedence matters here: an exact region name always wins, so a group
    called "west" can never shadow the West folder. Only if nothing matches a
    region exactly do we look at group names, and fuzzy region matching
    (alias/prefix/substring) is tried last.

    Returns ``(folders, "")`` on success or ``(None, reason)`` on failure.
    """
    key = normalise_region(name)
    folded = {normalise_region(r): r for r in available}

    if key in folded:
        return [folded[key]], ""

    group = next((g for g in groups if normalise_region(g) == key), None)
    if group is not None:
        folders: List[str] = []
        for member in groups[group]:
            folder, reason = resolve_region(member, available)
            if folder is None:
                return None, f"group {group!r} lists an unusable region - {reason}"
            if folder not in folders:
                folders.append(folder)
        return folders, ""

    folder, reason = resolve_region(name, available)
    if folder is None:
        extra = ""
        if groups:
            extra = (f". It is not a saved group either - defined groups are: "
                     f"{', '.join(sorted(groups))}")
        return None, reason + extra
    return [folder], ""


def split_region_args(values: Optional[Sequence[str]]) -> List[str]:
    """
    Flatten repeated flags and delimiter-separated lists into one plain list
    of names.

    All of these end up meaning the same thing::

        --region Rail --region Rockies
        --region Rail,Rockies
        --region "Rail; Rockies"
    """
    out: List[str] = []
    for raw in values or []:
        for piece in re.split(r"[,;/|]+", raw):
            piece = piece.strip()
            if piece:
                out.append(piece)
    return out


def load_groups(path: str) -> Dict[str, List[str]]:
    """
    Read the optional manager/group map: ``{group name: [region, ...]}``.

    It's plain JSON on purpose, so a manager can open it up and edit it
    without touching any Python::

        {
          "groups": {
            "jdoe":  ["Rail", "Rockies"],
            "asmith": ["South", "Southeast"]
          }
        }

    A missing or malformed file isn't fatal - it just means --group won't be
    available for this run, and the reason gets logged so it's not a mystery.
    """
    if not os.path.isfile(path):
        return {}
    try:
        import json

        with open(path, "r", encoding="utf-8-sig") as fh:
            blob = json.load(fh)
    except Exception as exc:  # noqa: BLE001 - a bad config file should never crash a run
        LOG.warning("could not read groups file %s (%s); --group unavailable", path, exc)
        return {}

    raw = blob.get("groups", blob) if isinstance(blob, dict) else {}
    groups: Dict[str, List[str]] = {}
    if isinstance(raw, dict):
        for key, members in raw.items():
            if isinstance(members, str):
                members = split_region_args([members])
            if isinstance(members, (list, tuple)):
                names = [str(m).strip() for m in members if str(m).strip()]
                if names:
                    groups[str(key)] = names
    if not groups:
        LOG.warning("groups file %s contains no usable groups", path)
    return groups


def discover(root: str, variant: str = "Working") -> Dict[str, List[AreaFile]]:
    """
    Walk *root* and return ``{region: [AreaFile, ...]}``.

    Each immediate sub-directory counts as a region. Within a region,
    workbooks get grouped by area, and exactly one file per area is chosen:
    the requested *variant*, preferring the plain name over a ``v2`` /
    ``Working1`` style sibling so a stray experiment never quietly becomes
    the official run.
    """
    if not os.path.isdir(root):
        raise SystemExit(f"model root does not exist: {root}")

    regions: Dict[str, List[AreaFile]] = {}
    for entry in sorted(os.listdir(root)):
        region_dir = os.path.join(root, entry)
        if not os.path.isdir(region_dir):
            continue
        if entry.lower() in NON_REGION_DIRS or entry.startswith("."):
            continue

        # Maps area name -> list of (version, sequence, filename) candidates.
        buckets: Dict[str, List[Tuple[int, int, str]]] = {}
        for fn in sorted(os.listdir(region_dir)):
            if not fn.lower().endswith(".xlsm"):
                continue
            if any(pat.search(fn) for pat in SKIP_FILE_PATTERNS):
                continue
            area, var, vers, seq = _classify(fn)
            if var.lower() != variant.lower():
                continue
            buckets.setdefault(area, []).append((vers, seq, fn))

        areas: List[AreaFile] = []
        for area, options in sorted(buckets.items()):
            options.sort()  # plain name (0, 0) sorts first
            _, _, chosen = options[0]
            areas.append(
                AreaFile(
                    region=entry,
                    area=area,
                    variant=variant,
                    path=os.path.join(region_dir, chosen),
                    candidates=[fn for _, _, fn in options],
                )
            )
            if len(options) > 1:
                LOG.info(
                    "%s / %s: %d candidates, using %r",
                    entry, area, len(options), chosen,
                )
        if areas:
            regions[entry] = areas
        else:
            LOG.debug("no %s workbooks in region %r", variant, entry)
    return regions


# --------------------------------------------------------------------------- #
# Writing the output workbook.
# --------------------------------------------------------------------------- #

_HDR_FILL = PatternFill("solid", fgColor="1F3864")
_HDR_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=13)
_WARN_FILL = PatternFill("solid", fgColor="FFF2CC")

_MONEY = '#,##0'
_VOL = '#,##0'            # gallons
_RATE = '$0.0000'         # $ per gallon
_PCT = '0.0%'

#: The model works in the workbooks' own units - thousands of gallons (kgal)
#: and $ per thousand gallons - because that is what every input cell holds.
#: Reports show plain GALLONS and $ PER GALLON instead, converted only at the
#: point of output, so no result changes: 300 kgal is shown as 300,000 gal
#: and $650.00/kgal as $0.6500/gal. Dollar totals are the same either way.
GAL_PER_KGAL = 1000.0

_KGAL_NUM_RE = re.compile(r"(-?\d[\d,]*(?:\.\d+)?)(\s*)kgal\b")
_PER_KGAL_RE = re.compile(r"\$(-?\d[\d,]*(?:\.\d+)?)\s*/\s*kgal\b")


def to_gal(v):
    """kgal -> gallons (None passes through)."""
    return None if v is None else float(v) * GAL_PER_KGAL


def to_per_gal(v):
    """$/kgal -> $/gal (None passes through)."""
    return None if v is None else float(v) / GAL_PER_KGAL


def gal_text(text: str) -> str:
    """
    Rewrite a message written in model units into gallons for people.

        '250.5 kgal of demand'      -> '250,500 gal of demand'
        'priced at $1,000/kgal'     -> 'priced at $1.0000/gal'
        'Volume (kgal)'             -> 'Volume (gal)'
    """
    if not text:
        return text

    def per(m):
        return f"${float(m.group(1).replace(',', '')) / GAL_PER_KGAL:,.4f}/gal"

    def vol(m):
        return f"{float(m.group(1).replace(',', '')) * GAL_PER_KGAL:,.0f}{m.group(2)}gal"

    text = _PER_KGAL_RE.sub(per, text)
    text = _KGAL_NUM_RE.sub(vol, text)
    return (text.replace("$/kgal", "$/gal").replace("kgal", "gal")
                .replace("thousands of gallons", "gallons"))


#: How each input converts for display: volumes x1000, prices /1000,
#: lift percentages unchanged.
_VOLUME_FIELDS = {"contract", "demand", "ship_min", "ship_max", "supply_inv_init",
                  "supply_inv_cap", "supply_inv_min", "demand_inv_init", "demand_inv_cap",
                  "demand_inv_min", "store_max", "store_min", "store_init"}
_PRICE_FIELDS = {"hub_price", "hub_premium", "price", "trans_cost", "diff", "freight"}


def gal_series_text(text: str, field: str) -> str:
    """An audit before/after string ('652.5' or '1 2 3') converted for *field*."""
    if not text or text == "-":
        return text
    if field in _VOLUME_FIELDS:
        f = GAL_PER_KGAL
    elif field in _PRICE_FIELDS:
        f = 1.0 / GAL_PER_KGAL
    else:
        return gal_text(text)
    out = []
    for tok in text.split():
        try:
            out.append(f"{float(tok) * f:g}")
        except ValueError:
            out.append(tok)
    return " ".join(out)


def slugify(text: str) -> str:
    """Collapse arbitrary text into a safe filename fragment."""
    return re.sub(r"[^A-Za-z0-9]+", "_", text or "").strip("_")


def region_filename(region: str, when: _dt.date, contract_label: str = "",
                    scenario: str = "") -> str:
    """
    ``West`` -> ``West_Model_09_17_26.xlsx``.

    With a scenario from an override deck the name carries it too, e.g.
    ``West_Model_WideOpenFirstPass_09_17_26.xlsx``. That matters because two
    scenarios run on the same day must not overwrite each other.

    Spaces turn into underscores so the name stays safe to use in a shell
    command or a URL without anyone having to quote it.
    """
    slug = slugify(region) or "Region"
    stamp = when.strftime("%m_%d_%y")
    parts = [slug, "Model"]
    if contract_label:
        parts.append(contract_label)
    if scenario:
        # CamelCase the scenario so it stays readable but compact in a filename.
        words = [w for w in re.split(r"[^A-Za-z0-9]+", scenario) if w]
        camel = "".join(w[:1].upper() + w[1:] for w in words)
        if camel:
            parts.append(camel)
    parts.append(stamp)
    return "_".join(parts) + ".xlsx"


class RegionWorkbookWriter:
    """
    Writes one .xlsx per region.

    Sheets
        Summary          region roll-up, per-area cost/volume, warnings
        Supply Plan      per area / supply point / period lift decision
        Shipments        the non-zero lanes, long format
        Demand Coverage  demand vs delivered vs inventory movement
        Supply Inventory begin / withdraw / inject / end / capacity
        Demand Inventory same, at the demand points
        Diagnostics      data-quality notes and feasibility checks
    """

    #: Lanes carrying less than this (in thousand gallons) get dropped from
    #: the Shipments sheet - it's solver rounding noise, not a real movement.
    SHIP_EPS = 1e-6

    def __init__(self, region: str, solutions: Sequence[Solution],
                 run_date: _dt.date, source_note: str = "",
                 scenario: str = "", compare_saved: bool = False):
        self.region = region
        self.solutions = list(solutions)
        self.run_date = run_date
        self.source_note = source_note
        #: Name from the override deck's [scenario] block, "" when no deck is
        #: in play. Shown on every sheet so a printed copy is traceable.
        self.scenario = scenario
        #: Whether to include the saved-vs-resolved cost comparison.
        self.compare_saved = compare_saved

    # -- helpers ----------------------------------------------------------- #

    @staticmethod
    def _header(ws, headers: Sequence[str], row: int = 1, freeze: bool = True) -> None:
        for c, text in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=c, value=text)
            cell.fill = _HDR_FILL
            cell.font = _HDR_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if freeze:
            ws.freeze_panes = ws.cell(row=row + 1, column=1)

    @staticmethod
    def _widths(ws, widths: Sequence[int]) -> None:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    @staticmethod
    def _fmt(ws, first_row: int, last_row: int, spec: Dict[int, str]) -> None:
        if last_row < first_row:
            return
        for col, fmt in spec.items():
            for r in range(first_row, last_row + 1):
                ws.cell(row=r, column=col).number_format = fmt

    # -- sheets ------------------------------------------------------------ #

    def _summary(self, wb) -> None:
        ws = wb.create_sheet("Summary")
        ws["A1"] = f"{self.region} - Propane Supply Optimization"
        ws["A1"].font = _TITLE_FONT
        ws["A2"] = (
            "Volumes in gallons. "
            "Prices, differentials and freight in $ per gallon."
        )
        ws["A2"].font = Font(bold=True)
        ws["A3"] = f"Generated {self.run_date:%Y-%m-%d} by model.py"
        if self.scenario:
            ws["A4"] = f"Scenario: {self.scenario}"
            ws["A4"].font = Font(bold=True)
        if self.source_note:
            ws["A5"] = self.source_note

        row = 7
        ws.cell(row=row, column=1, value="Area results").font = Font(bold=True)
        row += 1
        headers = [
            "Area", "Status", "Supply pts", "Demand pts",
            "Demand (gal)", "Purchased (gal)", "Shipped (gal)",
            "Unmet (gal)", "Surplus (gal)", "Other slack (gal)",
            "High freight >= $1/gal (gal)", "High freight >= $1/gal $",
            "Commodity $", "Freight $", "Total $",
            "Avg $/gal", "Solver", "Seconds", "Source workbook",
        ]
        self._header(ws, headers, row=row)
        first = row + 1
        r = first

        tot_demand = tot_purch = tot_ship = tot_comm = tot_frt = 0.0
        for sol in self.solutions:
            d = sol.data
            if sol.ok:
                purch = float(sol.purchased().sum())
                shipped = float(sol.ship.sum())
                comm = float(sol.commodity_cost().sum())
                frt = float(sol.freight_cost().sum())
                total = comm + frt
                tot_demand += d.total_demand
                tot_purch += purch
                tot_ship += shipped
                tot_comm += comm
                tot_frt += frt
                avg = total / to_gal(purch) if purch > 1e-9 else None
                values = [
                    d.area, "Relaxed" if sol.relaxed else "Optimal",
                    d.n_supply, d.n_demand,
                    to_gal(d.total_demand), to_gal(purch), to_gal(shipped),
                    to_gal(sol.unmet_total), to_gal(sol.surplus_total),
                    to_gal(sol.other_slack_total),
                    to_gal(sol.blocked_lane_volume()), sol.blocked_lane_cost(),
                    comm, frt, total, avg,
                    sol.solver, round(sol.solve_seconds, 2),
                    os.path.basename(d.source_path),
                ]
            else:
                values = [
                    d.area, sol.status.title(), d.n_supply, d.n_demand,
                    d.total_demand, None, None, None, None, None,
                    None, None,
                    None, None, None, None,
                    sol.solver, round(sol.solve_seconds, 2),
                    os.path.basename(d.source_path),
                ]
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row=r, column=c, value=v)
                if c == 2 and (not sol.ok or sol.relaxed):
                    cell.fill = _WARN_FILL
                    cell.font = Font(bold=True)
            if sol.ok and sol.relaxed and sol.relax_note:
                ws.cell(row=r, column=2).comment = openpyxl.comments.Comment(
                    sol.relax_note, "model.py"
                )
            r += 1

        last = r - 1
        tot_unmet = sum(s.unmet_total for s in self.solutions if s.ok)
        tot_surplus = sum(s.surplus_total for s in self.solutions if s.ok)
        tot_other = sum(s.other_slack_total for s in self.solutions if s.ok)
        tot_blk_vol = sum(s.blocked_lane_volume() for s in self.solutions if s.ok)
        tot_blk_cost = sum(s.blocked_lane_cost() for s in self.solutions if s.ok)
        money_cols = {5: _VOL, 6: _VOL, 7: _VOL, 8: _VOL, 9: _VOL, 10: _VOL,
                      11: _VOL, 12: _MONEY,
                      13: _MONEY, 14: _MONEY, 15: _MONEY, 16: _RATE}
        if any(s.ok for s in self.solutions):
            for c, v in (
                (1, f"{self.region} total"), (5, to_gal(tot_demand)), (6, to_gal(tot_purch)),
                (7, to_gal(tot_ship)), (8, to_gal(tot_unmet)), (9, to_gal(tot_surplus)),
                (10, to_gal(tot_other)),
                (11, to_gal(tot_blk_vol)), (12, tot_blk_cost),
                (13, tot_comm), (14, tot_frt),
                (15, tot_comm + tot_frt),
                (16, (tot_comm + tot_frt) / to_gal(tot_purch) if tot_purch > 1e-9 else None),
            ):
                cell = ws.cell(row=r, column=c, value=v)
                cell.font = Font(bold=True)
                cell.border = Border(top=Side(style="thin"))
            self._fmt(ws, r, r, money_cols)

        self._fmt(ws, first, last, money_cols)
        self._widths(ws, [22, 12, 11, 12, 15, 17, 15, 14, 15, 16,
                          20, 18,
                          14, 13, 14, 12, 14, 9, 34])

        # A monthly roll-up across the whole region, underneath the per-area table.
        ok = [s for s in self.solutions if s.ok]
        if ok:
            row = r + 3
            ws.cell(row=row, column=1, value="Region total by period").font = Font(bold=True)
            row += 1
            labels = ok[0].data.period_labels
            self._header(
                ws,
                ["Period", "Demand (gal)", "Purchased (gal)",
                 "Commodity $", "Freight $", "Total $"],
                row=row,
                freeze=False,          # leave the frozen pane on the area table above
            )
            start = row + 1
            for p, label in enumerate(labels):
                dem = sum(float(s.data.demand[p].sum()) for s in ok)
                pur = sum(float(s.purchased()[p].sum()) for s in ok)
                com = sum(float(s.commodity_cost()[p].sum()) for s in ok)
                frt = sum(float(s.freight_cost()[p].sum()) for s in ok)
                for c, v in enumerate([label, to_gal(dem), to_gal(pur), com, frt, com + frt], start=1):
                    ws.cell(row=start + p, column=c, value=v)
            self._fmt(ws, start, start + len(labels) - 1,
                      {2: _VOL, 3: _VOL, 4: _MONEY, 5: _MONEY, 6: _MONEY})

    def _supply_plan(self, wb) -> None:
        ws = wb.create_sheet("Supply Plan")
        has_surplus = any(s.ok and s.surplus_total > 1e-6 for s in self.solutions)
        headers = [
            "Area", "Supplier", "Supply point", "Hub", "Period",
            "Contract (gal)", "Min %", "Max %", "Lift %",
            "Purchased (gal)", "Hub price $/gal", "Premium $/gal", "Price $/gal",
            "Commodity $",
        ]
        if has_surplus:
            headers.append("Unplaceable (gal)")
        self._header(ws, headers)
        r = 2
        for sol in self.solutions:
            if not sol.ok:
                continue
            d = sol.data
            purch = sol.purchased()
            comm = sol.commodity_cost()
            for i in range(d.n_supply):
                for p in range(d.n_periods):
                    if d.contract[p, i] <= 0 and purch[p, i] <= 0:
                        continue
                    ws.cell(row=r, column=1, value=d.area)
                    ws.cell(row=r, column=2, value=d.supply_groups[i])
                    ws.cell(row=r, column=3, value=d.supply_names[i])
                    ws.cell(row=r, column=4, value=d.supply_hubs[i])
                    ws.cell(row=r, column=5, value=d.period_labels[p])
                    ws.cell(row=r, column=6, value=to_gal(d.contract[p, i]))
                    ws.cell(row=r, column=7, value=float(d.perc_min[p, i]))
                    ws.cell(row=r, column=8, value=float(d.perc_max[p, i]))
                    ws.cell(row=r, column=9, value=float(sol.perc[p, i]))
                    ws.cell(row=r, column=10, value=to_gal(purch[p, i]))
                    ws.cell(row=r, column=11, value=to_per_gal(d.hub_price[p, i]))
                    ws.cell(row=r, column=12, value=to_per_gal(d.hub_premium[p, i]))
                    ws.cell(row=r, column=13, value=to_per_gal(d.price[p, i]))
                    ws.cell(row=r, column=14, value=float(comm[p, i]))
                    if has_surplus:
                        ws.cell(
                            row=r, column=15,
                            value=to_gal(sol.surplus[p, i]) if sol.surplus is not None else 0.0,
                        )
                    r += 1
        spec = {
            6: _VOL, 7: _PCT, 8: _PCT, 9: _PCT, 10: _VOL,
            11: _RATE, 12: _RATE, 13: _RATE, 14: _MONEY,
        }
        if has_surplus:
            spec[15] = _VOL
        self._fmt(ws, 2, r - 1, spec)
        self._widths(ws, [18, 16, 26, 8, 10, 15, 9, 9, 9, 17, 11, 11, 13, 14, 17])

    def _shipments(self, wb) -> None:
        ws = wb.create_sheet("Shipments")
        self._header(ws, [
            "Area", "Period", "From (supply)", "To (demand)",
            "Volume (gal)", "Freight $/gal", "Freight $", "High freight?",
        ])
        r = 2
        for sol in self.solutions:
            if not sol.ok:
                continue
            d = sol.data
            blocked = d.blocked_lane_mask()
            nz = np.argwhere(sol.ship > self.SHIP_EPS)
            for p, i, j in nz:
                vol = float(sol.ship[p, i, j])
                rate = float(d.trans_cost[p, i, j])
                ws.cell(row=r, column=1, value=d.area)
                ws.cell(row=r, column=2, value=d.period_labels[int(p)])
                ws.cell(row=r, column=3, value=d.supply_names[int(i)])
                ws.cell(row=r, column=4, value=d.demand_names[int(j)])
                ws.cell(row=r, column=5, value=to_gal(vol))
                ws.cell(row=r, column=6, value=to_per_gal(rate))
                ws.cell(row=r, column=7, value=vol * rate)
                if blocked[p, i, j]:
                    # A warning flag only: the lane is used because it was
                    # the cheapest route; the rate is worth a second look.
                    cell = ws.cell(row=r, column=8, value="CHECK: >= $1/gal")
                    cell.fill = _WARN_FILL
                    cell.font = Font(bold=True)
                r += 1
        self._fmt(ws, 2, r - 1, {5: _VOL, 6: _RATE, 7: _MONEY})
        self._widths(ws, [18, 10, 26, 28, 15, 15, 14, 16])
        if r == 2:
            ws.cell(row=2, column=1, value="No shipments (no area solved successfully).")

    def _demand_coverage(self, wb) -> None:
        ws = wb.create_sheet("Demand Coverage")
        headers = [
            "Area", "Demand point", "Period", "Demand (gal)",
            "Delivered (gal)", "From storage (gal)", "To storage (gal)",
            "Ending inventory (gal)", "Storage capacity (gal)",
        ]
        has_short = any(s.ok and s.shortfall is not None for s in self.solutions)
        if has_short:
            headers.append("Unmet (gal)")
        self._header(ws, headers)

        r = 2
        for sol in self.solutions:
            if not sol.ok:
                continue
            d = sol.data
            delivered = sol.delivered()
            for j in range(d.n_demand):
                for p in range(d.n_periods):
                    if (
                        d.demand[p, j] <= 0
                        and delivered[p, j] <= self.SHIP_EPS
                        and d.demand_inv_cap[p, j] <= 0
                    ):
                        continue
                    ws.cell(row=r, column=1, value=d.area)
                    ws.cell(row=r, column=2, value=d.demand_names[j])
                    ws.cell(row=r, column=3, value=d.period_labels[p])
                    ws.cell(row=r, column=4, value=to_gal(d.demand[p, j]))
                    ws.cell(row=r, column=5, value=to_gal(delivered[p, j]))
                    ws.cell(row=r, column=6, value=to_gal(sol.demand_wd[p, j]))
                    ws.cell(row=r, column=7, value=to_gal(sol.demand_inj[p, j]))
                    ws.cell(row=r, column=8, value=to_gal(sol.demand_end[p, j]))
                    ws.cell(row=r, column=9, value=to_gal(d.demand_inv_cap[p, j]))
                    if has_short:
                        ws.cell(
                            row=r, column=10,
                            value=to_gal(sol.shortfall[p, j]) if sol.shortfall is not None else 0.0,
                        )
                    r += 1
        cols = {c: _VOL for c in range(4, 10 + (1 if has_short else 0))}
        self._fmt(ws, 2, r - 1, cols)
        self._widths(ws, [18, 28, 10, 15, 17, 14, 14, 17, 17, 14])

    def _inventory(self, wb, at_supply: bool) -> None:
        title = "Supply Inventory" if at_supply else "Demand Inventory"
        ws = wb.create_sheet(title)
        label = "Supply point" if at_supply else "Demand point"
        self._header(ws, [
            "Area", label, "Period", "Beginning (gal)", "Withdrawn (gal)",
            "Injected (gal)", "Ending (gal)", "Capacity (gal)", "Utilisation",
        ])
        r = 2
        for sol in self.solutions:
            if not sol.ok:
                continue
            d = sol.data
            if at_supply:
                names, cap, init = d.supply_names, d.supply_inv_cap, d.supply_inv_init
                wd, inj, end = sol.supply_wd, sol.supply_inj, sol.supply_end
                n = d.n_supply
            else:
                names, cap, init = d.demand_names, d.demand_inv_cap, d.demand_inv_init
                wd, inj, end = sol.demand_wd, sol.demand_inj, sol.demand_end
                n = d.n_demand

            for k in range(n):
                # No point cluttering the sheet with a point that has no
                # storage and nothing moving through it either.
                if cap[:, k].max() <= 0 and init[k] <= 0:
                    continue
                for p in range(d.n_periods):
                    begin = float(init[k]) if p == 0 else float(end[p - 1, k])
                    ws.cell(row=r, column=1, value=d.area)
                    ws.cell(row=r, column=2, value=names[k])
                    ws.cell(row=r, column=3, value=d.period_labels[p])
                    ws.cell(row=r, column=4, value=to_gal(begin))
                    ws.cell(row=r, column=5, value=to_gal(wd[p, k]))
                    ws.cell(row=r, column=6, value=to_gal(inj[p, k]))
                    ws.cell(row=r, column=7, value=to_gal(end[p, k]))
                    ws.cell(row=r, column=8, value=to_gal(cap[p, k]))
                    ws.cell(
                        row=r, column=9,
                        value=float(end[p, k] / cap[p, k]) if cap[p, k] > 1e-9 else None,
                    )
                    r += 1
        self._fmt(ws, 2, r - 1, {4: _VOL, 5: _VOL, 6: _VOL, 7: _VOL, 8: _VOL, 9: _PCT})
        self._widths(ws, [18, 28, 10, 13, 13, 12, 12, 12, 12])
        if r == 2:
            ws.cell(row=2, column=1, value="No storage configured in this region.")

    def _diagnostics(self, wb) -> None:
        ws = wb.create_sheet("Diagnostics")
        self._header(ws, ["Area", "Severity", "Detail"])
        r = 2

        def add(area: str, sev: str, text: str) -> None:
            nonlocal r
            ws.cell(row=r, column=1, value=area)
            c = ws.cell(row=r, column=2, value=sev)
            if sev in ("Error", "Warning"):
                c.fill = _WARN_FILL
            # Messages are written in model units; show people gallons.
            ws.cell(row=r, column=3, value=gal_text(text))
            r += 1

        for sol in self.solutions:
            d = sol.data
            add(d.area, "Info", f"Source workbook: {d.source_path}")
            add(
                d.area, "Info",
                f"{d.n_supply} supply points, {d.n_demand} demand points, "
                f"{d.n_periods} periods, total demand {d.total_demand:,.1f} kgal.",
            )
            if d.blocked_lane_pairs:
                add(
                    d.area, "Info",
                    f"{d.blocked_lane_pairs} supply/demand pair(s) have freight of "
                    f"$1/gal or more in at least one month. They are used like any "
                    f"other lane when they're the cheapest way. Only a blank freight "
                    f"cell keeps a lane unused.",
                )
            closed = d.closed_lane_mask()
            if closed.any():
                add(
                    d.area, "Info",
                    f"{int(closed.all(axis=0).sum())} supply/demand pair(s) have no freight "
                    f"rate in any month, and {int((closed.any(axis=0) & ~closed.all(axis=0)).sum())} "
                    f"more in some months. A blank freight cell means the lane is not used.",
                )
            for repair in d.repairs:
                add(d.area, "Error", "DATA ERROR in workbook - " + repair)
            if d.override_log:
                add(
                    d.area, "Info",
                    f"{len(d.override_log)} manual override(s) applied from the "
                    f"entry deck - see the Overrides sheet for the full list "
                    f"with before/after values.",
                )
            for note in d.notes:
                add(d.area, "Warning", note)
            for warn in d.feasibility_warnings():
                add(d.area, "Warning", warn)

            # -- saved-plan comparison (the regression oracle) --------------- #
            if self.compare_saved:
                self._add_saved_cost_rows(add, sol)

            if not sol.ok:
                add(d.area, "Error", f"Solve {sol.status}: {sol.message}")
                continue

            if sol.relaxed:
                add(d.area, "Error", sol.relax_note)
            for v in sol.violations:
                add(d.area, "Error", "UNSATISFIABLE CONSTRAINT - " + v)

            # Spell out the worst offenders so somebody can act on it without
            # having to pivot the detail sheets themselves.
            if sol.unmet_total > 1e-6:
                add(
                    d.area, "Error",
                    f"{sol.unmet_total:,.1f} kgal of demand is unservable in total. "
                    f"Either the contracted volume or the maximum lift % is too "
                    f"low in the periods listed below, or extra supply is needed.",
                )
                per_period = sol.shortfall.sum(axis=1)
                for p in np.argsort(-per_period):
                    if per_period[p] <= 1e-6:
                        continue
                    worst = np.argsort(-sol.shortfall[p])[:3]
                    detail = ", ".join(
                        f"{d.demand_names[j]} {sol.shortfall[p, j]:,.1f} kgal"
                        for j in worst if sol.shortfall[p, j] > 1e-6
                    )
                    add(
                        d.area, "Warning",
                        f"  {d.period_labels[p]}: {per_period[p]:,.1f} kgal short "
                        f"(largest: {detail})",
                    )

            if sol.surplus_total > 1e-6:
                add(
                    d.area, "Error",
                    f"{sol.surplus_total:,.1f} kgal is forced up by contract "
                    f"minimums with no demand or storage to absorb it. Lower the "
                    f"minimum lift % in the periods below, or add storage capacity.",
                )
                per_period = sol.surplus.sum(axis=1)
                for p in np.argsort(-per_period):
                    if per_period[p] <= 1e-6:
                        continue
                    worst = np.argsort(-sol.surplus[p])[:3]
                    detail = ", ".join(
                        f"{d.supply_names[i]} {sol.surplus[p, i]:,.1f} kgal"
                        for i in worst if sol.surplus[p, i] > 1e-6
                    )
                    add(
                        d.area, "Warning",
                        f"  {d.period_labels[p]}: {per_period[p]:,.1f} kgal surplus "
                        f"(largest: {detail})",
                    )

            if sol.excess_total > 1e-6:
                add(
                    d.area, "Error",
                    f"{sol.excess_total:,.1f} kgal is forced into demand points by "
                    f"MINIMUM TRANSPORT QUANTITIES beyond what those points can use "
                    f"or store. Lower the minimum transport quantity on the lanes "
                    f"feeding the locations below, or raise their demand/storage.",
                )
                nz = np.argwhere(sol.excess > 1e-6)
                for p, j in nz[:10]:
                    add(
                        d.area, "Warning",
                        f"  {d.demand_names[int(j)]} in {d.period_labels[int(p)]}: "
                        f"{sol.excess[p, j]:,.1f} kgal more delivered than usable "
                        f"(demand {d.demand[p, j]:,.1f} kgal, storage capacity "
                        f"{d.demand_inv_cap[p, j]:,.1f} kgal)",
                    )

            # Volume on lanes with high freight. This changes no number - it
            # just makes expensive lanes visible so the rates get checked.
            blk_vol = sol.high_freight_volume()
            if blk_vol >= HIGH_FREIGHT_REPORT_MIN:
                add(
                    d.area, "Warning",
                    f"HIGH FREIGHT: {blk_vol:,.1f} kgal moves on lanes with freight of "
                    f"$1/gal or more, costing ${sol.high_freight_cost():,.0f}. "
                    f"These are the cheapest routes the model found; check the "
                    f"rates are right. They are flagged on the Shipments sheet.",
                )
                mask = d.blocked_lane_mask()
                used = np.argwhere((sol.ship > self.SHIP_EPS) & mask)
                worst = sorted(
                    used, key=lambda t: -float(sol.ship[t[0], t[1], t[2]])
                )[:8]
                for p, i, j in worst:
                    add(
                        d.area, "Warning",
                        f"  {d.supply_names[int(i)]} -> {d.demand_names[int(j)]} "
                        f"in {d.period_labels[int(p)]}: "
                        f"{sol.ship[p, i, j]:,.1f} kgal",
                    )

            if sol.unsourced_total > 1e-6:
                add(
                    d.area, "Error",
                    f"{sol.unsourced_total:,.1f} kgal would need to ship with no "
                    f"contracted volume behind it - minimum transport quantities "
                    f"exceed what the contract maximums allow to be lifted.",
                )
                nz = np.argwhere(sol.unsourced > 1e-6)
                for p, i in nz[:10]:
                    add(
                        d.area, "Warning",
                        f"  {d.supply_names[int(i)]} in {d.period_labels[int(p)]}: "
                        f"{sol.unsourced[p, i]:,.1f} kgal beyond contract capacity",
                    )
        self._widths(ws, [20, 12, 130])
        for row in range(2, r):
            ws.cell(row=row, column=3).alignment = Alignment(wrap_text=True, vertical="top")

    #: A re-solve within this fraction of the saved cost means the workbook is
    #: current and the transcription is confirmed correct.
    SAVED_COST_MATCH_TOL = 0.001      # 0.1%

    def _add_saved_cost_rows(self, add, sol: "Solution") -> None:
        """
        Compare the freshly solved cost against the plan already stored in the
        workbook.

        This is the one check that actually proves fidelity, and all three
        outcomes are real and mean different things, so each gets its own
        interpretation rather than a bare number.
        """
        d = sol.data
        saved = d.saved_cost
        if not saved:
            add(d.area, "Info",
                "--compare-saved: no saved plan cost could be read from this "
                "workbook, so there is nothing to compare against.")
            return

        saved_total = saved.get("total", 0.0)
        if not sol.ok:
            add(d.area, "Info",
                f"Saved plan cost (as stored in workbook) : ${saved_total:,.0f}. "
                f"This area did not solve, so there is no re-solved cost to "
                f"compare it with.")
            return

        resolved = float(sol.commodity_cost().sum() + sol.freight_cost().sum())
        diff = resolved - saved_total
        pct = (diff / saved_total * 100.0) if abs(saved_total) > 1e-9 else float("nan")

        add(d.area, "Info",
            f"Saved plan cost (as stored in workbook) : ${saved_total:,.0f}")
        add(d.area, "Info",
            f"Re-solved optimal cost                  : ${resolved:,.0f}")
        sign = "-" if diff < 0 else "+"
        add(d.area, "Info",
            f"Difference                              : {sign}${abs(diff):,.0f}"
            f"   ({pct:+.1f}%)")

        # Component split, which is where a transcription error would show up
        # even when the totals happen to land close together.
        add(d.area, "Info",
            f"  saved split: commodity ${saved.get('commodity', 0):,.0f} + "
            f"premium ${saved.get('premium', 0):,.0f} + transport "
            f"${saved.get('transport', 0):,.0f}")
        add(d.area, "Info",
            f"  re-solved  : commodity+premium "
            f"${float(sol.commodity_cost().sum()):,.0f} + transport "
            f"${float(sol.freight_cost().sum()):,.0f}")

        if abs(saved_total) > 1e-9 and abs(diff) / abs(saved_total) <= self.SAVED_COST_MATCH_TOL:
            add(d.area, "Info",
                "Within 0.1% - the workbook is current and this model's "
                "transcription of it is confirmed correct.")
        elif diff < 0:
            add(d.area, "Info",
                "Re-solved cost is materially LOWER than the saved plan, which "
                "means the saved plan is stale or was solved against different "
                "inputs. This is expected for most areas and is not a problem "
                "with the model.")
        else:
            add(d.area, "Error",
                "Re-solved cost is HIGHER than the saved plan. That means the "
                "saved plan is INFEASIBLE against the inputs as they now stand - "
                "Excel saved a solution and then somebody edited an input. The "
                "workbook is currently showing a cost it cannot actually "
                "achieve. Re-solve it in Excel, or trust this run instead.")

    # -- entry point ------------------------------------------------------- #

    def _overrides(self, wb) -> None:
        """
        One row per applied override.

        Non-negotiable for a hand-edited scenario: a run whose inputs were
        changed by hand and which keeps no record of what was changed is not a
        usable artifact. Before/after values make it reviewable, and the deck
        line number makes it correctable.
        """
        ws = wb.create_sheet("Overrides")
        self._header(ws, [
            "Area", "Scope", "Target", "Field",
            "Value before (gal, $/gal)", "Value after (gal, $/gal)", "Deck line",
        ])
        r = 2
        for sol in self.solutions:
            for rec in sol.data.override_log:
                ws.cell(row=r, column=1, value=rec.area)
                ws.cell(row=r, column=2, value=rec.scope)
                ws.cell(row=r, column=3, value=rec.target)
                ws.cell(row=r, column=4, value=rec.field)
                ws.cell(row=r, column=5, value=gal_series_text(rec.before, rec.field))
                ws.cell(row=r, column=6, value=gal_series_text(rec.after, rec.field))
                ws.cell(row=r, column=7, value=rec.line_no)
                r += 1
        self._widths(ws, [20, 18, 44, 12, 30, 30, 10])
        if r == 2:
            ws.cell(row=2, column=1,
                    value="No manual overrides were applied to this region.")
        else:
            for row in range(2, r):
                for col in (5, 6):
                    ws.cell(row=row, column=col).alignment = Alignment(
                        wrap_text=True, vertical="top"
                    )

    def write(self, path: str) -> str:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        self._summary(wb)
        self._supply_plan(wb)
        self._shipments(wb)
        self._demand_coverage(wb)
        self._inventory(wb, at_supply=True)
        self._inventory(wb, at_supply=False)
        # Only add the audit sheet when a deck was actually in play, so a plain
        # run's workbook stays byte-for-byte what it was before this feature.
        if any(s.data.override_log for s in self.solutions):
            self._overrides(wb)
        self._diagnostics(wb)

        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        wb.save(path)
        return path


# --------------------------------------------------------------------------- #
# Orchestration - tying discovery, solving and writing together.
# --------------------------------------------------------------------------- #


def default_root() -> str:
    """
    Best guess at where the Model directory actually is.

    This file normally lives at ``Model/Test/model.py``, so the parent of its
    own folder is the model root.
    """
    here = os.path.dirname(os.path.abspath(__file__))
    parent = os.path.dirname(here)
    if os.path.basename(here).lower() == "test" and os.path.isdir(parent):
        return parent
    return here


def run_region(
    region: str,
    areas: Sequence[AreaFile],
    out_dir: str,
    run_date: _dt.date,
    allow_shortfall: bool = False,
    time_limit: Optional[float] = None,
    contract_label: str = "",
    auto_relax: bool = True,
    deck: Optional[OverrideDeck] = None,
    compare_saved: bool = False,
) -> Tuple[Optional[str], List[Solution], int]:
    """
    Read, solve, and write out results for every area in one region.

    Returns ``(workbook_path_or_None, solutions, load_failures)``. Areas that
    could not even be READ never make it into *solutions*, so they are counted
    separately - otherwise a refused workbook would leave the run looking
    completely clean.
    """
    LOG.info("=" * 72)
    LOG.info("Region %s  (%d area%s)", region, len(areas), "" if len(areas) == 1 else "s")

    solutions: List[Solution] = []
    load_failures = 0
    for af in areas:
        try:
            data = ModelBaseReader(
                af.path, af.area, region,
                deck=deck, want_saved_cost=compare_saved,
            ).read()
        except OverrideError:
            # A bad deck is the user's to fix and affects every area, so let it
            # escape to main() rather than quietly skipping this one area.
            raise
        except ModelStructureError as exc:
            LOG.error("%s: %s", af.area, exc)
            load_failures += 1
            continue
        except Exception as exc:  # noqa: BLE001 - one bad workbook shouldn't kill the batch
            LOG.error("%s: failed to read %s (%s)", af.area, af.path, exc)
            LOG.debug("%s", traceback.format_exc())
            load_failures += 1
            continue

        try:
            sol = solve_area(
                data,
                allow_shortfall=allow_shortfall,
                time_limit=time_limit,
                auto_relax=auto_relax,
            )
        except Exception as exc:  # noqa: BLE001
            LOG.error("%s: solve failed (%s)", af.area, exc)
            LOG.debug("%s", traceback.format_exc())
            sol = Solution(data=data, status="error", message=str(exc))
        solutions.append(sol)

        # Echo the saved-vs-resolved headline to the console as well as the
        # Diagnostics sheet, so the flag visibly does something when you run it.
        if compare_saved and data.saved_cost and sol.ok:
            saved = data.saved_cost.get("total", 0.0)
            got = float(sol.commodity_cost().sum() + sol.freight_cost().sum())
            delta = got - saved
            pct = (delta / saved * 100.0) if abs(saved) > 1e-9 else float("nan")
            verdict = ("MATCHES saved plan" if abs(saved) > 1e-9
                       and abs(delta) / abs(saved) <= 0.001
                       else ("cheaper than saved plan" if delta < 0
                             else "DEARER than saved plan - saved plan is infeasible"))
            LOG.info("%-16s saved $%s vs re-solved $%s  (%+.1f%%, %s)",
                     data.area, f"{saved:,.0f}", f"{got:,.0f}", pct, verdict)

    if not solutions:
        LOG.warning("Region %s produced no results; no workbook written.", region)
        return None, [], load_failures

    variants = sorted({af.variant for af in areas})
    note = (
        f"Built from the {'/'.join(variants)} workbook for each area: "
        + ", ".join(f"{s.data.area} ({os.path.basename(s.data.source_path)})" for s in solutions)
    )
    scenario = deck.scenario if deck else ""
    if deck is not None:
        applied = sum(len(s.data.override_log) for s in solutions)
        note += f"  |  {applied} manual override(s) from {os.path.basename(deck.path)}"

    # Each region gets its own subfolder, which keeps a day's worth of runs
    # (and several scenarios) navigable instead of one flat heap.
    region_dir = os.path.join(out_dir, slugify(region) or "Region")
    out_path = os.path.join(
        region_dir, region_filename(region, run_date, contract_label, scenario)
    )
    writer = RegionWorkbookWriter(
        region, solutions, run_date, source_note=note,
        scenario=scenario, compare_saved=compare_saved,
    )
    written = writer.write(out_path)
    LOG.info("wrote %s", written)
    return written, solutions, load_failures


#: Header written at the top of a generated deck. Documents the units, the
#: series rule and - required, because it is easy to get wrong - the precedence
#: order between the workbook, [all] and [area: X].
OVERRIDE_TEMPLATE_HEADER = '''\
# ===================================================================
#  Override deck - hand-entered changes applied on top of the workbooks
#
#  Volumes: kgal (1000 gal).  Prices / diffs / freight: $/kgal.
#  A monthly value is either ONE number (applies to all 12 months) or
#  exactly 12 numbers separated by spaces or commas.
#  Anything after a # is a comment.  Blank lines are ignored.
#  Names are matched loosely - case and spacing do not matter.
#
#  PRECEDENCE:  workbook value  ->  [all]  ->  [area: X],  last wins.
#  So an [area: X] line always beats an [all] line for that area.
#
#  The source workbooks are never modified. Nothing here writes back.
#
#  Anything that matches NOTHING is a hard error naming the line and the
#  nearest spellings - never a silent no-op. Run
#      py model.py --check-overrides overrides.txt
#  to validate without solving.
#
#  Fields you can set:
#    supply : contract, diff, hub_price, min_pct, max_pct,
#             store_max, store_min, store_init
#    demand : demand, store_max, store_min, store_init
#    lane   : freight, min_ship, max_ship
#    hub    : price
#    [all] / [area: X] bare keys : min_pct, max_pct  (blanket sweeps)
#
#  min_pct / max_pct are multipliers, not percents: 1.0 = 100%, 5 = 500%.
#  store_init is a single opening balance, not a monthly series.
#
#  Where one location has several suppliers, qualify it:
#      supply "Terminal A, ST" | Acme : min_pct = 0.9
#  An unqualified ambiguous name is an error for `supply`, because it changes
#  commercial terms. For `lane` it is allowed and applies to every matching
#  route (each one listed separately on the Overrides sheet).
# ===================================================================

[scenario]
name = Wide open first pass          # goes in the filename and on every sheet

[all]
# The classic first pass: run wide open at 500% of contract, see what the
# model wants, then tighten. Uncomment to use.
# max_pct = 5

'''


def write_override_template(dest: str, regions: Dict[str, List[AreaFile]]) -> str:
    """
    Write an annotated blank deck, with the real contents of the discovered
    workbooks listed as commented-out example lines.

    This is what makes the feature usable. Every supply point, demand point and
    hub is spelled correctly and already qualified where it needs to be, so the
    team copies a line, uncomments it and changes the number. Nobody should ever
    have to type a location name from scratch.

    The ``[area: ...]`` headers are written live rather than commented, because
    a section with every line commented out is inert anyway - and that way
    uncommenting a single example line actually works.
    """
    chunks: List[str] = [OVERRIDE_TEMPLATE_HEADER]
    total_areas = 0
    failed: List[str] = []

    for region in sorted(regions):
        for af in regions[region]:
            total_areas += 1
            try:
                data = ModelBaseReader(af.path, af.area, region).read()
            except Exception as exc:  # noqa: BLE001 - keep going, note it
                failed.append(f"{af.area}: {exc}")
                chunks.append(
                    f"# ==== {af.area}  (region {region}) ====\n"
                    f"#  NOT LISTED - this workbook could not be read:\n"
                    f"#  {exc}\n\n"
                )
                continue

            lines: List[str] = []
            lines.append(f"# =================================================="
                         f"=================\n")
            lines.append(f"#  {af.area}   (region {region};  "
                         f"{data.n_supply} supply, {data.n_demand} demand)\n")
            lines.append(f"#  source: {os.path.basename(af.path)}\n")
            lines.append(f"# =================================================="
                         f"=================\n")
            lines.append(f"[area: {af.area}]\n")
            lines.append("\n")

            # Which location names need a supplier qualifier to be unambiguous?
            counts: Dict[str, int] = {}
            for nm in data.supply_names:
                counts[fold_name(nm)] = counts.get(fold_name(nm), 0) + 1

            lines.append("#  -- supply points "
                         "-------------------------------------------\n")
            for i, nm in enumerate(data.supply_names):
                grp = data.supply_groups[i]
                qual = f' | {grp}' if (counts[fold_name(nm)] > 1 and grp) else ""
                hub = f"   hub {data.supply_hubs[i]}" if data.supply_hubs[i] else ""
                lines.append(f'#  supply "{nm}"{qual} : max_pct = 1.5{hub}\n')

            lines.append("#\n#  -- demand points "
                         "-------------------------------------------\n")
            for nm in data.demand_names:
                lines.append(f'#  demand "{nm}" : demand = 0\n')

            hubs = sorted({h for h in data.supply_hubs if h})
            if hubs:
                lines.append("#\n#  -- market hubs "
                             "---------------------------------------------\n")
                for h in hubs:
                    lines.append(f"#  hub {h} : price = 660\n")

            if data.supply_names and data.demand_names:
                lines.append("#\n#  -- lanes (a few examples; any supply -> any "
                             "demand works) ------\n")
                for i in range(min(3, data.n_supply)):
                    j = min(i, data.n_demand - 1)
                    lines.append(
                        f'#  lane "{data.supply_names[i]}" -> '
                        f'"{data.demand_names[j]}" : freight = 85.50\n'
                    )
                lines.append(
                    f'#  lane "{data.supply_names[0]}" -> '
                    f'"{data.demand_names[0]}" : freight = 1150'
                    f'   # a real rate ($1.15/gal) - used if cheapest, flagged as high freight\n'
                )

            lines.append("\n\n")
            chunks.append("".join(lines))

    os.makedirs(os.path.dirname(os.path.abspath(dest)) or ".", exist_ok=True)
    with open(dest, "w", encoding="utf-8") as fh:
        fh.write("".join(chunks))

    LOG.info("wrote override template: %s", dest)
    LOG.info("  %d area(s) listed%s", total_areas - len(failed),
             f", {len(failed)} could not be read" if failed else "")
    LOG.info("  Everything is commented out, so running now changes nothing.")
    LOG.info("  Uncomment a line, change the number, then: "
             "py model.py --check-overrides %s", os.path.basename(dest))
    return dest


def check_override_deck(deck: OverrideDeck,
                        regions: Dict[str, List[AreaFile]]) -> int:
    """
    Validate a deck and echo every resolved change, without solving anything.

    This is the "did I type it right" command, so it stays deliberately quick:
    only the areas the deck actually names by name get their workbook read,
    since those are the only places a name can fail to resolve. A deck that
    only carries blanket ``[all]`` sweeps needs no workbook at all - nothing in
    it can fail to match - so it is reported symbolically and returns instantly.
    """
    print(f"Deck      : {deck.path}")
    print(f"Scenario  : {deck.scenario or '(none set)'}")
    print(f"Rules     : {len(deck.rules)}")
    print()

    all_areas = [af.area for areas in regions.values() for af in areas]
    deck.validate_area_scopes(all_areas)

    blanket = [r for r in deck.rules if r.scope == "all"]
    named_scopes = deck.area_scopes

    if blanket:
        print("[all] - blanket sweeps, applied to every area in the run:")
        for r in blanket:
            target = r.target or "every supply point"
            sel = " ".join(r.selector) if r.selector else ""
            vals = _fmt_series(np.asarray(r.values))
            print(f"  line {r.line_no:>4}: {target} {sel}".rstrip()
                  + f" -> {r.field} = {vals}")
        print(f"  (these touch all {len(all_areas)} area(s))")
        print()

    if not named_scopes:
        print("No live [area: ...] rules, so there are no location names to "
              "resolve - nothing here can fail to match. (Section headers on "
              "their own assert nothing; only uncommented rules under them "
              "name a location.)")
        print("Deck is valid.")
        return 0

    # Only read the workbooks the deck actually names.
    wanted: List[Tuple[str, AreaFile]] = []
    for region in sorted(regions):
        for af in regions[region]:
            if any(_area_matches(s, af.area) for s in named_scopes):
                wanted.append((region, af))

    total = 0
    for region, af in wanted:
        print(f"[area: {af.area}]   (region {region})")
        data = ModelBaseReader(af.path, af.area, region, deck=deck).read()
        if not data.override_log:
            print("  (no change - the deck's values already match the workbook)")
        for rec in data.override_log:
            total += 1
            print(f"  line {rec.line_no:>4}: {rec.target}")
            print(f"              {rec.field}: {rec.before}  ->  {rec.after}")
        for rep in data.repairs:
            print(f"  ! repair after override: {rep}")
        print()

    print(f"Deck is valid. {total} change(s) would be applied across "
          f"{len(wanted)} area(s).")
    return 0


def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="model.py",
        description=(
            "Rebuild the Excel/OpenSolver propane supply optimization in Python "
            "and emit one workbook per region."
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "examples:\n"
            "  python model.py                                 every region\n"
            "  python model.py --region Rail                   just Rail\n"
            "  python model.py --region Rail,Rockies,PNW       several at once\n"
            "  python model.py --region \"North East\"            spelling is forgiving\n"
            "  python model.py --exclude Rail                  everything but Rail\n"
            "  python model.py --group jdoe                    a saved manager group\n"
            "  python model.py --list-regions                  what is available\n"
            "  python model.py --region South --allow-shortfall -v\n"
        ),
    )
    p.add_argument(
        "--root", default=None,
        help="Model directory containing the region folders (default: the parent "
             "of this script's folder).",
    )
    p.add_argument(
        "--out-dir", default=None,
        help="Where to write the region workbooks (default: <this script's "
             "folder>/Output).",
    )
    p.add_argument(
        "--region", "-r", action="append", default=None, metavar="NAME",
        help="Region(s) or saved group(s) to run. Repeatable and/or "
             "comma-separated, e.g. -r Rail -r Rockies  or  -r Rail,Rockies. "
             "Names are matched case- and spacing-insensitively, so "
             "'North East' finds Northeast. A group name from groups.json is "
             "accepted here too. Default: every region found.",
    )
    p.add_argument(
        "--exclude", "-x", action="append", default=None, metavar="NAME",
        help="Region(s) to skip; same flexible matching as --region. Applied "
             "after --region/--group.",
    )
    p.add_argument(
        "--group", "-g", action="append", default=None, metavar="NAME",
        help="Run a saved group of regions from groups.json (see --list-groups). "
             "Combines with --region.",
    )
    p.add_argument(
        "--groups-file", default=None, metavar="PATH",
        help="Location of the group definitions (default: groups.json next to "
             "this script).",
    )
    p.add_argument(
        "--variant", default="Working", choices=["Working", "Baseline"],
        help="Which workbook flavour to read for each area (default: Working).",
    )
    p.add_argument(
        "--date", default=None, metavar="YYYY-MM-DD",
        help="Date stamp for the output filenames (default: today).",
    )
    p.add_argument(
        "--contract-label", default="", metavar="TEXT",
        help="Optional extra token in the filename, e.g. 26_27 -> "
             "West_Model_26_27_09_17_26.xlsx.",
    )
    p.add_argument(
        "--allow-shortfall", action="store_true",
        help="Permit unmet demand from the very first solve, rather than only "
             "as a fallback. Useful for asking 'how short am I?' directly.",
    )
    p.add_argument(
        "--no-auto-relax", action="store_true",
        help="If the workbook's own limits are infeasible, report the failure "
             "instead of automatically retrying with unmet-demand and surplus "
             "allowances. Use this when you want strict spreadsheet fidelity.",
    )
    p.add_argument(
        "--time-limit", type=float, default=None, metavar="SECONDS",
        help="Per-area solver time limit.",
    )
    p.add_argument(
        "--overrides", default=None, metavar="PATH",
        help="Apply hand-entered input changes from a deck file. When omitted, "
             f"{DEFAULT_OVERRIDES_NAME} next to this script is used if it "
             "exists; if that is absent too, the run is exactly as it was "
             "before this feature existed. An explicitly named file that does "
             "not exist is an error, never a silent skip.",
    )
    p.add_argument(
        "--template-overrides", action="store_true",
        help="Write an annotated blank deck listing every real supply point, "
             "demand point, hub and some example lanes as commented-out lines, "
             "then exit. Use -o to choose where it goes.",
    )
    p.add_argument(
        "-o", "--out", default=None, metavar="PATH",
        help=f"Where --template-overrides writes to "
             f"(default: {DEFAULT_OVERRIDES_NAME} next to this script).",
    )
    p.add_argument(
        "--check-overrides", default=None, metavar="PATH",
        help="Validate a deck and echo every resolved change with before/after "
             "values, without solving anything. Then exit.",
    )
    p.add_argument(
        "--compare-saved", action="store_true",
        help="Also read the cost of the plan already stored in each workbook "
             "and report saved vs re-solved vs difference on Diagnostics. "
             "Reporting only - a difference never fails the run.",
    )
    p.add_argument(
        "--list", action="store_true",
        help="Show the regions/areas/workbooks that would be processed, then exit.",
    )
    p.add_argument(
        "--list-regions", action="store_true",
        help="List the region folders available, then exit.",
    )
    p.add_argument(
        "--list-groups", action="store_true",
        help="List the saved groups from groups.json, then exit.",
    )
    p.add_argument("-v", "--verbose", action="store_true", help="Debug logging.")
    p.add_argument("-q", "--quiet", action="store_true", help="Warnings and errors only.")
    return p


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = build_arg_parser().parse_args(argv)

    level = logging.INFO
    if args.verbose:
        level = logging.DEBUG
    elif args.quiet:
        level = logging.WARNING
    logging.basicConfig(level=level, format="%(levelname)-7s %(message)s", stream=sys.stdout)

    script_dir = os.path.dirname(os.path.abspath(__file__))
    root = os.path.abspath(args.root or default_root())
    # The default output location lives beside this script, so a run never
    # ends up writing into the same folders as the source workbooks.
    out_dir = os.path.abspath(args.out_dir or os.path.join(script_dir, "Output"))
    groups_file = os.path.abspath(args.groups_file or os.path.join(script_dir, "groups.json"))

    if args.date:
        try:
            run_date = _dt.datetime.strptime(args.date, "%Y-%m-%d").date()
        except ValueError:
            LOG.error("--date must look like YYYY-MM-DD, got %r", args.date)
            return 2
    else:
        run_date = _dt.date.today()

    default_deck = os.path.join(script_dir, DEFAULT_OVERRIDES_NAME)

    groups = load_groups(groups_file)

    if args.list_groups:
        if not groups:
            print(f"No groups defined. Create {groups_file} to add some, e.g.:")
            print('  { "groups": { "jdoe": ["Rail", "Rockies"] } }')
            return 0
        print(f"Groups defined in {groups_file}:\n")
        for name in sorted(groups):
            print(f"  {name:<18} {', '.join(groups[name])}")
        return 0

    LOG.info("model root : %s", root)
    LOG.info("output dir : %s", out_dir)
    LOG.info("variant    : %s", args.variant)

    regions = discover(root, variant=args.variant)
    if not regions:
        LOG.error("no %s workbooks found under %s", args.variant, root)
        return 1

    # ---- deck-only commands, which never solve --------------------------- #
    if args.template_overrides:
        dest = os.path.abspath(args.out or default_deck)
        try:
            write_override_template(dest, regions)
        except Exception as exc:  # noqa: BLE001
            LOG.error("could not write the template: %s", exc)
            LOG.debug("%s", traceback.format_exc())
            return 1
        return 0

    if args.check_overrides:
        path = os.path.abspath(args.check_overrides)
        if not os.path.isfile(path):
            LOG.error("override file not found: %s", path)
            return 1
        try:
            deck = parse_override_deck(path)
            return check_override_deck(deck, regions)
        except OverrideError as exc:
            LOG.error("%s", exc)
            LOG.error("Nothing was run. Fix the deck and check it again.")
            return EXIT_BAD_OVERRIDES

    # ---- load the deck for a real run ------------------------------------ #
    deck: Optional[OverrideDeck] = None
    if args.overrides:
        deck_path = os.path.abspath(args.overrides)
        if not os.path.isfile(deck_path):
            # An explicitly requested deck that is missing must never be
            # silently skipped - the whole run would be against the wrong
            # inputs while looking perfectly normal.
            LOG.error("override file not found: %s", deck_path)
            LOG.error("Remove --overrides to run without manual entries, or "
                      "create the file with --template-overrides.")
            return 1
    elif os.path.isfile(default_deck):
        deck_path = default_deck
        LOG.info("found %s beside model.py - applying it", DEFAULT_OVERRIDES_NAME)
    else:
        deck_path = ""

    if deck_path:
        try:
            deck = parse_override_deck(deck_path)
            deck.validate_area_scopes(
                [af.area for areas in regions.values() for af in areas]
            )
        except OverrideError as exc:
            LOG.error("%s", exc)
            LOG.error("Nothing was run. Fix the deck, or validate it with "
                      "--check-overrides.")
            return EXIT_BAD_OVERRIDES
        LOG.info("manual entries: %s (scenario: %s, %d rule(s))",
                 os.path.basename(deck_path),
                 deck.scenario or "unnamed", len(deck.rules))

    if args.list_regions:
        print(f"Regions found under {root}:\n")
        for name in sorted(regions):
            areas = regions[name]
            print(f"  {name:<14} {len(areas)} area(s): "
                  f"{', '.join(a.area for a in areas)}")
        return 0

    available = sorted(regions)

    # -- figure out which regions were actually asked for ------------------ #
    requested: List[str] = split_region_args(args.region)

    for gname in split_region_args(args.group):
        match = next((g for g in groups if normalise_region(g) == normalise_region(gname)), None)
        if match is None:
            LOG.error(
                "unknown group %r. Defined groups: %s",
                gname, ", ".join(sorted(groups)) or "(none)",
            )
            return 2
        LOG.info("group %s -> %s", match, ", ".join(groups[match]))
        requested.extend(groups[match])

    selected: List[str] = []          # keep the order the user typed things in
    if requested:
        errors: List[str] = []
        for name in requested:
            folders, reason = expand_selector(name, available, groups)
            if folders is None:
                errors.append(reason)
                continue
            for folder in folders:
                if folder not in selected:
                    selected.append(folder)
        if errors:
            for reason in errors:
                LOG.error("%s", reason)
            LOG.error("available regions: %s", ", ".join(available))
            return 2
    else:
        selected = available

    # -- now peel off anything the user wanted excluded ---------------------- #
    for name in split_region_args(args.exclude):
        folder, reason = resolve_region(name, available)
        if folder is None:
            LOG.error("%s", reason)
            LOG.error("available regions: %s", ", ".join(available))
            return 2
        if folder in selected:
            selected.remove(folder)
            LOG.info("excluding %s", folder)

    if not selected:
        LOG.error("no regions left to run after applying --region/--group/--exclude")
        return 2

    LOG.info("regions    : %s", ", ".join(selected))
    regions = {name: regions[name] for name in selected}

    if args.list:
        for region, areas in regions.items():
            print(f"\n{region}  ({len(areas)} area{'' if len(areas) == 1 else 's'})")
            print(f"  -> {region_filename(region, run_date, args.contract_label)}")
            for af in areas:
                extra = ""
                if len(af.candidates) > 1:
                    others = [c for c in af.candidates if c != os.path.basename(af.path)]
                    extra = f"   [also present: {', '.join(others)}]"
                print(f"     {af.area:<24} {os.path.basename(af.path)}{extra}")
        return 0

    written: List[str] = []
    failures = 0
    relaxed = 0
    for region, areas in regions.items():
        try:
            path, sols, load_failures = run_region(
                region, areas, out_dir, run_date,
                allow_shortfall=args.allow_shortfall,
                time_limit=args.time_limit,
                contract_label=args.contract_label,
                auto_relax=not args.no_auto_relax,
                deck=deck,
                compare_saved=args.compare_saved,
            )
        except OverrideError as exc:
            # A deck problem is global, not per-region: stop rather than write
            # some regions with overrides applied and others without.
            LOG.error("%s", exc)
            LOG.error("Nothing further was run. Fix the deck, then validate it "
                      "with --check-overrides.")
            return EXIT_BAD_OVERRIDES
        if path:
            written.append(path)
        failures += sum(1 for s in sols if not s.ok) + load_failures
        relaxed += sum(1 for s in sols if s.ok and s.relaxed)

    LOG.info("=" * 72)
    LOG.info("%d workbook(s) written to %s", len(written), out_dir)
    for w in written:
        LOG.info("  %s", os.path.basename(w))
    if relaxed:
        LOG.warning(
            "%d area(s) only solved after allowing unmet demand / surplus. Their "
            "workbook inputs are over-committed - see the Diagnostics sheets.",
            relaxed,
        )
    if failures:
        LOG.warning("%d area(s) did not solve at all; see the Diagnostics sheets.",
                    failures)
        return 3
    return 4 if relaxed else 0


if __name__ == "__main__":
    sys.exit(main())
